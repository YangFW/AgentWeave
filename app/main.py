from __future__ import annotations

from app.services.workspace_path_manager import default_path_manager

import asyncio
import csv
import hashlib
import hmac
import json
import os
import re
import shutil
import sqlite3
import subprocess
import time
import uuid
import zipfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping
from urllib.parse import urljoin, urlparse

import httpx
from fastapi import FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from app import db
from app.builtin_skill_catalog import BUILTIN_SKILL_CATALOG, get_builtin_skill
from app.schemas import (
    AgentCreate, AgentUpdate, ApprovalRequest, McpServerCreate, McpServerUpdate,
    ExpertInstallRequest, ExpertMemberRetryRequest, ExpertTeamCreate, ExpertTeamRunCreate,
    ExpertTeamUpdate, ExpertTemplateCreate, ExpertTemplateUpdate,
    ConversationSummaryUpdate, KnowledgeBaseCreate, KnowledgeBaseUpdate, KnowledgeDocumentUpload,
    LoopCreate, LoopUpdate, MemoryCreate, MemoryUpdate, ModelConfigCreate,
    ModelConfigUpdate, RemoteInstall, SkillCreate, SkillFileUpdate,
    SkillPathInstall, SkillUpdate,
    PresentationConfigureRequest, ExecutionEngineUpdate, ModelDiscoverRequest,
    CheckpointRestoreRequest, PolicyRuleCreate, PolicyRuleUpdate, TaskCommandRequest,
    TaskCreate, TaskResumeRequest, ToolInvokeRequest, WorkspaceCreate, WorkspaceUpdate,
    UserCreate, UserUpdate, WorkspaceMemberUpdate,
)
from app.seed import seed_agents
from app.services.agent_runtime import AgentRuntime, create_task_record
from app.services.context_service import ContextService, ExecutionScope, MemoryNotFoundError
from app.services.conversation_summary_service import ConversationSummaryConflictError
from app.services.diagnostic_service import DiagnosticService
from app.services.expert_team_service import (
    ExpertConflictError, ExpertNotFoundError, ExpertPermissionError,
    ExpertTeamService, ExpertValidationError,
)
from app.services.event_bus import emit
from app.services import task_queue
from app.services import event_notifications
from app.services import auth_service
from app.services.knowledge_base_service import (
    KnowledgeBaseError,
    KnowledgeBaseNotFoundError,
    KnowledgeBaseService,
)
from app.services.workspace_service import (
    WorkspaceError,
    WorkspaceNotFoundError,
    WorkspaceService,
)
from app.services.mcp_gateway import (
    ARTIFACT_DIR,
    BUILTIN_SERVERS,
    McpGateway,
    ToolError,
    presentation_configuration_status,
    presentation_generation_status,
    resolve_artifact_path,
)
from app.services.loop_scheduler import (
    IdempotencyConflictError, LoopScheduler, create_webhook_event, next_schedule_at,
    serialize_loop, serialize_notification, serialize_run, serialize_trigger_event,
    validate_trigger_config,
)
from app.services.model_gateway import ModelGateway
from app.services.network_policy import (
    env_flag,
    outbound_network_enabled,
    require_outbound_network,
    validate_outbound_http_url,
)
from app.services.policy_engine import PolicyConfigurationError, PolicyEngine, PolicyRule
from app.services.secret_store import secret_store
from app.services.execution_engine_service import (
    ExecutionEngineError,
    get_engine,
    get_engine_row,
    list_engines,
    update_engine,
    test_engine_connection,
    resolve_runtime_env,
)
from app.services.skill_registry import SkillRegistry, referenced_package_files
from app.services.task_state import (
    PublicationConflict,
    RunIntakeClosed,
    StateNotFoundError,
    TaskStateError,
    TaskStateService,
)

BASE_DIR = Path(__file__).resolve().parents[1]
WEB_DIR = BASE_DIR / "web"


def _load_local_env_file() -> None:
    """Load project-local settings when the server is started directly.

    Docker/Compose and shell-based launches can inject environment variables
    themselves, but a direct ``uvicorn`` launch should behave the same way for
    this platform.  Existing process variables win so deployment-level
    settings are never silently overridden by the local file.
    """

    raw_path = str(os.getenv("AGENTNEXUS_ENV_FILE") or ".env.local").strip()
    env_path = Path(raw_path).expanduser()
    if not env_path.is_absolute():
        env_path = BASE_DIR / env_path
    try:
        env_path = env_path.resolve()
        env_path.relative_to(BASE_DIR.resolve())
    except ValueError:
        # Keep the startup path bounded to the project directory.  An invalid
        # override is ignored rather than preventing the offline fallback from
        # starting.
        return
    if not env_path.is_file():
        return
    try:
        lines = env_path.read_text(encoding="utf-8").splitlines()
    except OSError:
        return
    for raw_line in lines:
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if line.startswith("export "):
            line = line[7:].lstrip()
        key, separator, value = line.partition("=")
        key = key.strip()
        if not separator or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", key):
            continue
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
            value = value.replace("\\\\", "\\").replace('\\"', '"').replace("\\'", "'")
        os.environ.setdefault(key, value)


_load_local_env_file()
UPLOAD_DIR = Path(os.getenv("APP_UPLOAD_DIR", str(BASE_DIR / "data" / "uploads")))

app = FastAPI(title="AgentWeave", version="0.1.0")


@app.middleware("http")
async def authentication_middleware(request: Request, call_next: Any) -> Any:
    identity = None
    audit_id = None
    if auth_service.enabled() and request.url.path.startswith("/api/"):
        public = {"/api/health", "/api/readiness", "/api/auth/login", "/api/auth/me"}
        signed_webhook = request.method == 'POST' and re.fullmatch(r'/api/loops/[A-Za-z0-9_-]{2,80}/webhook', request.url.path) is not None
        identity = None if signed_webhook else auth_service.get_session(request.cookies.get(auth_service.SESSION_COOKIE))
        mutating = request.method not in {"GET", "HEAD", "OPTIONS"}
        if mutating or request.url.path.endswith(("/download", "/preview")):
            audit_id = auth_service.start_audit(identity, request.method)
        if mutating and not signed_webhook:
            origin = request.headers.get("origin")
            expected = f"{request.url.scheme}://{request.url.netloc}"
            if (origin is not None and origin != expected) or request.headers.get("sec-fetch-site") == "cross-site":
                if audit_id:
                    auth_service.finish_audit(audit_id, "cross_origin_denied", 403)
                return JSONResponse({"detail": "不允许跨站修改请求"}, status_code=403)
        if request.url.path not in public and not signed_webhook and not identity:
            if audit_id:
                auth_service.finish_audit(audit_id, "authentication_denied", 401)
            return JSONResponse({"detail": "需要登录"}, status_code=401)
        management = {"users", "models", "agents", "skills", "mcp", "policies", "marketplace", "presentation", "execution-engines"}
        resource = request.url.path.split("/")[2]
        if identity:
            identity = {**identity, "request_method": request.method, "request_resource": resource}
        if identity and identity["role"] != "admin" and resource in management and request.method not in {"GET", "HEAD", "OPTIONS"}:
            if audit_id:
                auth_service.finish_audit(audit_id, "admin_required", 403)
            return JSONResponse({"detail": "需要管理员权限"}, status_code=403)
        shared_tables = {'memories': ('memory_entries', 'scope_type'), 'knowledge-bases': ('knowledge_bases', 'visibility')}
        parts = request.url.path.split('/')
        if identity and identity['role'] != 'admin' and mutating and resource in shared_tables and len(parts) > 3:
            table, column = shared_tables[resource]
            shared = db.query_one(f'SELECT {column} AS scope FROM {table} WHERE id=?', (parts[3],))
            if shared and shared['scope'] == 'organization':
                if audit_id:
                    auth_service.finish_audit(audit_id, 'organization_write_denied', 403)
                return JSONResponse({'detail': '修改组织共享资源需要管理员权限'}, status_code=403)
    token = auth_service.current_identity.set(identity)
    try:
        response = await call_next(request)
        if audit_id:
            # 只保存路由模板，不保存 URL 参数、正文、Cookie 或用户输入的路径。
            route = getattr(request.scope.get("route"), "path", "unmatched")
            auth_service.finish_audit(audit_id, route, response.status_code, getattr(request.state, "audit_user_id", None))
        return response
    finally:
        auth_service.current_identity.reset(token)

skill_registry = SkillRegistry()
mcp_gateway = McpGateway()
model_gateway = ModelGateway()
task_state = TaskStateService(auto_init=False)
policy_engine = PolicyEngine(
    http_enabled=outbound_network_enabled() and env_flag("APP_ALLOW_HTTP_POLICY"),
    http_allowlist=[
        item.strip()
        for item in os.getenv("APP_HTTP_POLICY_ALLOWLIST", "").split(",")
        if item.strip()
    ],
)
context_service = ContextService(auto_init=False)
knowledge_service = KnowledgeBaseService(auto_init=False)
workspace_service = WorkspaceService(auto_init=False)
diagnostic_service = DiagnosticService(
    workspace_service=workspace_service,
    knowledge_service=knowledge_service,
    capabilities_provider=lambda: capabilities(),
)


def _is_env_name(value: str) -> bool:
    return bool(re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value or ""))


def _migrate_legacy_model_keys() -> None:
    """Move keys accidentally entered in the old env-name field into encrypted storage."""
    for row in db.query_all("SELECT id, api_key_env, api_key_ciphertext FROM model_configs"):
        legacy = str(row.get("api_key_env") or "").strip()
        if legacy and not _is_env_name(legacy) and not row.get("api_key_ciphertext"):
            db.execute(
                "UPDATE model_configs SET api_key_env = '', api_key_ciphertext = ?, updated_at = ? WHERE id = ?",
                (secret_store.encrypt(legacy), db.utc_now(), row["id"]),
            )


def _remote_install_flag() -> bool:
    return env_flag("APP_ALLOW_REMOTE_INSTALL")


def _remote_install_url(url: str) -> str:
    require_outbound_network("下载链接安装", error_type=ValueError)
    if not _remote_install_flag():
        raise ValueError("下载链接安装尚未开启，请由管理员设置 APP_ALLOW_REMOTE_INSTALL=true 后重启平台")
    if urlparse(url).scheme.lower() != "https":
        raise ValueError("下载链接必须使用 HTTPS")
    return validate_outbound_http_url(
        url,
        capability="下载链接安装",
        allowlist_env="APP_REMOTE_INSTALL_HOST_ALLOWLIST",
        require_allowlist=True,
        allow_query=True,
        error_type=ValueError,
    )


async def _download_remote_install(url: str, max_bytes: int) -> tuple[bytes, str]:
    checked = _remote_install_url(url)
    async with httpx.AsyncClient(timeout=30, follow_redirects=False) as client:
        response = await client.get(checked, headers={"User-Agent": "AgentWeave/0.1"})
        response.raise_for_status()
    if len(response.content) > max_bytes:
        raise ValueError("远程安装包超过大小限制")
    filename = Path(urlparse(str(response.url)).path).name or "download"
    return response.content, filename


def _skill_package_from_bytes(
    raw: bytes, filename: str
) -> tuple[dict[str, bytes], str]:
    if len(raw) > 2 * 1024 * 1024:
        raise ValueError("Skill 安装包不能超过 2MB")
    if filename.lower().endswith(".zip") or raw.startswith(b"PK\x03\x04"):
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            files: dict[str, bytes] = {}
            for info in archive.infolist():
                if info.is_dir():
                    continue
                path = info.filename.replace("\\", "/")
                if path.startswith("/") or ".." in Path(path).parts:
                    raise ValueError("ZIP 包含不安全路径")
                if (info.external_attr >> 16) & 0o170000 == 0o120000:
                    raise ValueError("ZIP 不允许包含符号链接")
                files[path] = archive.read(info)
            candidates = [n for n in files if Path(n).name == "SKILL.md"]
            if not candidates:
                raise ValueError("ZIP 中未找到 SKILL.md")
            if len(candidates) > 1:
                raise ValueError("ZIP 中只能包含一个 Skill 包")
            fallback = Path(candidates[0]).parent.name or Path(filename).stem
            return files, fallback
    raw.decode("utf-8")
    fallback = (
        Path(filename).stem
        if filename.lower().endswith(".md")
        else "downloaded_skill"
    )
    return {"SKILL.md": raw}, fallback


def _install_skill_bytes(raw: bytes, filename: str) -> dict[str, Any]:
    files, fallback = _skill_package_from_bytes(raw, filename)
    return skill_registry.install_package(files, fallback_id=fallback)


async def _load_skill_remote_url(url: str) -> dict[str, Any]:
    """Download and validate a Skill package without mutating the registry."""

    raw, filename = await _download_remote_install(url, 2 * 1024 * 1024)
    if not filename.lower().endswith(".zip") and not raw.startswith(b"PK\x03\x04"):
        content = raw.decode("utf-8")
        files: dict[str, bytes] = {"SKILL.md": raw}
        total = len(raw)
        for relative in referenced_package_files(content):
            stored_relative = relative
            try:
                child, _ = await _download_remote_install(urljoin(url, relative), 1024 * 1024)
            except (ValueError, httpx.HTTPError):
                alternate = str(Path(relative).with_name(Path(relative).name.lower())).replace("\\", "/")
                if alternate == relative:
                    continue
                try:
                    child, _ = await _download_remote_install(urljoin(url, alternate), 1024 * 1024)
                    stored_relative = alternate
                except (ValueError, httpx.HTTPError):
                    continue
            total += len(child)
            if total > 2 * 1024 * 1024:
                raise ValueError("Skill 包及引用文件合计超过 2MB")
            files[stored_relative] = child
        return {"files": files, "fallback_id": Path(filename).stem}
    files, fallback = _skill_package_from_bytes(raw, filename)
    return {"files": files, "fallback_id": fallback}


async def _install_skill_remote_url(url: str) -> dict[str, Any]:
    package = await _load_skill_remote_url(url)
    return skill_registry.install_package(
        package["files"], fallback_id=str(package.get("fallback_id") or "")
    )


async def _load_mcp_remote_url(url: str) -> Any:
    """Download and validate MCP JSON without changing installed servers."""

    raw, _ = await _download_remote_install(url, 1024 * 1024)
    return json.loads(raw.decode("utf-8"))


async def _install_mcp_remote_url(url: str) -> list[dict[str, Any]]:
    return mcp_gateway.import_config(await _load_mcp_remote_url(url))


runtime = AgentRuntime(
    skill_registry,
    mcp_gateway,
    model_gateway,
    skill_url_loader=_load_skill_remote_url,
    mcp_url_loader=_load_mcp_remote_url,
    task_state=task_state,
    policy_engine=policy_engine,
    context_service=context_service,
    knowledge_service=knowledge_service,
)
loop_scheduler = LoopScheduler(runtime)
expert_team_service = ExpertTeamService(runtime, task_state=task_state)
_runtime_tasks: set[asyncio.Task[Any]] = set()


def _schedule_runtime(
    task_id: str,
    run_id: str | None = None,
    *,
    activation_result: dict[str, Any] | None = None,
) -> asyncio.Task[Any]:
    if activation_result is None and task_queue.enabled():
        async def enqueue_runtime() -> None:
            # 新运行已原子写入 outbox；兼容旧运行的恢复调度。
            db.execute("INSERT OR IGNORE INTO dispatch_outbox(run_id,task_id,payload_json,created_at) VALUES(?,?,?,?)", (run_id, task_id, json.dumps({"task_id": task_id, "run_id": run_id}), db.utc_now()))

        background = asyncio.create_task(enqueue_runtime())
        _runtime_tasks.add(background)
        background.add_done_callback(_runtime_tasks.discard)
        return background
    continuation = (
        runtime.run_task(
            task_id,
            run_id=run_id,
            activation_result=activation_result,
        )
        if activation_result is not None
        else runtime.run_task(task_id, run_id=run_id)
    )
    background = asyncio.create_task(continuation)
    _runtime_tasks.add(background)
    background.add_done_callback(_runtime_tasks.discard)
    return background


def _schedule_team_run(team_run_id: str) -> asyncio.Task[Any]:
    if task_queue.enabled():
        async def enqueue_team() -> None:
            row = db.query_one("SELECT parent_task_id,parent_run_id FROM team_runs WHERE id=?", (team_run_id,))
            if not row:
                raise RuntimeError("专家团运行不存在")
            db.execute("INSERT OR IGNORE INTO dispatch_outbox(run_id,task_id,payload_json,created_at) VALUES(?,?,?,?)", (row['parent_run_id'], row['parent_task_id'], db.json_dumps({'task_id': row['parent_task_id'], 'run_id': row['parent_run_id'], 'kind': 'team', 'team_run_id': team_run_id}), db.utc_now()))
        background = asyncio.create_task(enqueue_team())
        _runtime_tasks.add(background)
        background.add_done_callback(_runtime_tasks.discard)
        return background
    background = asyncio.create_task(expert_team_service.run_team(team_run_id))
    _runtime_tasks.add(background)
    background.add_done_callback(_runtime_tasks.discard)
    return background


def _schedule_member_retry(
    team_run_id: str, member_run_id: str, scope: ExecutionScope
) -> asyncio.Task[Any]:
    if task_queue.enabled():
        row = db.query_one('SELECT parent_task_id,parent_run_id FROM team_runs WHERE id=?', (team_run_id,))
        job_id = 'member_retry_' + uuid.uuid4().hex
        message = {'kind':'member_retry', 'job_id':job_id, 'dispatch_id':job_id, 'team_run_id':team_run_id, 'member_run_id':member_run_id, 'task_id':row['parent_task_id'], 'run_id':row['parent_run_id'], 'scope':{'organization_id':scope.organization_id, 'workspace_id':scope.workspace_id, 'user_id':scope.user_id}}
        try:
            db.execute('INSERT INTO member_retry_dispatch(job_id,team_run_id,payload_json,created_at) VALUES(?,?,?,?)', (job_id,team_run_id,db.json_dumps(message),db.utc_now()))
        except sqlite3.IntegrityError as exc:
            raise HTTPException(status_code=409, detail='该专家团已有成员重试正在排队或运行') from exc
        async def accepted():
            return {'accepted':True, 'job_id':job_id}
        return asyncio.create_task(accepted())
    background = asyncio.create_task(
        expert_team_service.retry_member(team_run_id, member_run_id, scope)
    )
    _runtime_tasks.add(background)
    background.add_done_callback(_runtime_tasks.discard)
    return background


def _reload_policy_rules() -> None:
    rules: list[dict[str, Any]] = []
    for row in db.query_all("SELECT rule_json FROM policy_rules WHERE enabled = 1 ORDER BY priority DESC, id"):
        value = db.json_loads(row.get("rule_json"), {})
        if isinstance(value, dict):
            rules.append(value)
    env_rules = os.getenv("APP_POLICY_RULES_JSON", "").strip()
    if env_rules:
        parsed = json.loads(env_rules)
        if not isinstance(parsed, list):
            raise ValueError("APP_POLICY_RULES_JSON 必须是规则数组")
        rules.extend(item for item in parsed if isinstance(item, dict))
    policy_engine.set_rules(rules)


def _fail_running_nodes(run_id: str, reason: str) -> None:
    for node in task_state.list_nodes(run_id):
        if node["status"] != "running":
            continue
        try:
            task_state.fail_node(
                node["id"],
                {"message": reason, "error_type": "ServiceRestart"},
                metadata={"interrupted": True},
            )
        except TaskStateError:
            continue


_ORCHESTRATED_EXECUTOR_TYPES = frozenset(
    {"team", "team_member", "team_supervisor", "automation"}
)


def _persisted_run_executor(run: Mapping[str, Any]) -> tuple[str, str, bool]:
    """Resolve durable execution ownership without trusting startup memory.

    The Task row is the authoritative dispatcher contract.  Run metadata is a
    second durable fence: an old or partially migrated row that advertises an
    orchestrated owner in either place is never handed to the ordinary Agent
    runtime.  ``consistent`` is false when both projections are populated but
    disagree, allowing the owning orchestrator to quarantine the row.
    """

    task = db.query_one(
        "SELECT executor_type, executor_id FROM tasks WHERE id = ?",
        (str(run.get("task_id") or ""),),
    ) or {}
    task_type = str(task.get("executor_type") or "agent").strip() or "agent"
    task_id = str(task.get("executor_id") or "").strip()
    metadata = run.get("metadata") if isinstance(run.get("metadata"), Mapping) else {}
    metadata_type = str((metadata or {}).get("executor_type") or "").strip()
    metadata_id = str((metadata or {}).get("executor_id") or "").strip()
    advertised_types = {item for item in (task_type, metadata_type) if item}
    orchestrated = advertised_types & _ORCHESTRATED_EXECUTOR_TYPES
    executor_type = task_type
    executor_id = task_id
    if task_type == "agent" and orchestrated:
        executor_type = sorted(orchestrated)[0]
        executor_id = metadata_id
    consistent = not (
        metadata_type
        and task_type != metadata_type
        and (task_type in _ORCHESTRATED_EXECUTOR_TYPES or metadata_type in _ORCHESTRATED_EXECUTOR_TYPES)
    )
    if task_id and metadata_id and task_id != metadata_id:
        consistent = False
    return executor_type, executor_id, consistent


def _ordinary_runtime_owned(run: Mapping[str, Any]) -> bool:
    if (run.get("metadata") or {}).get("dispatch_backend") == "redis":
        return False
    executor_type, _, consistent = _persisted_run_executor(run)
    return consistent and executor_type == "agent"


def _approval_command_for_restart(
    task_id: str,
    run_id: str,
    result: dict[str, Any],
) -> dict[str, Any] | None:
    """Find the one durable non-policy approval continuation, if decided."""

    commands = task_state.list_commands(
        task_id=task_id,
        run_id=run_id,
        command_types=["approval"],
        limit=100,
    )
    active = [
        item for item in commands if item.get("status") in {"queued", "claimed"}
    ]
    if active:
        return active[0]

    proof_command_ids: set[str] = set()
    for key in (
        "skill_recommendation_decision",
        "generic_approval_decision",
        "approval_decision",
    ):
        proof = result.get(key)
        if isinstance(proof, dict) and str(proof.get("command_id") or ""):
            proof_command_ids.add(str(proof["command_id"]))
    for item in commands:
        if item.get("status") != "completed":
            continue
        command_result = item.get("result") or {}
        if str(item.get("id") or "") in proof_command_ids or (
            isinstance(command_result, dict)
            and str(command_result.get("action") or "")
            in {"install_recommended_skill", "generic_approval"}
        ):
            return item
    return None


def _prepare_waiting_approval_recovery(
    preserve_task_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    """Reconcile durable approval input without disturbing undecided waits.

    Policy decisions become immutable proof first; their now-running attempt is
    then handled by normal checkpoint recovery.  Recommendation and generic
    decisions remain on the same waiting Run and are returned for safe
    continuation scheduling.  A pending cancel always wins at startup.
    """

    continuations: list[dict[str, Any]] = []
    preserved = preserve_task_ids or set()
    for waiting_run in task_state.list_runs(status="waiting_approval", limit=1000):
        if (waiting_run.get("metadata") or {}).get("dispatch_backend") == "redis":
            continue
        task_id = str(waiting_run["task_id"])
        if task_id in preserved:
            continue
        task = db.query_one(
            "SELECT status, result_json FROM tasks WHERE id = ?", (task_id,)
        )
        if not task or str(task.get("status") or "") != "waiting_approval":
            continue
        if task_state.is_cancel_requested(task_id, run_id=waiting_run["id"]):
            task_state.commit_cancellation(
                task_id=task_id,
                run_id=waiting_run["id"],
                result={
                    "cancelled": True,
                    "recovered_after_restart": True,
                },
            )
            continue

        result = db.json_loads(task.get("result_json"), {})
        metadata = waiting_run.get("metadata") or {}
        pending_policy = (
            metadata.get("pending_policy_approval")
            if isinstance(metadata, dict)
            else None
        )
        is_policy = result.get("pending_action") == "policy_approval" or isinstance(
            pending_policy, dict
        )
        if is_policy:
            approval_id = str(
                result.get("policy_approval_id")
                or (
                    pending_policy.get("approval_id")
                    if isinstance(pending_policy, dict)
                    else ""
                )
                or ""
            )
            # An undecided approval is a healthy durable wait.  Do not fail the
            # Run and do not recreate its public approval event.
            if not approval_id:
                continue
            decision = task_state.commit_policy_approval_decision(
                task_id=task_id,
                run_id=waiting_run["id"],
                approval_id=approval_id,
                worker_id=f"restart-recovery:{waiting_run['id']}",
            )
            if decision is None:
                continue
            # The transaction above changed Task/Run to running and persisted
            # the exact command proof.  _recover_interrupted_runs will now
            # create a checkpoint-bound retry without asking again.
            continue

        command = _approval_command_for_restart(
            task_id, str(waiting_run["id"]), result
        )
        if command is None:
            continue
        payload = command.get("payload") or {}
        command_result = command.get("result") or {}
        approved_value = (
            command_result.get("approved")
            if isinstance(command_result, dict)
            and isinstance(command_result.get("approved"), bool)
            else payload.get("approved")
        )
        if not isinstance(approved_value, bool):
            raise TaskStateError(
                "Persisted approval continuation is missing a boolean decision"
            )
        continuations.append(
            {
                "task_id": task_id,
                "run_id": str(waiting_run["id"]),
                "command_id": str(command["id"]),
                "approved": approved_value,
                "note": str(payload.get("note") or ""),
            }
        )
    return continuations


def _recover_interrupted_runs(
    exclude_task_ids: set[str] | None = None,
    preserve_task_ids: set[str] | None = None,
) -> list[dict[str, Any]]:
    """Recover only runs durably owned by the ordinary Agent runtime.

    ``exclude_task_ids`` are interrupted automation attempts and are closed as
    failed. ``preserve_task_ids`` remains a compatibility quarantine for
    callers, but ownership is always reloaded from ``tasks.executor_type`` /
    ``executor_id`` and the Run metadata.  An orchestrated row can therefore
    never become an ordinary Agent run merely because a transient startup list
    was empty on a later restart.
    """
    excluded = exclude_task_ids or set()
    preserved = preserve_task_ids or set()
    protected = excluded | preserved

    # Early platform revisions could commit a terminal Task projection before
    # closing its Run. Those split-state rows are not interrupted work: the
    # user already received a terminal result, so rerunning them could duplicate
    # side effects or overwrite published artifacts. Reconcile every such Run
    # transactionally before ordinary restart recovery decides what to resume.
    legacy_terminal_runs = db.query_all(
        """
        SELECT r.id
        FROM task_runs AS r
        JOIN tasks AS t ON t.id = r.task_id
        WHERE r.status IN ('running', 'paused', 'waiting_approval')
          AND t.status IN ('completed', 'failed', 'cancelled')
        ORDER BY r.created_at, r.attempt, r.id
        """
    )
    for legacy_run in legacy_terminal_runs:
        task_state.reconcile_legacy_terminal_projection(str(legacy_run["id"]))

    for task_id in excluded:
        task = db.query_one("SELECT status FROM tasks WHERE id = ?", (task_id,))
        if task and task.get("status") not in {"completed", "failed", "cancelled"}:
            db.update_task_status(
                task_id, "failed",
                result={"error": "自动化尝试因平台服务重启而中断", "error_type": "ServiceRestart"},
            )
    recovered: list[dict[str, Any]] = []
    interrupted_runs = [
        *task_state.list_runs(status="running", limit=1000),
        *task_state.list_runs(status="paused", limit=1000),
    ]
    for old_run in interrupted_runs:
        task = db.query_one("SELECT id FROM tasks WHERE id = ?", (old_run["task_id"],))
        if not task:
            continue
        if not _ordinary_runtime_owned(old_run):
            continue
        if old_run["task_id"] in preserved:
            continue
        if old_run["task_id"] in excluded:
            _fail_running_nodes(old_run["id"], "平台服务重启，旧执行尝试已中断")
            task_state.finish_run(
                old_run["id"],
                status="failed",
                error={"message": "自动化尝试因平台服务重启而中断", "error_type": "ServiceRestart"},
                metadata={"interrupted": True, "automation_run": True},
            )
            db.update_task_status(
                old_run["task_id"], "failed",
                result={"error": "自动化尝试因平台服务重启而中断", "error_type": "ServiceRestart"},
            )
            continue
        if task_state.is_cancel_requested(
            old_run["task_id"], run_id=old_run["id"]
        ):
            task_state.commit_cancellation(
                task_id=old_run["task_id"],
                run_id=old_run["id"],
                result={
                    "cancelled": True,
                    "recovered_after_restart": True,
                },
            )
            continue
        recovery = task_state.recover_interrupted_attempt(old_run["id"])
        recovered.append(recovery["run"])

    # Queued runs survive a restart unchanged. Legacy queued tasks without a
    # task_run receive one before they are scheduled.
    all_queued_runs = task_state.list_runs(status="queued", limit=1000)
    for queued in all_queued_runs:
        if queued["task_id"] in excluded:
            task_state.finish_run(
                queued["id"], status="cancelled",
                error={"message": "自动化尝试因平台服务重启而中断", "error_type": "ServiceRestart"},
                metadata={"interrupted": True, "automation_run": True},
            )
    queued_runs = [
        item
        for item in all_queued_runs
        if item["task_id"] not in protected and _ordinary_runtime_owned(item)
    ]
    known_queued_tasks = {item["task_id"] for item in queued_runs}
    tasks_with_runs = {
        item["task_id"] for item in task_state.list_runs(limit=10_000)
    }
    for task in db.query_all(
        "SELECT id, executor_type FROM tasks WHERE status = 'running'"
    ):
        if task["id"] in preserved:
            continue
        if str(task.get("executor_type") or "agent") != "agent":
            continue
        if task["id"] in excluded:
            db.update_task_status(
                task["id"], "failed",
                result={"error": "自动化尝试因平台服务重启而中断", "error_type": "ServiceRestart"},
            )
            continue
        if task["id"] in tasks_with_runs:
            continue
        db.update_task_status(task["id"], "queued")
        legacy_run = task_state.create_run(
            task["id"], metadata={"legacy_interrupted_task": True, "recovered_after_restart": True}
        )
        queued_runs.append(legacy_run)
        known_queued_tasks.add(task["id"])
        db.insert_event(
            task["id"],
            "recovery_scheduled",
            "已安排旧任务恢复",
            "检测到升级前遗留的运行中任务，将从任务起点重新执行。",
            {"run_id": legacy_run["id"]},
        )
    for task in db.query_all(
        "SELECT id, executor_type FROM tasks WHERE status = 'queued'"
    ):
        if (
            str(task.get("executor_type") or "agent") == "agent"
            and task["id"] not in protected
            and task["id"] not in known_queued_tasks
            and task["id"] not in tasks_with_runs
        ):
            queued_runs.append(task_state.create_run(task["id"], metadata={"legacy_task": True}))
    by_id = {item["id"]: item for item in [*queued_runs, *recovered]}
    return list(by_id.values())


@app.on_event("startup")
async def on_startup() -> None:
    db.init_db()
    auth_service.init_schema()
    auth_service.validate_deployment_admin()
    workspace_service.init_schema()
    context_service.init_schema()
    knowledge_service.init_schema()
    task_state.init_schema()
    if task_queue.enabled():
        dispatcher = asyncio.create_task(task_queue.dispatch_forever())
        _runtime_tasks.add(dispatcher)
        dispatcher.add_done_callback(_runtime_tasks.discard)
        relay = asyncio.create_task(event_notifications.relay_forever())
        _runtime_tasks.add(relay)
        relay.add_done_callback(_runtime_tasks.discard)
    runtime.tool_effect_journal.init_schema()
    all_runs = db.query_all("SELECT id,task_id,metadata_json FROM task_runs")
    worker_run_ids = {
        row["id"] for row in all_runs
        if db.json_loads(row["metadata_json"], {}).get("dispatch_backend") == "redis"
    }
    worker_task_ids = {row['task_id'] for row in all_runs if row['id'] in worker_run_ids}
    task_parents = db.query_all('SELECT id,parent_task_id FROM tasks')
    while True:
        children = {row['id'] for row in task_parents if row['parent_task_id'] in worker_task_ids}
        if children.issubset(worker_task_ids):
            break
        worker_task_ids.update(children)
    worker_run_ids.update(row['id'] for row in all_runs if row['task_id'] in worker_task_ids)
    runtime.tool_effect_journal.recover_interrupted_executions(
        reason="service_restart", exclude_run_ids=worker_run_ids,
    )
    worker_loop_ids = {row['loop_id'] for row in db.query_all("SELECT loop_id FROM automation_dispatch WHERE status IN ('queued','running')")}
    interrupted_loop_tasks = loop_scheduler.recover_interrupted_runs(exclude_loop_ids=worker_loop_ids)
    queued_team_runs = expert_team_service.reconcile_interrupted_orchestrated_runs(exclude_task_ids=worker_task_ids)
    queued_team_task_ids = {item["parent_task_id"] for item in queued_team_runs}
    _migrate_legacy_model_keys()
    skill_registry.load_builtin_skills()
    mcp_gateway.seed_builtin_servers()
    seed_agents()
    _reload_policy_rules()
    loop_scheduler.start()
    approval_continuations = _prepare_waiting_approval_recovery(
        queued_team_task_ids
    )
    for run in _recover_interrupted_runs(interrupted_loop_tasks, queued_team_task_ids):
        recovery_activation = (run.get("metadata") or {}).get(
            "recovery_activation_result"
        )
        if isinstance(recovery_activation, dict):
            _schedule_runtime(
                run["task_id"],
                run["id"],
                activation_result=recovery_activation,
            )
        else:
            _schedule_runtime(run["task_id"], run["id"])
    for continuation in approval_continuations:
        background = asyncio.create_task(
            _resume_after_approval_safely(
                continuation["task_id"],
                continuation["approved"],
                continuation["note"],
                continuation["command_id"],
            )
        )
        _runtime_tasks.add(background)
        background.add_done_callback(_runtime_tasks.discard)
    for team_run in queued_team_runs:
        _schedule_team_run(team_run["id"])


@app.on_event("shutdown")
async def on_shutdown() -> None:
    await loop_scheduler.stop()
    pending = [item for item in _runtime_tasks if not item.done()]
    for item in pending:
        item.cancel()
    if pending:
        await asyncio.gather(*pending, return_exceptions=True)


@app.get("/api/health")
def health() -> dict[str, Any]:
    result: dict[str, Any] = {"ok": True, "name": "AgentWeave", "product": "AgentWeave", "display_name": "智织"}
    if task_queue.enabled():
        try:
            import redis

            client = redis.Redis.from_url(task_queue.redis_url(), socket_timeout=1)
            client.ping()
            client.close()
            result["redis"] = {"ok": True}
        except Exception:
            result["ok"] = False
            result["redis"] = {"ok": False}
    else:
        result["redis"] = {"ok": False, "configured": False}
    return result


@app.post("/api/auth/login")
def auth_login(payload: dict[str, str], response: Response, request: Request) -> dict[str, Any]:
    if not auth_service.enabled():
        return {"authenticated": False, "enabled": False}
    peer = request.client.host if request.client else "unknown"
    if not auth_service.reserve_login_attempt(str(payload.get("username") or ""), peer):
        raise HTTPException(status_code=429, detail="登录尝试过于频繁，请稍后重试", headers={"Retry-After": "60"})
    result = auth_service.login(str(payload.get("username") or ""), str(payload.get("password") or ""))
    if not result:
        raise HTTPException(status_code=401, detail="用户名或密码错误")
    token, user = result
    request.state.audit_user_id = user["user_id"]
    response.set_cookie(auth_service.SESSION_COOKIE, token, httponly=True, secure=request.url.scheme == "https", samesite="lax", max_age=86400)
    return {"authenticated": True, "user": user}


@app.get("/api/readiness")
def readiness() -> JSONResponse:
    status = task_queue.readiness()
    try:
        db.query_one("SELECT 1 FROM task_runs LIMIT 1")
        status["database"] = True
    except sqlite3.Error:
        status.update(ready=False, database=False)
    return JSONResponse(status, status_code=200 if status["ready"] else 503)


@app.post("/api/auth/logout")
def auth_logout(request: Request, response: Response) -> dict[str, bool]:
    auth_service.logout(request.cookies.get(auth_service.SESSION_COOKIE))
    response.delete_cookie(auth_service.SESSION_COOKIE)
    return {"authenticated": False}


@app.get("/api/auth/me")
def auth_me(request: Request) -> dict[str, Any]:
    session = auth_service.get_session(request.cookies.get(auth_service.SESSION_COOKIE))
    return {"authenticated": bool(session), "enabled": auth_service.enabled(), "user": session or {}}


def _require_admin(request: Request) -> dict[str, Any]:
    if not auth_service.enabled():
        raise HTTPException(status_code=403, detail="请先启用认证")
    session = auth_service.get_session(request.cookies.get(auth_service.SESSION_COOKIE))
    if not session:
        raise HTTPException(status_code=401, detail="需要登录")
    if session["role"] != "admin":
        raise HTTPException(status_code=403, detail="需要管理员权限")
    return session


@app.get("/api/users")
def list_users(request: Request) -> list[dict[str, Any]]:
    _require_admin(request)
    return auth_service.list_users()


@app.get("/api/audit-events")
def list_audit_events(request: Request, limit: int = 100) -> list[dict[str, Any]]:
    _require_admin(request)
    return db.query_all("SELECT * FROM audit_events ORDER BY id DESC LIMIT ?", (max(1, min(limit, 500)),))


@app.post("/api/users", status_code=201)
def create_user(payload: UserCreate, request: Request) -> dict[str, Any]:
    _require_admin(request)
    try:
        return auth_service.create_user(payload.username, payload.password, payload.role)
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="用户名已存在") from exc


@app.put("/api/users/{user_id}")
def update_user(user_id: str, payload: UserUpdate, request: Request) -> dict[str, Any]:
    _require_admin(request)
    try:
        return auth_service.update_user(user_id, payload.model_dump(exclude_none=True))
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


def _presentation_env_file() -> Path:
    """Resolve the local env file used by the one-click PPTX setup."""

    raw = str(os.getenv("AGENTNEXUS_ENV_FILE") or ".env.local").strip()
    candidate = Path(raw).expanduser()
    if not candidate.is_absolute():
        candidate = BASE_DIR / candidate
    resolved = candidate.resolve()
    try:
        resolved.relative_to(BASE_DIR.resolve())
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="配置文件必须位于项目目录内") from exc
    return resolved


def _persist_env_values(values: Mapping[str, str]) -> Path:
    path = _presentation_env_file()
    text = path.read_text(encoding="utf-8") if path.exists() else ""
    if text and not text.endswith("\n"):
        text += "\n"
    for key, value in values.items():
        escaped = str(value).replace("\\", "\\\\").replace('"', '\\"')
        line = f'{key}="{escaped}"'
        pattern = re.compile(rf"(?m)^\s*{re.escape(key)}\s*=.*$")
        if pattern.search(text):
            text = pattern.sub(line, text, count=1)
        else:
            text += f"{line}\n"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def _resolve_local_binary(value: str, *, label: str) -> str:
    raw = str(value or "").strip()
    if not raw:
        raise HTTPException(status_code=400, detail=f"请填写{label}")
    candidate = Path(raw).expanduser()
    resolved = str(candidate.resolve()) if candidate.is_file() else shutil.which(raw) or ""
    if not resolved or not Path(resolved).is_file():
        raise HTTPException(status_code=400, detail=f"找不到{label}：{raw}")
    return resolved


@app.get("/api/presentation/configuration")
def get_presentation_configuration() -> dict[str, Any]:
    return presentation_configuration_status()


@app.post("/api/presentation/configure")
def configure_presentation(payload: PresentationConfigureRequest) -> dict[str, Any]:
    if not payload.confirmed:
        raise HTTPException(status_code=400, detail="请先确认将配置写入本机 .env.local")
    mode = str(payload.mode or "python").strip().lower()
    if mode == "python":
        try:
            import pptx  # noqa: F401
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail="平台内置 Python PPTX 组件不可用，请在当前 .venv 中安装 python-pptx 后重试",
            ) from exc
        values = {"APP_PPTX_GENERATOR": "python"}
        os.environ.update(values)
    elif mode == "artifact_tool":
        node_binary = _resolve_local_binary(payload.node_binary, label="Node.js 可执行文件")
        entrypoint = Path(str(payload.entrypoint or "").strip()).expanduser()
        if not entrypoint.is_absolute():
            entrypoint = (BASE_DIR / entrypoint).resolve()
        else:
            entrypoint = entrypoint.resolve()
        if not entrypoint.is_file():
            raise HTTPException(status_code=400, detail=f"找不到 Artifact Tool 入口文件：{entrypoint}")
        check = subprocess.run(
            [
                node_binary,
                "--input-type=module",
                "-e",
                "const m=await import(process.argv[1]); if(!m.Presentation || !m.PresentationFile) process.exit(2);",
                str(entrypoint),
            ],
            capture_output=True,
            text=True,
            timeout=20,
            check=False,
        )
        if check.returncode != 0:
            raise HTTPException(
                status_code=400,
                detail="Artifact Tool 入口无法加载，或未导出 Presentation 与 PresentationFile",
            )
        values = {
            "APP_PPTX_GENERATOR": "artifact_tool",
            "APP_NODE_BINARY": node_binary,
            "APP_ARTIFACT_TOOL_ENTRYPOINT": str(entrypoint),
        }
        os.environ.update(values)
    else:
        raise HTTPException(status_code=400, detail="PPTX 生成模式只能是 python 或 artifact_tool")
    env_file = _persist_env_values(values)
    status = presentation_configuration_status()
    return {
        "ok": bool(status.get("configured")),
        "configuration": status,
        "env_file": str(env_file),
        "restart_required": False,
        "message": "PPTX 生成器已配置，当前服务立即生效；以后重启平台也会保留此配置。",
    }


@app.get("/api/capabilities")
def capabilities() -> dict[str, Any]:
    presentation = presentation_generation_status()
    presentation_setup = presentation_configuration_status()
    outbound_enabled = outbound_network_enabled()
    return {
        "outbound_network": {"supported": True, "enabled": outbound_enabled},
        "file_upload": {
            "supported": True,
            "max_mb": int(os.getenv("APP_MAX_UPLOAD_MB", "20")),
            "text_extraction": [
                "txt", "md", "csv", "json", "yaml", "code",
                "docx", "xlsx", "pptx", "pdf",
            ],
            "max_files_per_task": runtime.ATTACHMENT_MAX_FILES,
            "max_chars_per_file": runtime.ATTACHMENT_MAX_FILE_CHARS,
            "max_context_chars": runtime.ATTACHMENT_MAX_CONTEXT_CHARS,
        },
        "web_search": {
            "supported": True,
            "enabled": outbound_enabled and env_flag("APP_ALLOW_WEB_SEARCH"),
            "configured": bool(os.getenv("TAVILY_API_KEY") or os.getenv("BRAVE_SEARCH_API_KEY")),
            "provider": "tavily/brave",
        },
        "stdio_mcp": {"supported": True, "enabled": os.getenv("APP_ALLOW_STDIO_MCP", "false").lower() in {"1", "true", "yes"}},
        "remote_mcp": {"supported": True, "enabled": outbound_enabled and env_flag("APP_ALLOW_REMOTE_MCP")},
        "http_tools": {"supported": True, "enabled": outbound_enabled and env_flag("APP_ALLOW_HTTP_TOOLS")},
        "remote_install": {"supported": True, "enabled": outbound_enabled and _remote_install_flag()},
        "direct_api_key": {"supported": True, "encrypted": True, "storage": "local"},
        "models": ["deterministic", "openai", "openai_compatible"],
        "workspaces": {
            "supported": True,
            "default_id": "default",
            "scope": "organization",
            "current_user": "local-user",
        },
        "document_output": {
            "formats": ["markdown", "docx", "pdf", "xlsx", "csv", "html"],
            "optional_formats": ["pptx"],
            "pptx_configured": bool(presentation.get("configured")),
            "pptx_reason": str(presentation.get("reason") or ""),
            "pptx_setup": {
                "mode": presentation_setup.get("mode"),
                "node_binary": presentation_setup.get("node_binary"),
                "entrypoint": presentation_setup.get("entrypoint"),
                "native_python_available": bool(presentation_setup.get("native_python_available")),
            },
        },
        "memory": {"supported": True, "scopes": ["organization", "workspace", "user", "agent", "conversation"], "revision_history": True, "conversation_summary": {"automatic": True, "viewable": True, "editable": True, "deletable": True}},
        "knowledge_base": {
            "supported": True,
            "scopes": ["private", "workspace", "organization"],
            "retrieval": "keyword",
            "runtime_injection": True,
            "indexed_upload_formats": ["txt", "md", "csv", "json", "yaml", "html", "docx", "xlsx", "pptx", "pdf"],
        },
        "expert_teams": {
            "supported": True,
            "template_installation": True,
            "parallel_members": True,
            "isolated_member_context": True,
            "supervisor_aggregation": True,
            "single_member_retry": True,
        },
        "automation": {
            "supported": True,
            "triggers": ["manual", "interval", "cron", "once", "webhook"],
            "persistent_history": True,
            "signed_webhooks": True,
            "idempotency": True,
            "notifications": True,
            "structured_state_diff": True,
            "legacy_api": "/api/loops",
        },
        "policy_hooks": {
            "supported": True,
            "handlers": ["builtin_rule", "http"],
            "http_enabled": policy_engine.http_enabled,
            "arbitrary_shell": False,
        },
    }


@app.get("/api/diagnostics")
def diagnostics(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    return diagnostic_service.run(
        organization_id=organization_id,
        workspace_id=workspace_id,
        user_id=user_id,
    )


def _skill_marketplace_plan(item: dict[str, Any], installed: dict[str, Any] | None) -> dict[str, Any]:
    content = str(item.get("content") or "")
    required_mcps = list(item.get("required_mcps") or [])
    if not required_mcps:
        for line in content.splitlines()[:20]:
            if line.strip().startswith("required_mcps:"):
                raw = line.split(":", 1)[1].strip().strip("[]")
                required_mcps = [part.strip().strip("'\"") for part in raw.split(",") if part.strip()]
                break
    writes_documents = any(keyword in content for keyword in ("生成 Word", "生成 DOCX", "生成 Word、PDF", "生成 Word、PDF 或 Markdown", "调用报告工具"))
    return {
        "method": "builtin_catalog",
        "requires_approval": False,
        "will_create": [] if installed else ["skill"],
        "will_enable": ["skill"] if installed and not installed.get("enabled") else [],
        "required_mcps": required_mcps,
        "permissions": {
            "reads_uploaded_files": True,
            "writes_artifacts": writes_documents,
            "runs_local_process": False,
            "uses_network": False,
        },
        "impact": "安装后会进入技能中心并可被普通模式自动匹配；不会自动执行包内脚本。",
        "post_install": "在“技能中心”查看、停用、编辑或导出，也可以绑定到指定智能体。",
    }


def _mcp_marketplace_plan(server: dict[str, Any], installed: dict[str, Any] | None) -> dict[str, Any]:
    kind = str(server.get("kind") or "builtin")
    tools = [tool for tool in server.get("tools", []) if isinstance(tool, dict)]
    effects = sorted({str(tool.get("effect") or "read") for tool in tools})
    uses_network = server.get("id") in {"weather", "web-search"} or kind in {"mcp_http", "http"}
    writes_artifacts = any(effect in {"write", "side_effect"} for effect in effects)
    return {
        "method": "builtin_mcp",
        "requires_approval": False,
        "will_create": [] if installed else ["mcp_server"],
        "will_enable": ["mcp_server"] if not installed or not installed.get("enabled") else [],
        "tools": [str(tool.get("name") or "") for tool in tools if tool.get("name")],
        "tool_effects": effects,
        "permissions": {
            "reads_uploaded_files": False,
            "writes_artifacts": writes_artifacts,
            "runs_local_process": kind == "mcp_stdio",
            "uses_network": uses_network,
        },
        "impact": "启用后工具服务会出现在工具接入页；智能体仍需拥有对应 MCP 权限才会在任务中调用。",
        "post_install": "在“工具接入”查看工具 Schema、连接状态和调用测试；再到智能体配置里绑定服务 ID。",
    }


@app.get("/api/marketplace")
def marketplace() -> dict[str, Any]:
    skills: list[dict[str, Any]] = []
    for item in BUILTIN_SKILL_CATALOG:
        installed = skill_registry.get_skill(item["id"])
        skills.append(
            {
                "id": item["id"],
                "name": item["name"],
                "description": item["description"],
                "keywords": list(item.get("keywords") or []),
                "source_label": item.get("source_label", "智枢内置目录"),
                "installed": bool(installed),
                "enabled": bool(installed.get("enabled")) if installed else False,
                "category": installed.get("category") if installed else item["id"],
                "install_plan": _skill_marketplace_plan(item, installed),
            }
        )
    mcps: list[dict[str, Any]] = []
    for server in BUILTIN_SERVERS:
        installed = mcp_gateway.get_server(server["id"])
        mcps.append(
            {
                "id": server["id"],
                "name": server["name"],
                "description": server["description"],
                "tools": [tool.get("name") for tool in server.get("tools", [])],
                "source_label": "平台内置 MCP",
                "installed": bool(installed),
                "enabled": bool(installed.get("enabled")) if installed else False,
                "kind": installed.get("kind") if installed else server.get("kind", "builtin"),
                "install_plan": _mcp_marketplace_plan(server, installed),
            }
        )
    return {"skills": skills, "mcp_servers": mcps}


@app.post("/api/marketplace/skills/{skill_id}/install")
def install_marketplace_skill(skill_id: str) -> dict[str, Any]:
    builtin = get_builtin_skill(skill_id)
    if not builtin:
        raise HTTPException(status_code=404, detail="市场中没有这个 Skill")
    current = skill_registry.get_skill(skill_id)
    if current:
        if not current.get("enabled"):
            return skill_registry.update_skill(skill_id, {"enabled": True}) or current
        return current
    payload = {
        "id": builtin["id"],
        "name": builtin["name"],
        "description": builtin["description"],
        "category": "recommended",
        "version": "1.0.0",
        "content": builtin["content"],
        "enabled": True,
        "required_mcps": [],
    }
    return skill_registry.create_skill(payload)


@app.post("/api/marketplace/mcp/{server_id}/enable")
def enable_marketplace_mcp(server_id: str) -> dict[str, Any]:
    builtin = next((item for item in BUILTIN_SERVERS if item["id"] == server_id), None)
    if not builtin:
        raise HTTPException(status_code=404, detail="市场中没有这个 MCP")
    if not mcp_gateway.get_server(server_id):
        mcp_gateway.seed_builtin_servers()
    updated = mcp_gateway.update_server(server_id, {"enabled": True})
    if not updated:
        raise HTTPException(status_code=404, detail="MCP server not found")
    return updated


def _api_scope(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> ExecutionScope:
    identity = auth_service.current_identity.get()
    if identity:
        organization_id, user_id = "local-org", identity["user_id"]
        if identity.get("request_resource") != "workspaces":
            _require_workspace_access(workspace_id, write=identity.get("request_method") not in {"GET", "HEAD", "OPTIONS"})
    try:
        return ExecutionScope(
            organization_id=organization_id,
            workspace_id=workspace_id,
            user_id=user_id,
            agent_id=agent_id,
            conversation_id=conversation_id,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _actor_id(fallback: str) -> str:
    identity = auth_service.current_identity.get()
    return identity['user_id'] if identity else fallback


def _require_shared_creation_scope(scope: str) -> None:
    identity = auth_service.current_identity.get()
    if identity and identity['role'] != 'admin' and scope in {'organization', 'public'}:
        raise HTTPException(status_code=403, detail='发布组织共享资源需要管理员权限')


def _require_workspace_access(workspace_id: str, *, write: bool = False, manage: bool = False) -> None:
    identity = auth_service.current_identity.get()
    if not identity:
        return
    access = auth_service.workspace_access(workspace_id, identity)
    if not access or (write and access == "viewer") or (manage and access != "owner"):
        raise HTTPException(status_code=403, detail="无权访问或修改此工作区")


@app.get("/api/workspaces/{workspace_id}/members")
def list_workspace_members(workspace_id: str) -> list[dict[str, Any]]:
    _require_workspace_access(workspace_id, manage=True)
    if not auth_service.current_identity.get():
        raise HTTPException(status_code=403, detail="请先启用认证")
    return db.query_all("SELECT m.user_id,u.username,m.role FROM workspace_members m JOIN users u ON u.id=m.user_id WHERE m.workspace_id=? ORDER BY u.username", (workspace_id,))


@app.put("/api/workspaces/{workspace_id}/members/{user_id}")
def set_workspace_member(workspace_id: str, user_id: str, payload: WorkspaceMemberUpdate) -> dict[str, Any]:
    _require_workspace_access(workspace_id, manage=True)
    if not auth_service.current_identity.get():
        raise HTTPException(status_code=403, detail="请先启用认证")
    if not db.query_one("SELECT id FROM users WHERE id=? AND enabled=1", (user_id,)):
        raise HTTPException(status_code=404, detail="用户不存在或已停用")
    db.execute("INSERT INTO workspace_members(workspace_id,user_id,role) VALUES(?,?,?) ON CONFLICT(workspace_id,user_id) DO UPDATE SET role=excluded.role", (workspace_id,user_id,payload.role))
    return {"workspace_id": workspace_id, "user_id": user_id, "role": payload.role}


@app.delete("/api/workspaces/{workspace_id}/members/{user_id}")
def remove_workspace_member(workspace_id: str, user_id: str) -> dict[str, bool]:
    _require_workspace_access(workspace_id, manage=True)
    if not auth_service.current_identity.get():
        raise HTTPException(status_code=403, detail="请先启用认证")
    db.execute("DELETE FROM workspace_members WHERE workspace_id=? AND user_id=?", (workspace_id,user_id))
    return {"removed": True}


@app.put('/api/workspaces/{workspace_id}/members/by-username/{username}')
def set_workspace_member_by_username(workspace_id: str, username: str, payload: WorkspaceMemberUpdate) -> dict[str, Any]:
    _require_workspace_access(workspace_id, manage=True)
    if not auth_service.current_identity.get():
        raise HTTPException(status_code=403, detail='请先启用认证')
    user = db.query_one('SELECT id FROM users WHERE username=? AND enabled=1', (username,))
    if not user:
        raise HTTPException(status_code=404, detail='用户不存在或已停用')
    return set_workspace_member(workspace_id, user['id'], payload)


def _expert_http_error(exc: Exception) -> HTTPException:
    if isinstance(exc, ExpertNotFoundError):
        return HTTPException(status_code=404, detail=str(exc))
    if isinstance(exc, ExpertPermissionError):
        return HTTPException(status_code=403, detail=str(exc))
    if isinstance(exc, ExpertConflictError):
        return HTTPException(status_code=409, detail=str(exc))
    return HTTPException(status_code=400, detail=str(exc))


@app.get("/api/workspaces")
def list_workspaces(
    organization_id: str = "local-org",
    user_id: str = "local-user",
    include_disabled: bool = False,
) -> list[dict[str, Any]]:
    try:
        items = workspace_service.list_workspaces(
            _api_scope(organization_id, "default", user_id),
            include_disabled=include_disabled,
        )
        identity = auth_service.current_identity.get()
        return [item for item in items if not identity or auth_service.workspace_access(item["id"], identity)]
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/workspaces", status_code=201)
def create_workspace(payload: WorkspaceCreate) -> dict[str, Any]:
    try:
        return workspace_service.create_workspace(
            _api_scope(payload.organization_id, "default", payload.user_id),
            workspace_id=payload.id,
            name=payload.name,
            description=payload.description,
            default_agent_id=payload.default_agent_id,
            default_model_id=payload.default_model_id,
            settings=payload.settings,
            enabled=payload.enabled,
        )
    except sqlite3.IntegrityError as exc:
        raise HTTPException(status_code=409, detail="项目 ID 已存在") from exc
    except (WorkspaceError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/workspaces/{workspace_id}")
def get_workspace(
    workspace_id: str,
    organization_id: str = "local-org",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_workspace_access(workspace_id)
    try:
        item = workspace_service.get_workspace(
            workspace_id, _api_scope(organization_id, "default", user_id)
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not item:
        raise HTTPException(status_code=404, detail="项目不存在或当前实例不可见")
    return item


@app.put("/api/workspaces/{workspace_id}")
def update_workspace(
    workspace_id: str,
    payload: WorkspaceUpdate,
    organization_id: str = "local-org",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_workspace_access(workspace_id, manage=True)
    try:
        return workspace_service.update_workspace(
            workspace_id,
            _api_scope(organization_id, "default", user_id),
            name=payload.name,
            description=payload.description,
            default_agent_id=payload.default_agent_id,
            default_model_id=payload.default_model_id,
            settings=payload.settings,
            enabled=payload.enabled,
        )
    except WorkspaceNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (WorkspaceError, TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc



@app.get("/api/workspaces/{workspace_id}/codex-sessions")
def list_workspace_codex_sessions(
    workspace_id: str,
    organization_id: str = "local-org",
    user_id: str = "local-user",
) -> list[dict[str, Any]]:
    _require_workspace_access(workspace_id)
    paths = default_path_manager.get_paths(organization_id, user_id, workspace_id)
    sessions_dir = paths.codex_state_dir / "sessions"
    if not sessions_dir.exists():
        return []
    results = []
    for f in sessions_dir.rglob("*.jsonl"):
        try:
            stat = f.stat()
            session_id = ""
            first_prompt = ""
            last_reply = ""
            turn_count = 0
            model = ""
            with open(f, "r", encoding="utf-8", errors="replace") as fh:
                for line in fh:
                    if not line.strip():
                        continue
                    try:
                        entry = json.loads(line)
                        t = entry.get("type")
                        p = entry.get("payload", {})
                        if t == "session_meta":
                            session_id = str(p.get("id") or "") or session_id
                            model = str(p.get("model") or "") or model
                        elif t == "response_item":
                            role = p.get("role") or p.get("type")
                            content = p.get("content") or p.get("message")
                            if role == "user":
                                turn_count += 1
                                if not first_prompt:
                                    if isinstance(content, list) and content:
                                        first_prompt = str(content[0].get("text") or "")
                                    else:
                                        first_prompt = str(content)
                            elif role == "assistant":
                                if isinstance(content, list) and content:
                                    last_reply = str(content[-1].get("text") or "")
                                else:
                                    last_reply = str(content)
                    except Exception:
                        continue
            if not session_id:
                session_id = f.stem.replace("rollout-", "")
            clean_first = first_prompt
            if "<environment_context>" in clean_first:
                clean_first = clean_first.split("</environment_context>")[-1].strip()
            results.append({
                "session_id": session_id,
                "file_name": f.name,
                "created_at": datetime.fromtimestamp(stat.st_ctime, timezone.utc).isoformat(),
                "updated_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
                "turn_count": turn_count,
                "first_prompt": clean_first[:150],
                "last_reply": last_reply[:150],
                "model": model,
                "size_bytes": stat.st_size,
            })
        except Exception:
            continue
    results.sort(key=lambda x: x["updated_at"], reverse=True)
    return results


@app.get("/api/workspaces/{workspace_id}/codex-sessions/{session_id}")
def get_workspace_codex_session(
    workspace_id: str,
    session_id: str,
    organization_id: str = "local-org",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_workspace_access(workspace_id)
    paths = default_path_manager.get_paths(organization_id, user_id, workspace_id)
    sessions_dir = paths.codex_state_dir / "sessions"
    target_file = None
    if sessions_dir.exists():
        for f in sessions_dir.rglob("*.jsonl"):
            if session_id in f.name:
                target_file = f
                break
    if not target_file or not target_file.exists():
        raise HTTPException(status_code=404, detail="未找到该 Codex 会话记录")

    messages = []
    meta = {}
    with open(target_file, "r", encoding="utf-8", errors="replace") as fh:
        for line in fh:
            try:
                entry = json.loads(line)
                t = entry.get("type")
                p = entry.get("payload", {})
                if t == "session_meta":
                    meta = p
                elif t == "response_item":
                    role = p.get("role") or p.get("type")
                    content = p.get("content") or p.get("message")
                    text = ""
                    if isinstance(content, list):
                        parts = [str(item.get("text") or "") for item in content if isinstance(item, dict) and item.get("text")]
                        text = "\n".join(parts)
                    elif isinstance(content, str):
                        text = content
                    clean_text = text.strip()
                    if role in ("user", "assistant") and clean_text:
                        if clean_text.startswith("<environment_context>"):
                            clean_text = clean_text.split("</environment_context>")[-1].strip()
                        if clean_text:
                            messages.append({
                                "role": role,
                                "content": clean_text,
                                "timestamp": entry.get("timestamp"),
                            })
            except Exception:
                continue
    return {
        "session_id": session_id,
        "workspace_id": workspace_id,
        "meta": meta,
        "messages": messages,
    }


@app.delete("/api/workspaces/{workspace_id}")
def delete_workspace(
    workspace_id: str,
    organization_id: str = "local-org",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_workspace_access(workspace_id, manage=True)
    try:
        return workspace_service.delete_workspace(
            workspace_id, _api_scope(organization_id, "default", user_id)
        )
    except WorkspaceNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except WorkspaceError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _public_expert_selection(
    team: dict[str, Any],
    scope: ExecutionScope,
    *,
    automatic: bool,
    recommendation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Project an expert-team routing decision onto safe, user-facing fields."""

    supervisor_id = str(team.get("supervisor_agent_id") or "")
    supervisor = expert_team_service.get_agent(supervisor_id, scope) or {}
    members: list[dict[str, str]] = []
    for member in team.get("members") or []:
        agent_id = str(member.get("agent_id") or "")
        agent = expert_team_service.get_agent(agent_id, scope) or {}
        members.append(
            {
                "agent_id": agent_id,
                "agent_name": str(agent.get("name") or agent_id),
                "role": str(member.get("role") or "专家"),
            }
        )
    recommendation = recommendation or {}
    return {
        "selection_mode": "automatic" if automatic else "manual",
        "team_id": str(team.get("id") or ""),
        "team_name": str(team.get("name") or team.get("id") or "专家团"),
        "reason": str(
            recommendation.get("reason")
            or ("已根据当前目标自动匹配专家团" if automatic else "使用用户指定的专家团")
        ),
        "matched_terms": [str(item) for item in recommendation.get("matched_terms") or []][:8],
        "supervisor": {
            "agent_id": supervisor_id,
            "agent_name": str(supervisor.get("name") or supervisor_id),
        },
        "members": members,
    }


@app.get("/api/expert-templates")
def list_expert_templates(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    include_disabled: bool = False,
) -> list[dict[str, Any]]:
    return expert_team_service.list_templates(
        _api_scope(organization_id, workspace_id, user_id),
        include_disabled=include_disabled,
    )


@app.post("/api/expert-templates", status_code=201)
def create_expert_template(payload: ExpertTemplateCreate) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility)
    scope = _api_scope(payload.organization_id, payload.workspace_id, payload.owner_user_id)
    values = payload.model_dump()
    values.update(organization_id=scope.organization_id, owner_user_id=scope.user_id)
    try:
        return expert_team_service.create_template(values)
    except (ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc


@app.get("/api/expert-templates/{template_id}")
def get_expert_template(
    template_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    item = expert_team_service.get_template(
        template_id, _api_scope(organization_id, workspace_id, user_id)
    )
    if not item:
        raise HTTPException(status_code=404, detail="专家模板不存在或当前作用域不可见")
    return item


@app.put("/api/expert-templates/{template_id}")
def update_expert_template(
    template_id: str,
    payload: ExpertTemplateUpdate,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility or '')
    try:
        return expert_team_service.update_template(
            template_id,
            _api_scope(organization_id, workspace_id, user_id),
            payload.model_dump(exclude_unset=True),
        )
    except (ExpertNotFoundError, ExpertPermissionError, ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc


@app.delete("/api/expert-templates/{template_id}")
def delete_expert_template(
    template_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        expert_team_service.delete_template(
            template_id, _api_scope(organization_id, workspace_id, user_id)
        )
        return {"ok": True, "id": template_id}
    except (ExpertNotFoundError, ExpertPermissionError, ExpertConflictError) as exc:
        raise _expert_http_error(exc) from exc


@app.post("/api/expert-templates/{template_id}/install", status_code=201)
def install_expert_template(
    template_id: str, payload: ExpertInstallRequest
) -> dict[str, Any]:
    scope = _api_scope(payload.organization_id, payload.workspace_id, payload.user_id)
    try:
        return expert_team_service.install_template(template_id, scope, payload.model_dump())
    except (ExpertNotFoundError, ExpertPermissionError, ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc


@app.get("/api/expert-installations")
def list_expert_installations(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    include_disabled: bool = False,
) -> list[dict[str, Any]]:
    return expert_team_service.list_installations(
        _api_scope(organization_id, workspace_id, user_id),
        include_disabled=include_disabled,
    )


@app.delete("/api/expert-installations/{installation_id}")
def uninstall_expert(
    installation_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        return expert_team_service.disable_installation(
            installation_id, _api_scope(organization_id, workspace_id, user_id)
        )
    except (ExpertNotFoundError, ExpertConflictError) as exc:
        raise _expert_http_error(exc) from exc


@app.get("/api/expert-teams")
def list_expert_teams(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    include_disabled: bool = False,
) -> list[dict[str, Any]]:
    return expert_team_service.list_teams(
        _api_scope(organization_id, workspace_id, user_id),
        include_disabled=include_disabled,
    )


@app.post("/api/expert-teams", status_code=201)
def create_expert_team(payload: ExpertTeamCreate) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility)
    scope = _api_scope(payload.organization_id, payload.workspace_id, payload.owner_user_id)
    values = payload.model_dump()
    values.update(organization_id=scope.organization_id, owner_user_id=scope.user_id)
    try:
        return expert_team_service.create_team(values)
    except (ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc


@app.get("/api/expert-teams/{team_id}")
def get_expert_team(
    team_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    item = expert_team_service.get_team(
        team_id, _api_scope(organization_id, workspace_id, user_id)
    )
    if not item:
        raise HTTPException(status_code=404, detail="专家团不存在或当前作用域不可见")
    return item


@app.put("/api/expert-teams/{team_id}")
def update_expert_team(
    team_id: str,
    payload: ExpertTeamUpdate,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility or '')
    try:
        return expert_team_service.update_team(
            team_id,
            _api_scope(organization_id, workspace_id, user_id),
            payload.model_dump(exclude_unset=True),
        )
    except (ExpertNotFoundError, ExpertPermissionError, ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc


@app.delete("/api/expert-teams/{team_id}")
def delete_expert_team(
    team_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        expert_team_service.delete_team(
            team_id, _api_scope(organization_id, workspace_id, user_id)
        )
        return {"ok": True, "id": team_id}
    except (ExpertNotFoundError, ExpertPermissionError, ExpertConflictError) as exc:
        raise _expert_http_error(exc) from exc


@app.post("/api/expert-teams/{team_id}/runs", status_code=202)
async def run_expert_team(team_id: str, payload: ExpertTeamRunCreate) -> dict[str, Any]:
    if payload.model_id and payload.model_id != "deterministic":
        configured = db.query_one(
            "SELECT id FROM model_configs WHERE id = ? AND enabled = 1", (payload.model_id,)
        )
        if not configured:
            raise HTTPException(status_code=400, detail="所选模型不存在或未启用")
    scope = _api_scope(payload.organization_id, payload.workspace_id, payload.user_id)
    try:
        task, parent_run, team_run = expert_team_service.create_task_and_run(
            team_id, scope, message=payload.message, model_id=payload.model_id,
            conversation_id=payload.conversation_id,
        )
    except (ExpertNotFoundError, ExpertConflictError, ExpertValidationError) as exc:
        raise _expert_http_error(exc) from exc
    _schedule_team_run(team_run["id"])
    return {"accepted": True, "task": task, "run": parent_run, "team_run": team_run}


@app.get("/api/expert-teams/{team_id}/runs")
def list_expert_team_runs(
    team_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> list[dict[str, Any]]:
    try:
        return expert_team_service.list_team_runs(
            team_id, _api_scope(organization_id, workspace_id, user_id)
        )
    except ExpertNotFoundError as exc:
        raise _expert_http_error(exc) from exc


@app.get("/api/expert-team-runs/{team_run_id}")
def get_expert_team_run(
    team_run_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    item = expert_team_service.get_team_run(
        team_run_id, _api_scope(organization_id, workspace_id, user_id)
    )
    if not item:
        raise HTTPException(status_code=404, detail="专家团运行不存在或当前作用域不可见")
    return item


@app.post(
    "/api/expert-team-runs/{team_run_id}/members/{member_run_id}/retry",
    status_code=202,
)
async def retry_expert_team_member(
    team_run_id: str,
    member_run_id: str,
    payload: ExpertMemberRetryRequest,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _ = payload
    scope = _api_scope(organization_id, workspace_id, user_id)
    try:
        expert_team_service.validate_member_retry(team_run_id, member_run_id, scope)
    except (ExpertNotFoundError, ExpertConflictError) as exc:
        raise _expert_http_error(exc) from exc
    _schedule_member_retry(team_run_id, member_run_id, scope)
    return {
        "accepted": True,
        "team_run_id": team_run_id,
        "member_run_id": member_run_id,
        "scope": {
            "organization_id": scope.organization_id,
            "workspace_id": scope.workspace_id,
            "user_id": scope.user_id,
        },
    }


@app.get("/api/memories")
def list_memories(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
    scope_type: str | None = None,
    include_disabled: bool = True,
    include_expired: bool = True,
) -> list[dict[str, Any]]:
    try:
        return context_service.list_memories(
            _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
            scope_type=scope_type,
            include_disabled=include_disabled,
            include_expired=include_expired,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/memories", status_code=201)
def create_memory(payload: MemoryCreate) -> dict[str, Any]:
    _require_shared_creation_scope(payload.scope_type)
    scope = _api_scope(
        payload.organization_id,
        payload.workspace_id,
        payload.user_id,
        payload.agent_id,
        payload.conversation_id,
    )
    try:
        return context_service.create_memory(
            scope,
            scope_type=payload.scope_type,
            kind=payload.kind,
            title=payload.title,
            content=payload.content,
            tags=payload.tags,
            source_type=payload.source_type,
            source_ref=payload.source_ref,
            trust_level=payload.trust_level,
            enabled=payload.enabled,
            expires_at=payload.expires_at,
            created_by=_actor_id(payload.user_id),
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/memories/{memory_id}")
def get_memory(
    memory_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    memory = context_service.get_memory(
        memory_id,
        _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
    )
    if not memory:
        raise HTTPException(status_code=404, detail="记忆不存在或不属于当前执行作用域")
    return memory


@app.put("/api/memories/{memory_id}")
def update_memory(
    memory_id: str,
    payload: MemoryUpdate,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    changes = payload.model_dump(exclude_unset=True)
    reason = str(changes.pop("reason", "updated"))
    try:
        return context_service.update_memory(
            memory_id,
            _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
            actor_id=_actor_id(user_id),
            reason=reason,
            **changes,
        )
    except MemoryNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/memories/{memory_id}/enable")
def enable_memory(
    memory_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    try:
        return context_service.enable_memory(
            memory_id,
            _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
            actor_id=_actor_id(user_id),
        )
    except MemoryNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.post("/api/memories/{memory_id}/disable")
def disable_memory(
    memory_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    try:
        return context_service.disable_memory(
            memory_id,
            _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
            actor_id=_actor_id(user_id),
        )
    except MemoryNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.delete("/api/memories/{memory_id}")
def delete_memory(
    memory_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    try:
        return context_service.delete_memory(
            memory_id,
            _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
            actor_id=_actor_id(user_id),
            reason="api_delete",
        )
    except MemoryNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@app.get("/api/memories/{memory_id}/revisions")
def list_memory_revisions(
    memory_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> list[dict[str, Any]]:
    return context_service.list_revisions(
        memory_id,
        _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id),
    )


@app.get("/api/context/effective")
def get_effective_context(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    agent_id: str = "",
    conversation_id: str = "",
) -> dict[str, Any]:
    return context_service.get_effective_context(
        _api_scope(organization_id, workspace_id, user_id, agent_id, conversation_id)
    )


@app.get("/api/knowledge-bases")
def list_knowledge_bases(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    include_disabled: bool = True,
) -> list[dict[str, Any]]:
    try:
        return knowledge_service.list_bases(
            _api_scope(organization_id, workspace_id, user_id),
            include_disabled=include_disabled,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/knowledge-bases", status_code=201)
def create_knowledge_base(payload: KnowledgeBaseCreate) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility)
    try:
        return knowledge_service.create_base(
            _api_scope(payload.organization_id, payload.workspace_id, payload.user_id),
            base_id=payload.id,
            name=payload.name,
            description=payload.description,
            visibility=payload.visibility,
            enabled=payload.enabled,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/knowledge-bases/{base_id}")
def get_knowledge_base(
    base_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        item = knowledge_service.get_base(
            base_id, _api_scope(organization_id, workspace_id, user_id)
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not item:
        raise HTTPException(status_code=404, detail="知识库不存在或当前作用域不可见")
    return item


@app.put("/api/knowledge-bases/{base_id}")
def update_knowledge_base(
    base_id: str,
    payload: KnowledgeBaseUpdate,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    _require_shared_creation_scope(payload.visibility)
    try:
        return knowledge_service.update_base(
            base_id,
            _api_scope(organization_id, workspace_id, user_id),
            name=payload.name,
            description=payload.description,
            visibility=payload.visibility,
            enabled=payload.enabled,
        )
    except KnowledgeBaseNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/knowledge-bases/{base_id}")
def delete_knowledge_base(
    base_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        return knowledge_service.delete_base(
            base_id, _api_scope(organization_id, workspace_id, user_id)
        )
    except KnowledgeBaseNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/knowledge-bases/{base_id}/documents")
def list_knowledge_documents(
    base_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> list[dict[str, Any]]:
    try:
        return knowledge_service.list_documents(
            base_id, _api_scope(organization_id, workspace_id, user_id)
        )
    except KnowledgeBaseNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/knowledge-bases/{base_id}/documents/upload", status_code=201)
def index_knowledge_upload(
    base_id: str,
    payload: KnowledgeDocumentUpload,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    identity = auth_service.current_identity.get()
    if identity and identity["role"] != "admin":
        owner = db.query_one("SELECT user_id FROM upload_owners WHERE upload_id=?", (payload.upload_id,))
        if not owner or owner["user_id"] != identity["user_id"]:
            raise HTTPException(status_code=403, detail="无权使用此附件")
    upload = db.query_one("SELECT * FROM uploads WHERE id = ?", (payload.upload_id,))
    if not upload:
        raise HTTPException(status_code=404, detail="上传文件不存在")
    try:
        return knowledge_service.index_upload(
            base_id,
            _api_scope(organization_id, workspace_id, user_id),
            upload=upload,
        )
    except KnowledgeBaseNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except KnowledgeBaseError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/knowledge/search")
def search_knowledge(
    q: str,
    base_id: str = "",
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    limit: int = 5,
) -> dict[str, Any]:
    try:
        return knowledge_service.search(
            _api_scope(organization_id, workspace_id, user_id),
            query=q,
            base_id=base_id,
            limit=limit,
        )
    except KnowledgeBaseNotFoundError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/conversation-summaries")
def list_conversation_summaries(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    limit: int = 100,
) -> list[dict[str, Any]]:
    try:
        return runtime.conversation_summary_service.list(
            _api_scope(organization_id, workspace_id, user_id), limit=limit
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/conversation-summaries/{conversation_id}")
def get_conversation_summary(
    conversation_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        item = runtime.conversation_summary_service.get(
            _api_scope(
                organization_id,
                workspace_id,
                user_id,
                conversation_id=conversation_id,
            ),
            conversation_id,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not item:
        raise HTTPException(status_code=404, detail="对话摘要不存在或不属于当前执行作用域")
    return item


@app.put("/api/conversation-summaries/{conversation_id}")
def update_conversation_summary(
    conversation_id: str,
    payload: ConversationSummaryUpdate,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    scope = _api_scope(
        organization_id,
        workspace_id,
        user_id,
        conversation_id=conversation_id,
    )
    try:
        return runtime.conversation_summary_service.upsert(
            scope,
            conversation_id=conversation_id,
            summary=payload.summary,
            preserved_constraints=payload.preserved_constraints,
            through_task_id=payload.through_task_id,
            model_id=payload.model_id,
        )
    except ConversationSummaryConflictError as exc:
        raise HTTPException(status_code=409, detail="该对话 ID 已属于其他用户或工作区") from exc
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.delete("/api/conversation-summaries/{conversation_id}")
def delete_conversation_summary(
    conversation_id: str,
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> dict[str, Any]:
    try:
        deleted = runtime.conversation_summary_service.delete(
            _api_scope(
                organization_id,
                workspace_id,
                user_id,
                conversation_id=conversation_id,
            ),
            conversation_id,
        )
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not deleted:
        raise HTTPException(status_code=404, detail="对话摘要不存在或不属于当前执行作用域")
    return {"conversation_id": conversation_id, "deleted": True}


def _policy_rule_to_api(value: dict[str, Any]) -> dict[str, Any]:
    return PolicyRule.from_dict(value).to_dict()


@app.get("/api/policies")
def list_policy_rules() -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for row in db.query_all("SELECT rule_json FROM policy_rules ORDER BY priority DESC, id"):
        value = db.json_loads(row.get("rule_json"), {})
        if isinstance(value, dict):
            result.append(_policy_rule_to_api(value))
    return result


@app.get("/api/policies/{rule_id}")
def get_policy_rule(rule_id: str) -> dict[str, Any]:
    row = db.query_one("SELECT rule_json FROM policy_rules WHERE id = ?", (rule_id,))
    if not row:
        raise HTTPException(status_code=404, detail="策略规则不存在")
    return _policy_rule_to_api(db.json_loads(row.get("rule_json"), {}))


@app.post("/api/policies", status_code=201)
def create_policy_rule(payload: PolicyRuleCreate) -> dict[str, Any]:
    if db.query_one("SELECT id FROM policy_rules WHERE id = ?", (payload.id,)):
        raise HTTPException(status_code=409, detail="策略规则 ID 已存在")
    raw = payload.model_dump(exclude_none=True)
    raw["name"] = str(raw.get("name") or payload.id)
    try:
        rule = PolicyRule.from_dict(raw)
    except PolicyConfigurationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    now = db.utc_now()
    event_label = ",".join(rule.events)
    db.execute(
        """
        INSERT INTO policy_rules(id, name, event, scope, scope_id, priority, enabled, rule_json, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            rule.id,
            rule.name,
            event_label,
            rule.scope,
            rule.scope_id or "",
            rule.priority,
            1 if rule.enabled else 0,
            db.json_dumps(raw),
            now,
            now,
        ),
    )
    _reload_policy_rules()
    return rule.to_dict()


@app.put("/api/policies/{rule_id}")
def update_policy_rule(rule_id: str, payload: PolicyRuleUpdate) -> dict[str, Any]:
    row = db.query_one("SELECT rule_json FROM policy_rules WHERE id = ?", (rule_id,))
    if not row:
        raise HTTPException(status_code=404, detail="策略规则不存在")
    current = db.json_loads(row.get("rule_json"), {})
    incoming = payload.model_dump(exclude_unset=True)
    if "event" in incoming:
        current.pop("events", None)
    if "events" in incoming:
        current.pop("event", None)
    merged = {**current, **incoming, "id": rule_id}
    try:
        rule = PolicyRule.from_dict(merged)
    except PolicyConfigurationError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.execute(
        """
        UPDATE policy_rules
        SET name = ?, event = ?, scope = ?, scope_id = ?, priority = ?, enabled = ?, rule_json = ?, updated_at = ?
        WHERE id = ?
        """,
        (
            rule.name,
            ",".join(rule.events),
            rule.scope,
            rule.scope_id or "",
            rule.priority,
            1 if rule.enabled else 0,
            db.json_dumps(merged),
            db.utc_now(),
            rule_id,
        ),
    )
    _reload_policy_rules()
    return rule.to_dict()


@app.delete("/api/policies/{rule_id}")
def delete_policy_rule(rule_id: str) -> dict[str, Any]:
    if not db.query_one("SELECT id FROM policy_rules WHERE id = ?", (rule_id,)):
        raise HTTPException(status_code=404, detail="策略规则不存在")
    db.execute("DELETE FROM policy_rules WHERE id = ?", (rule_id,))
    _reload_policy_rules()
    return {"ok": True, "id": rule_id}


@app.get("/api/skills")
def list_skills() -> list[dict[str, Any]]:
    return skill_registry.list_skills()


@app.get("/api/skills/{skill_id}")
def get_skill(skill_id: str) -> dict[str, Any]:
    skill = skill_registry.get_skill(skill_id)
    if not skill:
        raise HTTPException(status_code=404, detail="Skill not found")
    return skill


@app.get("/api/skills/{skill_id}/export")
def export_skill_package(skill_id: str) -> StreamingResponse:
    skill = skill_registry.get_skill(skill_id)
    if not skill:
        raise HTTPException(status_code=404, detail="Skill not found")
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for item in skill_registry.list_files(skill_id):
            row = db.query_one(
                "SELECT content FROM skill_files WHERE skill_id = ? AND path = ?",
                (skill_id, item["path"]),
            )
            raw = (row or {}).get("content") or b""
            if isinstance(raw, str):
                raw = raw.encode("utf-8")
            archive.writestr(item["path"], raw)
    buffer.seek(0)
    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", skill_id) or "skill"
    return StreamingResponse(
        buffer,
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="{safe_name}.zip"'},
    )


@app.get("/api/skills/{skill_id}/files")
def list_skill_files(skill_id: str) -> list[dict[str, Any]]:
    if not skill_registry.get_skill(skill_id):
        raise HTTPException(status_code=404, detail="Skill not found")
    return skill_registry.list_files(skill_id)


@app.post("/api/skills/{skill_id}/files/upload")
async def upload_skill_file(skill_id: str, path: str, file: UploadFile = File(...)) -> dict[str, Any]:
    raw = await file.read(1024 * 1024 + 1)
    if len(raw) > 1024 * 1024:
        raise HTTPException(status_code=413, detail="单个 Skill 文件不能超过 1MB")
    try:
        return skill_registry.put_file(skill_id, path, raw)
    except ValueError as exc:
        status = 404 if str(exc) == "Skill not found" else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@app.get("/api/skills/{skill_id}/files/{file_path:path}")
def get_skill_file(skill_id: str, file_path: str) -> dict[str, Any]:
    try:
        item = skill_registry.get_file(skill_id, file_path)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if not item:
        raise HTTPException(status_code=404, detail="Skill file not found")
    return item


@app.put("/api/skills/{skill_id}/files/{file_path:path}")
def put_skill_file(skill_id: str, file_path: str, payload: SkillFileUpdate) -> dict[str, Any]:
    try:
        return skill_registry.put_file(skill_id, file_path, payload.content.encode("utf-8"))
    except ValueError as exc:
        status = 404 if str(exc) == "Skill not found" else 400
        raise HTTPException(status_code=status, detail=str(exc)) from exc


@app.delete("/api/skills/{skill_id}/files/{file_path:path}")
def delete_skill_file(skill_id: str, file_path: str) -> dict[str, Any]:
    try:
        skill_registry.delete_file(skill_id, file_path)
        return {"ok": True, "path": file_path}
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/skills")
def create_skill(payload: SkillCreate) -> dict[str, Any]:
    if skill_registry.get_skill(payload.id):
        raise HTTPException(status_code=409, detail="Skill already exists")
    missing_mcps = [item for item in payload.required_mcps if not mcp_gateway.server_exists(item)]
    if missing_mcps:
        raise HTTPException(status_code=400, detail="Skill 引用了不存在的 MCP：" + "、".join(missing_mcps))
    return skill_registry.create_skill(payload.model_dump())


@app.post("/api/skills/install/upload")
async def install_skill_upload(file: UploadFile = File(...)) -> dict[str, Any]:
    raw = await file.read(2 * 1024 * 1024 + 1)
    if len(raw) > 2 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Skill 安装包不能超过 2MB")
    try:
        filename = file.filename or "SKILL.md"
        if filename != "SKILL.md" and not filename.lower().endswith((".md", ".zip")):
            raise ValueError("仅支持 SKILL.md 或 ZIP 安装包")
        return _install_skill_bytes(raw, filename)
    except (ValueError, UnicodeDecodeError, zipfile.BadZipFile) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/skills/install/url")
async def install_skill_url(payload: RemoteInstall) -> dict[str, Any]:
    try:
        return await _install_skill_remote_url(payload.url)
    except (ValueError, UnicodeDecodeError, zipfile.BadZipFile, httpx.HTTPError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/skills/install/path")
def install_skill_path(payload: SkillPathInstall) -> dict[str, Any]:
    try:
        return skill_registry.install_from_path(payload.path, payload.enabled)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/api/skills/{skill_id}")
def update_skill(skill_id: str, payload: SkillUpdate) -> dict[str, Any]:
    if payload.required_mcps is not None:
        missing_mcps = [item for item in payload.required_mcps if not mcp_gateway.server_exists(item)]
        if missing_mcps:
            raise HTTPException(status_code=400, detail="Skill 引用了不存在的 MCP：" + "、".join(missing_mcps))
    updated = skill_registry.update_skill(skill_id, payload.model_dump(exclude_unset=True))
    if not updated:
        raise HTTPException(status_code=404, detail="Skill not found")
    return updated


@app.delete("/api/skills/{skill_id}")
def delete_skill(skill_id: str) -> dict[str, Any]:
    skill = skill_registry.get_skill(skill_id)
    if not skill:
        raise HTTPException(status_code=404, detail="Skill not found")
    if skill.get("category") == "builtin":
        raise HTTPException(status_code=400, detail="平台内置 Skill 不能卸载，可在编辑器中停用")
    for agent in list_agents():
        bound = [item for item in agent.get("skills", []) if item != skill_id]
        if bound != agent.get("skills", []):
            db.execute("UPDATE agents SET skills_json = ?, updated_at = ? WHERE id = ?", (db.json_dumps(bound), db.utc_now(), agent["id"]))
    skill_registry.delete_package(skill_id)
    db.execute("DELETE FROM skills WHERE id = ?", (skill_id,))
    return {"ok": True, "id": skill_id}


@app.get("/api/mcp")
def list_mcp_servers() -> list[dict[str, Any]]:
    return mcp_gateway.list_servers()


@app.post("/api/mcp")
def create_mcp_server(payload: McpServerCreate) -> dict[str, Any]:
    if mcp_gateway.server_exists(payload.id):
        raise HTTPException(status_code=409, detail="MCP server already exists")
    return mcp_gateway.create_server(payload.model_dump())


@app.post("/api/mcp/import")
async def import_mcp_config(file: UploadFile = File(...)) -> list[dict[str, Any]]:
    raw = await file.read(1024 * 1024 + 1)
    if len(raw) > 1024 * 1024:
        raise HTTPException(status_code=413, detail="MCP 配置文件不能超过 1MB")
    try:
        payload = json.loads(raw.decode("utf-8"))
        return mcp_gateway.import_config(payload)
    except (ValueError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/api/mcp/install/url")
async def install_mcp_url(payload: RemoteInstall) -> list[dict[str, Any]]:
    try:
        return await _install_mcp_remote_url(payload.url)
    except (ValueError, UnicodeDecodeError, json.JSONDecodeError, httpx.HTTPError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.put("/api/mcp/{server_id}")
def update_mcp_server(server_id: str, payload: McpServerUpdate) -> dict[str, Any]:
    updated = mcp_gateway.update_server(server_id, payload.model_dump(exclude_unset=True))
    if not updated:
        raise HTTPException(status_code=404, detail="MCP server not found")
    return updated


@app.delete("/api/mcp/{server_id}")
def delete_mcp_server(server_id: str) -> dict[str, Any]:
    server = mcp_gateway.get_server(server_id)
    if not server:
        raise HTTPException(status_code=404, detail="MCP server not found")
    if server.get("kind") == "builtin":
        raise HTTPException(status_code=400, detail="平台内置 MCP 不能卸载，可在编辑器中停用")
    for agent in list_agents():
        bound = [item for item in agent.get("mcp_servers", []) if item != server_id]
        if bound != agent.get("mcp_servers", []):
            db.execute("UPDATE agents SET mcp_servers_json = ?, updated_at = ? WHERE id = ?", (db.json_dumps(bound), db.utc_now(), agent["id"]))
    db.execute("DELETE FROM mcp_servers WHERE id = ?", (server_id,))
    return {"ok": True, "id": server_id}


@app.get("/api/mcp/{server_id}")
def get_mcp_server(server_id: str) -> dict[str, Any]:
    server = mcp_gateway.get_server(server_id)
    if not server:
        raise HTTPException(status_code=404, detail="MCP server not found")
    return server


@app.get("/api/mcp/{server_id}/tools")
def list_mcp_tools(server_id: str) -> list[dict[str, Any]]:
    if not mcp_gateway.get_server(server_id):
        raise HTTPException(status_code=404, detail="MCP server not found")
    return mcp_gateway.list_tools(server_id)


@app.post("/api/mcp/{server_id}/discover")
async def discover_mcp_tools(server_id: str) -> list[dict[str, Any]]:
    try:
        return await mcp_gateway.discover_tools(server_id)
    except (ToolError, ImportError, OSError) as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _model_capabilities(row: dict[str, Any]) -> dict[str, Any]:
    provider = str(row.get("provider") or "")
    config = db.json_loads(row.get("config_json"), {})
    if provider == "deterministic":
        return {
            "protocol": "offline",
            "streaming": True,
            "tool_calling": False,
            "online_required": False,
            "credential_required": False,
            "context_window": None,
            "notes": ["离线流程检查模型，不适合复杂内容生成。"],
        }
    return {
        "protocol": "openai_chat_completions",
        "streaming": True,
        "tool_calling": True,
        "online_required": True,
        "credential_required": True,
        "context_window": config.get("context_window") or config.get("max_context_tokens"),
        "notes": [
            "需要 APP_ALLOW_OUTBOUND_NETWORK 和 APP_MODEL_HOST_ALLOWLIST 允许目标主机。",
            "逐段实时输出取决于上游是否正确支持 SSE stream=true。",
            "工具调用取决于上游是否兼容 Chat Completions tools 字段。",
        ],
    }


def _model_readiness(row: dict[str, Any]) -> dict[str, Any]:
    provider = str(row.get("provider") or "")
    enabled = bool(row.get("enabled"))
    if provider == "deterministic":
        return {"state": "ready", "label": "内置可用", "detail": "无需密钥和网络，适合检查流程。"}
    has_direct_key = bool(row.get("api_key_ciphertext"))
    api_key_env = str(row.get("api_key_env") or "").strip()
    if not enabled:
        return {"state": "off", "label": "已停用", "detail": "启用后才会进入模型选择器。"}
    if not has_direct_key and not api_key_env:
        return {"state": "needs_config", "label": "缺少密钥", "detail": "请选择环境变量或直接 API Key。"}
    if not str(row.get("base_url") or "").strip():
        return {"state": "needs_config", "label": "缺少 Base URL", "detail": "请填写 OpenAI-compatible API 根地址。"}
    if api_key_env and not os.getenv(api_key_env):
        return {"state": "needs_config", "label": "环境变量未检测", "detail": f"服务进程当前未检测到 {api_key_env}。"}
    if not outbound_network_enabled():
        return {"state": "needs_config", "label": "联网未开启", "detail": "需要开启 APP_ALLOW_OUTBOUND_NETWORK 并配置模型主机白名单。"}
    return {"state": "ready", "label": "配置完整", "detail": "建议点击测试连接确认供应商接口可用。"}


def model_to_api(row: dict[str, Any]) -> dict[str, Any]:
    legacy_key = str(row.get("api_key_env") or "").strip()
    has_direct_key = bool(row.get("api_key_ciphertext")) or bool(legacy_key and not _is_env_name(legacy_key))
    safe = {k: v for k, v in row.items() if k not in {"api_key_ciphertext", "config_json"}}
    safe["api_key_env"] = legacy_key if _is_env_name(legacy_key) else ""
    public = {**safe, "enabled": bool(row.get("enabled")), "config": db.json_loads(row.get("config_json"), {}), "has_api_key": has_direct_key, "api_key_mode": "direct" if has_direct_key else "env"}
    last_test = {
        "status": str(row.get("last_test_status") or ""),
        "message": str(row.get("last_test_message") or ""),
        "tested_at": str(row.get("last_test_at") or ""),
    }
    if not last_test["status"]:
        last_test = {"status": "untested", "message": "尚未测试连接", "tested_at": ""}
    return {**public, "last_test": last_test, "capabilities": _model_capabilities(row), "readiness": _model_readiness(row)}


@app.get("/api/models")
def list_models() -> list[dict[str, Any]]:
    deterministic = {"id": "deterministic", "name": "离线确定性模型", "provider": "deterministic", "model": "deterministic-offline", "base_url": "", "api_key_env": "", "enabled": True, "config": {}, "has_api_key": False, "api_key_mode": "env", "last_test": {"status": "pass", "message": "内置模型无需连接测试", "tested_at": ""}, "capabilities": _model_capabilities({"provider": "deterministic"}), "readiness": _model_readiness({"provider": "deterministic", "enabled": True})}
    return [deterministic] + [model_to_api(r) for r in db.query_all("SELECT * FROM model_configs ORDER BY name")]


def _ensure_model_ready(model_id: str | None, *, label: str = "所选模型") -> None:
    if not model_id or model_id == "deterministic":
        return
    row = db.query_one("SELECT * FROM model_configs WHERE id = ?", (model_id,))
    if not row or not row.get("enabled"):
        raise HTTPException(status_code=400, detail=f"{label}不存在或未启用")
    readiness = _model_readiness(row)
    if readiness.get("state") != "ready":
        detail = str(readiness.get("detail") or "").strip()
        suffix = f"：{detail}" if detail else ""
        raise HTTPException(
            status_code=400,
            detail=f"{label}暂不可用（{readiness.get('label', '需要配置')}）{suffix}",
        )


@app.post("/api/models")
def create_model(payload: ModelConfigCreate) -> dict[str, Any]:
    if payload.id == "deterministic" or db.query_one("SELECT id FROM model_configs WHERE id = ?", (payload.id,)):
        raise HTTPException(status_code=409, detail="Model config already exists")
    now = db.utc_now()
    encrypted = ""
    api_key_env = ""
    base_url = payload.base_url

    if payload.copy_credentials_from:
        src = db.query_one("SELECT * FROM model_configs WHERE id = ?", (payload.copy_credentials_from,))
        if src:
            encrypted = src.get("api_key_ciphertext") or ""
            api_key_env = src.get("api_key_env") or ""
            if not base_url:
                base_url = src.get("base_url") or ""

    if not encrypted and not api_key_env:
        if payload.api_key_mode not in {"env", "direct"}:
            raise HTTPException(status_code=400, detail="api_key_mode 必须是 env 或 direct")
        if payload.api_key_mode == "direct" and not payload.api_key:
            raise HTTPException(status_code=400, detail="直接密钥模式必须填写 API Key")
        if payload.api_key_mode == "env" and not _is_env_name(payload.api_key_env):
            raise HTTPException(status_code=400, detail="环境变量模式必须填写合法变量名，例如 OPENAI_API_KEY")
        encrypted = secret_store.encrypt(payload.api_key) if payload.api_key_mode == "direct" and payload.api_key else ""
        api_key_env = payload.api_key_env if payload.api_key_mode == "env" else ""

    db.execute(
        "INSERT INTO model_configs(id, name, provider, model, base_url, api_key_env, api_key_ciphertext, enabled, config_json, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (payload.id, payload.name, payload.provider, payload.model, base_url, api_key_env, encrypted, 1 if payload.enabled else 0, db.json_dumps(payload.config), now, now),
    )
    return model_to_api(db.query_one("SELECT * FROM model_configs WHERE id = ?", (payload.id,)) or {})


@app.put("/api/models/{model_id}")
def update_model(model_id: str, payload: ModelConfigUpdate) -> dict[str, Any]:
    current = db.query_one("SELECT * FROM model_configs WHERE id = ?", (model_id,))
    if not current:
        raise HTTPException(status_code=404, detail="Model config not found")
    current_api = model_to_api(current)
    merged = {**current_api, **payload.model_dump(exclude_unset=True)}
    mode = payload.api_key_mode or current_api.get("api_key_mode", "env")
    if mode not in {"env", "direct"}:
        raise HTTPException(status_code=400, detail="api_key_mode 必须是 env 或 direct")
    encrypted = current.get("api_key_ciphertext") or ""
    if mode == "env":
        encrypted = ""
        if not _is_env_name(str(merged.get("api_key_env") or "")):
            raise HTTPException(status_code=400, detail="环境变量模式必须填写合法变量名，例如 OPENAI_API_KEY")
    elif payload.api_key:
        encrypted = secret_store.encrypt(payload.api_key)
    elif not encrypted:
        raise HTTPException(status_code=400, detail="直接密钥模式必须填写 API Key")
    db.execute(
        "UPDATE model_configs SET name = ?, provider = ?, model = ?, base_url = ?, api_key_env = ?, api_key_ciphertext = ?, enabled = ?, config_json = ?, updated_at = ? WHERE id = ?",
        (merged["name"], merged["provider"], merged["model"], merged.get("base_url", ""), merged.get("api_key_env", "") if mode == "env" else "", encrypted, 1 if merged.get("enabled") else 0, db.json_dumps(merged.get("config", {})), db.utc_now(), model_id),
    )
    return model_to_api(db.query_one("SELECT * FROM model_configs WHERE id = ?", (model_id,)) or {})


@app.delete("/api/models/{model_id}")
def delete_model(model_id: str) -> dict[str, Any]:
    if model_id == "deterministic":
        raise HTTPException(status_code=400, detail="内置离线模型不能删除")
    if not db.query_one("SELECT id FROM model_configs WHERE id = ?", (model_id,)):
        raise HTTPException(status_code=404, detail="Model config not found")
    db.execute("UPDATE agents SET model = 'deterministic', updated_at = ? WHERE model = ?", (db.utc_now(), model_id))
    db.execute("DELETE FROM model_configs WHERE id = ?", (model_id,))
    return {"ok": True, "id": model_id}


@app.post("/api/models/{model_id}/test")
async def test_model(model_id: str) -> dict[str, Any]:
    if model_id == "deterministic":
        return {
            "ok": True,
            "model_id": model_id,
            "response": "内置模型无需连接测试",
            "model": next(item for item in list_models() if item["id"] == "deterministic"),
        }
    if not db.query_one("SELECT id FROM model_configs WHERE id = ?", (model_id,)):
        raise HTTPException(status_code=404, detail="Model config not found")
    try:
        response = await model_gateway.summarize(
            "这是连接测试。请只回复 OK。",
            {"system_prompt": "你正在执行模型连接测试。"},
            model_config_id=model_id,
        )
        message = str(response or "OK")[:1000]
        db.execute(
            "UPDATE model_configs SET last_test_status = ?, last_test_message = ?, last_test_at = ?, updated_at = ? WHERE id = ?",
            ("pass", message, db.utc_now(), db.utc_now(), model_id),
        )
        return {
            "ok": True,
            "model_id": model_id,
            "response": message,
            "model": model_to_api(db.query_one("SELECT * FROM model_configs WHERE id = ?", (model_id,)) or {}),
        }
    except Exception as exc:
        message = str(exc)[:1000]
        db.execute(
            "UPDATE model_configs SET last_test_status = ?, last_test_message = ?, last_test_at = ?, updated_at = ? WHERE id = ?",
            ("fail", message, db.utc_now(), db.utc_now(), model_id),
        )
        raise HTTPException(status_code=400, detail=message) from exc




@app.post("/api/models/discover")
async def discover_remote_models(payload: ModelDiscoverRequest) -> dict[str, Any]:
    base_url = payload.base_url.rstrip("/")
    key = ""
    if payload.api_key_mode == "direct" and payload.api_key:
        key = payload.api_key
    elif payload.api_key_mode == "env" and payload.api_key_env:
        key = os.getenv(payload.api_key_env, "")
    elif payload.api_key:
        key = payload.api_key

    if not key and payload.engine_id:
        engine_env = resolve_runtime_env(payload.engine_id)
        key = engine_env.get("OPENAI_API_KEY") or engine_env.get("ANTHROPIC_API_KEY") or ""
    if not key and payload.model_id:
        model_row = db.query_one("SELECT * FROM model_configs WHERE id = ?", (payload.model_id,))
        if model_row:
            if model_row.get("api_key_ciphertext"):
                try:
                    key = secret_store.decrypt(model_row["api_key_ciphertext"])
                except Exception:
                    pass
            elif model_row.get("api_key_env"):
                key = os.getenv(model_row["api_key_env"], "")

    headers = {}
    if key:
        headers["Authorization"] = f"Bearer {key}"

    url = f"{base_url}/models"
    try:
        async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
            resp = await client.get(url, headers=headers)
            if resp.status_code == 404 and not base_url.endswith("/v1"):
                url = f"{base_url}/v1/models"
                resp = await client.get(url, headers=headers)
            if resp.status_code == 401:
                raise HTTPException(status_code=400, detail="获取失败：密钥无效或未授权 (401 Unauthorized)")
            if resp.status_code != 200:
                raise HTTPException(status_code=400, detail=f"上游服务返回状态码 {resp.status_code}：{resp.text[:200]}")
            data = resp.json()
            items = data.get("data", [])
            if isinstance(items, list):
                models = sorted({str(m.get("id") or m) for m in items if isinstance(m, dict) and m.get("id") or isinstance(m, str)})
            else:
                models = []
            if not models:
                raise HTTPException(status_code=400, detail="上游服务返回成功但未包含任何可用模型")
            return {"ok": True, "models": models, "count": len(models)}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"连接模型接口失败：{exc}")

@app.get("/api/execution-engines")
def list_execution_engines() -> list[dict[str, Any]]:
    return list_engines()


@app.get("/api/execution-engines/{engine_id}")
def get_execution_engine(engine_id: str) -> dict[str, Any]:
    engine = get_engine(engine_id)
    if not engine:
        raise HTTPException(status_code=404, detail="执行引擎不存在")
    return engine


@app.put("/api/execution-engines/{engine_id}")
def update_execution_engine(engine_id: str, payload: ExecutionEngineUpdate) -> dict[str, Any]:
    try:
        return update_engine(engine_id, payload.model_dump(exclude_unset=True))
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ExecutionEngineError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc



@app.post("/api/execution-engines/{engine_id}/test")
async def test_execution_engine(engine_id: str) -> dict[str, Any]:
    try:
        return await test_engine_connection(engine_id)
    except LookupError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc
    except ExecutionEngineError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

@app.post("/api/mcp/{server_id}/tools/{tool_name}/invoke")
async def invoke_mcp_tool(server_id: str, tool_name: str, payload: ToolInvokeRequest) -> dict[str, Any]:
    definition = mcp_gateway.get_tool_definition(server_id, tool_name)
    if not definition:
        raise HTTPException(status_code=404, detail="MCP 工具不存在，请先同步工具清单")
    # Validate required fields before the page-level read-only guard.  A user
    # who clicked "调用工具" with an incomplete form needs one concise,
    # actionable parameter hint; the permission message is only useful after
    # the invocation is otherwise well-formed.  Support both the platform's
    # snake_case schema and the MCP protocol's camelCase spelling because
    # imported configurations commonly use ``inputSchema``.
    schema = definition.get("input_schema") or definition.get("inputSchema") or {}
    if isinstance(schema, dict):
        required = schema.get("required")
        arguments = payload.arguments if isinstance(payload.arguments, dict) else {}
        if isinstance(required, list):
            missing = [
                str(item)
                for item in required
                if str(item) not in arguments
                or arguments.get(str(item)) in (None, "")
            ]
            if missing:
                raise HTTPException(
                    status_code=400,
                    detail="工具调用缺少必填参数：" + "、".join(missing),
                )
    server_kind = str(definition.get("server_kind") or "")
    trusted_read_only = (
        server_kind == "builtin" and definition.get("effect") == "read"
    ) or (
        server_kind != "builtin"
        and definition.get("annotations", {}).get("readOnlyHint") is True
    )
    if not trusted_read_only:
        raise HTTPException(
            status_code=403,
            detail=(
                "页面测试调用只允许明确标注的只读工具；写入、破坏性或未标注工具"
                "请通过正式对话任务执行，以应用智能体权限、Policy 和人工审批。"
            ),
        )
    policy_context = {
        "organization_id": "local-org",
        "workspace_id": "default",
        "user_id": "local-user",
        "task_id": payload.task_id or "",
        "tool": {
            "server": server_id,
            "server_id": server_id,
            "name": tool_name,
            "tool_name": tool_name,
            "arguments": payload.arguments,
            "direct_test": True,
        },
    }
    try:
        evaluation = await policy_engine.evaluate("tool.before", policy_context)
        if evaluation.denied:
            raise HTTPException(status_code=403, detail=evaluation.summary)
        if evaluation.requires_approval:
            raise HTTPException(
                status_code=403,
                detail=(
                    f"{evaluation.summary}。页面测试调用不会创建审批；"
                    "请通过正式对话任务执行并完成审批。"
                ),
            )
        modified = evaluation.apply(policy_context)
        arguments = modified.get("tool", {}).get("arguments", payload.arguments)
        if not isinstance(arguments, dict):
            raise HTTPException(status_code=400, detail="Policy 修改后的工具参数必须是 JSON 对象")
        return await mcp_gateway.invoke_tool(server_id, tool_name, arguments, task_id=payload.task_id)
    except ToolError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/agents")
def list_agents(workspace_id: str = 'default') -> list[dict[str, Any]]:
    _require_workspace_access(workspace_id)
    rows = db.query_all("SELECT * FROM agents ORDER BY name")
    identity = auth_service.current_identity.get()
    if identity:
        rows = [row for row in rows if auth_service.agent_access(row['id'], identity, workspace_id)]
    return [
        {
            **r,
            "skills": db.json_loads(r.get("skills_json"), []),
            "mcp_servers": db.json_loads(r.get("mcp_servers_json"), []),
            "permissions": db.json_loads(r.get("permissions_json"), {}),
        }
        for r in rows
    ]


def _validate_agent_bindings(model_id: str, skill_ids: list[str], mcp_ids: list[str]) -> None:
    if model_id != "deterministic" and not db.query_one(
        "SELECT id FROM model_configs WHERE id = ? AND enabled = 1", (model_id,)
    ):
        raise HTTPException(status_code=400, detail="智能体选择的模型不存在或未启用")
    missing_skills = [item for item in dict.fromkeys(skill_ids) if not skill_registry.get_skill(item)]
    if missing_skills:
        raise HTTPException(status_code=400, detail="智能体引用了不存在的 Skill：" + "、".join(missing_skills))
    missing_mcps = [item for item in dict.fromkeys(mcp_ids) if not mcp_gateway.server_exists(item)]
    if missing_mcps:
        raise HTTPException(status_code=400, detail="智能体引用了不存在的 MCP：" + "、".join(missing_mcps))


@app.post("/api/agents")
def create_agent(payload: AgentCreate) -> dict[str, Any]:
    if db.query_one("SELECT id FROM agents WHERE id = ?", (payload.id,)):
        raise HTTPException(status_code=409, detail="Agent already exists")
    _validate_agent_bindings(payload.model, payload.skills, payload.mcp_servers)
    now = db.utc_now()
    db.execute(
        """
        INSERT INTO agents(id, name, description, model, system_prompt, skills_json, mcp_servers_json, permissions_json, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            payload.id, payload.name, payload.description, payload.model, payload.system_prompt,
            db.json_dumps(payload.skills), db.json_dumps(payload.mcp_servers), db.json_dumps(payload.permissions), now, now,
        ),
    )
    return next(a for a in list_agents() if a["id"] == payload.id)


@app.put("/api/agents/{agent_id}")
def update_agent(agent_id: str, payload: AgentUpdate) -> dict[str, Any]:
    current = db.query_one("SELECT * FROM agents WHERE id = ?", (agent_id,))
    if not current:
        raise HTTPException(status_code=404, detail="Agent not found")
    current_api = {
        **current,
        "skills": db.json_loads(current.get("skills_json"), []),
        "mcp_servers": db.json_loads(current.get("mcp_servers_json"), []),
        "permissions": db.json_loads(current.get("permissions_json"), {}),
    }
    incoming = payload.model_dump(exclude_unset=True)
    merged = {**current_api, **incoming}
    _validate_agent_bindings(merged["model"], merged.get("skills", []), merged.get("mcp_servers", []))
    db.execute(
        """
        UPDATE agents SET name = ?, description = ?, model = ?, system_prompt = ?, skills_json = ?, mcp_servers_json = ?, permissions_json = ?, updated_at = ?
        WHERE id = ?
        """,
        (
            merged["name"], merged["description"], merged["model"], merged["system_prompt"],
            db.json_dumps(merged.get("skills", [])), db.json_dumps(merged.get("mcp_servers", [])), db.json_dumps(merged.get("permissions", {})),
            db.utc_now(), agent_id,
        ),
    )
    return next(a for a in list_agents() if a["id"] == agent_id)


@app.get("/api/tasks")
def list_tasks(workspace_id: str = "", organization_id: str = "", user_id: str = "") -> list[dict[str, Any]]:
    identity = auth_service.current_identity.get()
    if identity and identity["role"] != "admin":
        organization_id, user_id = "local-org", identity["user_id"]
    clauses: list[str] = []
    params: list[Any] = []
    if identity and identity["role"] != "admin":
        clauses.append("workspace IN (SELECT w.id FROM workspaces w LEFT JOIN workspace_members m ON m.workspace_id=w.id AND m.user_id=? WHERE w.organization_id='local-org' AND (w.owner_user_id=? OR (w.enabled=1 AND m.user_id IS NOT NULL)))")
        params.extend([identity["user_id"], identity["user_id"]])
    for column, value in (
        ("workspace", workspace_id),
        ("organization_id", organization_id),
        ("user_id", user_id),
    ):
        if value:
            clauses.append(f"{column} = ?")
            params.append(value)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    rows = db.query_all(
        f"SELECT * FROM tasks{where} ORDER BY created_at DESC LIMIT 100",  # noqa: S608 - fixed columns
        tuple(params),
    )
    return [_public_task(r) for r in rows]


def _public_attachment(value: dict[str, Any]) -> dict[str, Any]:
    """Expose attachment metadata without leaking the server storage path."""
    allowed = {"id", "name", "content_type", "size", "created_at"}
    result = {
        key: value.get(key)
        for key in allowed
        if value.get(key) not in (None, "")
    }
    result["context_status"] = _attachment_context_status(value)
    return result


def _attachment_context_status(value: dict[str, Any]) -> dict[str, Any]:
    """Describe whether an upload can be injected into task context without exposing content/path."""
    name = str(value.get("name") or "")
    suffix = Path(name).suffix.lower()
    content_type = str(value.get("content_type") or "application/octet-stream")
    size = int(value.get("size") or 0)
    text_suffixes = {
        ".txt", ".md", ".csv", ".json", ".yaml", ".yml",
        ".py", ".js", ".ts", ".html", ".css",
    }
    document_suffixes = {".docx", ".xlsx", ".pptx", ".pdf"}
    supported = suffix in text_suffixes or suffix in document_suffixes or content_type.startswith("text/")
    limit = int(getattr(runtime, "ATTACHMENT_MAX_FILE_BYTES", 2 * 1024 * 1024))
    if not supported:
        return {
            "state": "unsupported",
            "label": "不支持正文解析",
            "detail": "文件会保留为附件记录，但不会自动进入模型上下文。",
            "supported": False,
            "extractable": False,
            "max_bytes": limit,
        }
    if size > limit:
        return {
            "state": "too_large",
            "label": "超过解析上限",
            "detail": f"文件超过 {max(1, limit // 1024 // 1024)}MB 正文解析上限，不会自动进入模型上下文。",
            "supported": True,
            "extractable": False,
            "max_bytes": limit,
        }
    return {
        "state": "ready",
        "label": "可进入上下文",
        "detail": "发送任务后，平台会尝试提取正文并作为附件上下文使用。",
        "supported": True,
        "extractable": True,
        "max_bytes": limit,
    }


def _public_artifact(value: dict[str, Any]) -> dict[str, Any]:
    artifact_id = str(value.get("id") or "")
    declared_status = str(value.get("delivery_status") or "")
    authoritative_status = ""
    if artifact_id:
        row = db.query_one(
            "SELECT delivery_status FROM artifacts WHERE id = ?", (artifact_id,)
        )
        if row is not None:
            authoritative_status = str(row.get("delivery_status") or "")
    # Event payloads and historic task projections are not an authority for
    # publication.  A usable link is projected only when the current Artifact
    # row says it was published by the finalization transaction.
    delivery_status = authoritative_status or (
        declared_status
        if declared_status in {"pending_verification", "rejected"}
        else "unavailable"
    )
    allowed = {
        "id", "task_id", "run_id", "workspace_id", "name", "kind",
        "mime_type", "size", "sha256", "version", "created_at",
        "delivery_status",
    }
    result = {key: value.get(key) for key in allowed if value.get(key) not in (None, "")}
    result["delivery_status"] = delivery_status
    if artifact_id and authoritative_status == "published":
        result["download_url"] = f"/api/artifacts/{artifact_id}/download"
        result["preview_url"] = f"/api/artifacts/{artifact_id}/preview"
    return result


def _sanitize_public_payload(value: Any) -> Any:
    if isinstance(value, list):
        return [_sanitize_public_payload(item) for item in value]
    if not isinstance(value, dict):
        return value
    looks_like_artifact = bool(
        str(value.get("id") or "").startswith("art_")
        or value.get("download_url")
        or (value.get("name") and value.get("kind") and ("path" in value or "relative_path" in value))
    )
    if looks_like_artifact:
        return _public_artifact(value)
    return {
        key: _sanitize_public_payload(item)
        for key, item in value.items()
        if key not in {"api_key", "api_key_ciphertext", "path", "relative_path"}
    }


def _public_goal_ref(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {}
    return {
        key: value.get(key)
        for key in ("id", "goal_id", "version", "spec_hash")
        if isinstance(value.get(key), (str, int)) and value.get(key) not in (None, "")
    }


def _public_goal_summary(value: Any) -> dict[str, Any]:
    """Re-project a stored GoalSpec summary instead of trusting its JSON shape."""

    if not isinstance(value, dict):
        return {}
    objective = value.get("objective") if isinstance(value.get("objective"), dict) else {}
    capabilities = (
        value.get("capabilities") if isinstance(value.get("capabilities"), dict) else {}
    )
    confirmation = (
        value.get("confirmation") if isinstance(value.get("confirmation"), dict) else {}
    )
    return {
        "schema_version": value.get("schema_version"),
        "goal_id": value.get("goal_id"),
        "task_id": value.get("task_id"),
        "version": value.get("version"),
        "status": value.get("status"),
        "objective": {
            "statement": str(objective.get("statement") or ""),
            "intent": str(objective.get("intent") or ""),
            "in_scope": [item for item in objective.get("in_scope", []) if isinstance(item, str)],
            "out_of_scope": [item for item in objective.get("out_of_scope", []) if isinstance(item, str)],
            "constraints": [item for item in objective.get("constraints", []) if isinstance(item, str)],
        },
        "missing_inputs": [
            {
                key: item.get(key)
                for key in ("key", "label", "ask")
                if item.get(key) not in (None, "")
            }
            for item in value.get("missing_inputs", [])
            if isinstance(item, dict)
        ],
        "deliverables": [
            {
                key: item.get(key)
                for key in (
                    "id", "kind", "format", "title", "filename", "required",
                    "download_required",
                )
                if item.get(key) not in (None, "")
            }
            for item in value.get("deliverables", [])
            if isinstance(item, dict)
        ],
        "capabilities": {
            "skills": [
                {
                    key: item.get(key)
                    for key in ("skill_id", "name", "version", "purpose")
                    if item.get(key) not in (None, "")
                }
                for item in capabilities.get("skills", [])
                if isinstance(item, dict)
            ],
            "tools": [
                {
                    key: item.get(key)
                    for key in ("server_id", "tool_name", "effect", "purpose")
                    if item.get(key) not in (None, "")
                }
                for item in capabilities.get("tools", [])
                if isinstance(item, dict)
            ],
            "network_access": capabilities.get("network_access"),
        },
        "context_ref_count": value.get("context_ref_count", 0),
        "acceptance": [
            {
                key: item.get(key)
                for key in ("id", "title", "kind", "severity")
                if item.get(key) not in (None, "")
            }
            for item in value.get("acceptance", [])
            if isinstance(item, dict)
        ],
        "confirmation": {
            **{
                key: confirmation.get(key)
                for key in ("status", "mode", "confidence")
                if isinstance(confirmation.get(key), (str, int, float, bool))
                and confirmation.get(key) not in (None, "")
            },
            "ambiguities": [
                item
                for item in confirmation.get("ambiguities", [])
                if isinstance(item, str)
            ],
        },
        "supersedes": _public_goal_ref(value.get("supersedes")),
        "spec_hash": value.get("spec_hash"),
    }


def _public_verification_report(value: Any) -> dict[str, Any]:
    """Keep only the intentionally public verifier contract."""

    if not isinstance(value, dict):
        return {}
    semantic = value.get("semantic") if isinstance(value.get("semantic"), dict) else {}
    return {
        key: item
        for key, item in {
            "schema_version": value.get("schema_version"),
            "mode": value.get("mode"),
            "verdict": value.get("verdict"),
            "passed": value.get("passed"),
            "coverage": value.get("coverage"),
            "semantic_attempted": value.get("semantic_attempted"),
            "semantic_verified": value.get("semantic_verified"),
            "public_reason": value.get("public_reason"),
            "rules": [
                {
                    field: rule.get(field)
                    for field in (
                        "id", "title", "status", "public_reason", "repair_instruction",
                    )
                    if isinstance(rule.get(field), (str, int, float, bool))
                    and rule.get(field) not in (None, "")
                }
                for rule in value.get("rules", [])
                if isinstance(rule, dict)
            ],
            "semantic": {
                **{
                    field: semantic.get(field)
                    for field in ("status", "public_reason")
                    if isinstance(semantic.get(field), (str, int, float, bool))
                    and semantic.get(field) not in (None, "")
                },
                "repair_instructions": [
                    item
                    for item in semantic.get("repair_instructions", [])
                    if isinstance(item, str)
                ],
            },
            "repair_instructions": [
                item
                for item in value.get("repair_instructions", [])
                if isinstance(item, str)
            ],
        }.items()
        if item is not None
    }


def _public_plan_node(value: Any, *, include_children: bool = True) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {}
    node = {
        key: value.get(key)
        for key in ("id", "title", "kind", "status")
        if value.get(key) not in (None, "")
    }
    if include_children:
        node["children"] = [
            _public_plan_node(item, include_children=False)
            for item in value.get("children", [])
            if isinstance(item, dict)
        ]
    return node


def _public_plan(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {}
    confirmation = (
        value.get("goal_confirmation")
        if isinstance(value.get("goal_confirmation"), dict)
        else {}
    )
    return {
        "goal": value.get("goal"),
        "goal_confirmation": {
            key: confirmation.get(key)
            for key in ("status", "label", "message")
            if confirmation.get(key) not in (None, "")
        },
        "intent": value.get("intent"),
        "steps": [str(item) for item in value.get("steps", []) if str(item).strip()],
        "nodes": [
            _public_plan_node(item)
            for item in value.get("nodes", [])
            if isinstance(item, dict)
        ],
        "allowed_servers": [
            str(item) for item in value.get("allowed_servers", []) if str(item).strip()
        ],
        "output_format": value.get("output_format"),
        "output_formats": [
            str(item) for item in value.get("output_formats", []) if str(item).strip()
        ],
        "requested_formats": [
            str(item) for item in value.get("requested_formats", []) if str(item).strip()
        ],
        "unavailable_formats": [
            str(item) for item in value.get("unavailable_formats", []) if str(item).strip()
        ],
        "requires_artifact": bool(value.get("requires_artifact")),
        "acceptance_criteria": [
            {
                key: item.get(key)
                for key in ("id", "title", "status")
                if item.get(key) not in (None, "")
            }
            for item in value.get("acceptance_criteria", [])
            if isinstance(item, dict)
        ],
    }


def _public_event_data(event_type: str, value: Any) -> dict[str, Any]:
    """Project each public event type; arbitrary nested payloads are never copied."""

    data = value if isinstance(value, dict) else {}
    scalar_fields: dict[str, tuple[str, ...]] = {
        "start": ("run_id", "attempt"),
        "goal_spec_progress": ("status",),
        "recovery": ("run_id", "checkpoint_id"),
        "recovery_scheduled": ("run_id", "checkpoint_id", "attempt"),
        "checkpoint": ("checkpoint_id", "run_id", "node_id"),
        "clarification": (),
        "progress": ("status", "percent"),
        "plan_progress": (
            "plan_id", "node_id", "status", "child_id", "child_title",
            "child_kind", "elapsed_seconds",
        ),
        "plan_check": ("tool", "passed", "plan_id", "reason", "source"),
        "tool_call": ("server_id", "tool_name"),
        "tool_result": ("server_id", "tool_name", "duration_ms"),
        "tool_error": ("server_id", "tool_name"),
        "tool_blocked": ("server_id", "tool_name", "reason", "source"),
        "tool_reused": ("server_id", "tool_name"),
        "answer_delta": ("draft", "delivery_state", "goal_spec_version"),
        "answer_reset": ("reason", "verification_id"),
        "verification_started": ("mode", "delivery_state"),
        "verification_result": ("verification_id", "delivery_state"),
        "candidate_verified": ("verification_id", "delivery_state"),
        "answer": ("verification_id", "delivery_state", "team_run_id"),
        "done": ("verification_id", "delivery_state"),
        "output_check": (
            "passed", "verification_id", "expected_format", "artifact_count",
        ),
        "approval_required": ("action",),
        # Raw exception class names are internal diagnostics.  The public
        # projector exposes only an intentionally assigned stable error code.
        "error": ("error_code", "team_run_id"),
        "steering": ("count",),
        "resume": ("run_id", "checkpoint_id"),
        "resume_scheduled": ("run_id", "checkpoint_id", "restore_count"),
        "retry_scheduled": ("run_id", "attempt"),
        "command_queued": ("command_id", "run_id", "intake_generation"),
        "team_queued": ("team_id", "team_run_id", "member_count"),
        "team_parallel_start": ("team_run_id",),
        "team_aggregating": ("team_run_id",),
        "team_completed": ("team_run_id",),
        "team_member_retry": ("team_run_id", "member_run_id"),
        "team_partial_failed": ("team_run_id",),
        "team_acceptance_failed": ("team_run_id",),
    }
    result = {
        key: data.get(key)
        for key in scalar_fields.get(event_type, ())
        if isinstance(data.get(key), (str, int, float, bool))
    }
    if event_type == "agent" and isinstance(data.get("agent"), dict):
        result["agent"] = {
            key: data["agent"].get(key)
            for key in ("id", "name", "description", "icon")
            if data["agent"].get(key) not in (None, "")
        }
    if event_type in {"skill", "install"}:
        if isinstance(data.get("skill"), dict):
            result["skill"] = {
                key: data["skill"].get(key)
                for key in ("id", "name", "description", "version", "enabled")
                if data["skill"].get(key) not in (None, "")
            }
        if isinstance(data.get("skills"), list):
            result["skills"] = [
                {
                    key: item.get(key)
                    for key in ("id", "name", "version", "score")
                    if item.get(key) not in (None, "")
                }
                for item in data["skills"]
                if isinstance(item, dict)
            ]
    if event_type == "install" and isinstance(data.get("mcp_servers"), list):
        result["mcp_servers"] = [
            {
                key: item.get(key)
                for key in ("id", "name", "transport", "enabled", "status")
                if item.get(key) not in (None, "")
            }
            for item in data["mcp_servers"]
            if isinstance(item, dict)
        ]
    if event_type == "knowledge":
        result["knowledge_base_ids"] = [
            str(item) for item in data.get("knowledge_base_ids", []) if str(item).strip()
        ][:20]
        result["matches"] = _public_knowledge_matches(data.get("matches"))
    if event_type == "goal_spec":
        result["goal_spec"] = _public_goal_summary(data.get("goal_spec"))
        result["goal_spec_ref"] = _public_goal_ref(data.get("goal_spec_ref"))
    if event_type in {"plan_check", "tool_blocked", "verification_started", "verification_result"}:
        goal_ref = _public_goal_ref(data.get("goal_spec_ref"))
        if goal_ref:
            result["goal_spec_ref"] = goal_ref
    if event_type == "plan":
        result["plan"] = _public_plan(data.get("plan"))
    if event_type == "tool_result" and isinstance(data.get("artifact"), dict):
        result["artifact"] = _public_artifact(data["artifact"])
    if event_type in {"answer", "tool_result"} and isinstance(data.get("artifacts"), list):
        result["artifacts"] = [
            _public_artifact(item) for item in data["artifacts"] if isinstance(item, dict)
        ]
    if event_type in {"verification_result", "output_check"}:
        report = _public_verification_report(data.get("report"))
        if report:
            result["report"] = report
        if isinstance(data.get("criteria"), list):
            result["criteria"] = [
                {
                    key: item.get(key)
                    for key in ("id", "title", "status", "detail", "public_reason")
                    if item.get(key) not in (None, "")
                }
                for item in data["criteria"]
                if isinstance(item, dict)
            ]
    if event_type == "approval_required" and isinstance(data.get("recommendations"), list):
        result["recommendations"] = [
            {
                key: item.get(key)
                for key in ("id", "name", "description", "source_label")
                if item.get(key) not in (None, "")
            }
            for item in data["recommendations"]
            if isinstance(item, dict)
        ]
    if event_type == "clarification" and isinstance(data.get("missing_information"), list):
        result["missing_information"] = [
            str(item) for item in data["missing_information"] if str(item).strip()
        ]
    if event_type == "expert_selection":
        result.update(
            {
                key: data.get(key)
                for key in ("selection_mode", "team_id", "team_name", "reason")
                if isinstance(data.get(key), (str, int, float, bool))
            }
        )
        result["matched_terms"] = [
            str(item) for item in data.get("matched_terms", []) if str(item).strip()
        ][:8]
        if isinstance(data.get("supervisor"), dict):
            result["supervisor"] = {
                key: data["supervisor"].get(key)
                for key in ("agent_id", "agent_name")
                if data["supervisor"].get(key) not in (None, "")
            }
        result["members"] = [
            {
                key: item.get(key)
                for key in ("agent_id", "agent_name", "role")
                if item.get(key) not in (None, "")
            }
            for item in data.get("members", [])
            if isinstance(item, dict)
        ]
    return result


def _public_error_code(value: Any) -> str:
    """Map internal exception classes to a small, user-facing error vocabulary."""

    data = value if isinstance(value, dict) else {}
    message = str(data.get("message") or "").strip().lower()
    # PPTX is an optional capability. Preserve that distinction at the public
    # boundary so a missing presentation component is actionable instead of
    # looking like an unrelated MCP/network failure. Only match known
    # capability messages; arbitrary provider text stays redacted.
    pptx_unavailable_markers = (
        "powerpoint 生成组件尚未安装",
        "powerpoint 生成需要 node.js",
        "powerpoint 生成脚本缺失",
        "powerpoint 生成功能尚未配置",
        "pptx 生成器尚未配置",
    )
    if any(marker in message for marker in pptx_unavailable_markers):
        return "artifact_pptx_unavailable"
    error_type = str(data.get("error_type") or "").lower().replace("_", "")
    if error_type in {
        "httpstatusexception", "connecterror", "readtimeout", "connecttimeout",
        "timeoutexception", "remoteprotocolerror", "networkerror",
    } or "model" in error_type:
        return "model_unavailable"
    # Network-policy failures are raised as a generic RuntimeError before an
    # HTTP client exception exists.  They still belong to the selected model,
    # not to the generic task-failure bucket; otherwise the conversation only
    # says “任务未完成” and leaves the user guessing why the run stopped.
    model_network_markers = (
        "模型 api",
        "模型api",
        "模型主机",
        "api主机",
        "联网未开启",
        "联网访问",
        "在线模型",
        "模型供应商",
        "模型配置不存在",
    )
    if any(marker in message for marker in model_network_markers):
        return "model_unavailable"
    if "tool" in error_type or "mcp" in error_type:
        return "tool_unavailable"
    return "task_execution_failed"


_PUBLIC_TASK_EVENT_TYPES = frozenset(
    {
        "agent", "answer", "answer_delta", "answer_reset", "approval",
        "approval_required", "cancelled", "candidate_verified", "checkpoint",
        "clarification", "command_queued", "conversation_summary", "done", "error",
        "expert_selection", "goal_spec", "goal_spec_progress", "install", "intent", "knowledge",
        "interrupted", "memory", "memory_deleted", "memory_saved", "model", "notice",
        "output_check", "permissions", "plan", "plan_check", "plan_progress", "progress",
        "recovery", "recovery_scheduled", "resume", "resume_scheduled", "retry_scheduled",
        "skill", "start", "steering", "team_acceptance_failed", "team_aggregating",
        "team_completed", "team_member_retry", "team_parallel_start", "team_partial_failed",
        "team_queued", "tool_blocked", "tool_call", "tool_error", "tool_result",
        "tool_reused", "verification_result", "verification_started",
    }
)


def _public_event(value: dict[str, Any]) -> dict[str, Any]:
    """Expose one known event through a type-specific, default-deny projector."""

    event_type = str(value.get("type") or "").lower()
    raw_data = db.json_loads(value.get("data_json"), {})
    error_code = _public_error_code(raw_data) if event_type == "error" else ""
    if event_type == "error":
        classification = dict(raw_data)
        # ``commit_failure`` stores the raw message in the event content and
        # keeps only a redacted data payload. Use it solely for selecting a
        # stable public code; never return it directly.
        if value.get("content") not in (None, ""):
            classification["message"] = value.get("content")
        error_code = _public_error_code(classification)
        error_copy = {
            "model_unavailable": (
                "模型暂时不可用或网络访问失败。请检查模型配置、联网开关后重试。"
            ),
            "tool_unavailable": (
                "工具服务调用失败。请检查 MCP 配置、授权或网络白名单后重试。"
            ),
            "artifact_pptx_unavailable": (
                "PowerPoint 生成功能尚未配置。请先配置平台的 PowerPoint 生成组件后重试。"
            ),
            "task_execution_failed": (
                "任务执行未完成。请检查模型、参数或工具配置后重试。"
            ),
        }
        title = {
            "model_unavailable": "模型暂时不可用",
            "tool_unavailable": "工具服务调用失败",
            "artifact_pptx_unavailable": "PowerPoint 生成功能不可用",
            "task_execution_failed": "任务未完成",
        }.get(error_code, "任务未完成")
        content = error_copy.get(error_code, error_copy["task_execution_failed"])
    elif event_type == "tool_error":
        title = "工具调用未完成"
        content = "工具调用未成功。请检查工具配置、授权或网络状态后重试。"
    elif event_type == "tool_blocked":
        title = value.get("title") or "工具调用已阻止"
        content = "该工具调用不符合当前目标、权限或执行计划，平台已停止执行。"
    elif event_type in {"plan_progress", "progress"} and str(
        raw_data.get("status") or ""
    ).lower() == "failed":
        title = value.get("title") or "执行步骤未完成"
        content = "该执行步骤未完成。请检查公开错误提示、参数或能力配置后重试。"
    else:
        title = value.get("title")
        content = value.get("content")
    result = {
        key: value.get(key)
        for key in ("id", "task_id", "ts", "type")
        if value.get(key) is not None
    }
    if title is not None:
        result["title"] = title
    if content is not None:
        result["content"] = content
    result["data"] = _public_event_data(event_type, raw_data)
    result['schema_version'] = 1
    if error_code:
        result["data"]["error_code"] = error_code
    return result


_TASK_STREAM_TERMINAL_STATUSES = frozenset({"completed", "failed", "cancelled"})


def _public_knowledge_matches(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    matches: list[dict[str, Any]] = []
    for item in value[:20]:
        if not isinstance(item, dict):
            continue
        public: dict[str, Any] = {}
        for key in ("chunk_id", "document_id", "knowledge_base_id", "document_name"):
            raw = item.get(key)
            if raw not in (None, ""):
                public[key] = str(raw)[:240]
        ordinal = item.get("ordinal")
        if isinstance(ordinal, int):
            public["ordinal"] = ordinal
        elif isinstance(ordinal, str) and ordinal.isdigit():
            public["ordinal"] = int(ordinal)
        if public:
            matches.append(public)
    return matches


def _task_knowledge_summary(task_id: str) -> dict[str, Any]:
    rows = db.query_all(
        "SELECT data_json FROM task_events WHERE task_id = ? AND type = 'knowledge' ORDER BY id",
        (task_id,),
    )
    by_chunk: dict[str, dict[str, Any]] = {}
    knowledge_base_ids: set[str] = set()
    documents: dict[str, dict[str, Any]] = {}
    for row in rows:
        data = db.json_loads(row.get("data_json"), {})
        for base_id in data.get("knowledge_base_ids", []) if isinstance(data, dict) else []:
            if str(base_id).strip():
                knowledge_base_ids.add(str(base_id))
        for match in _public_knowledge_matches(data.get("matches") if isinstance(data, dict) else None):
            chunk_key = str(match.get("chunk_id") or f"{match.get('document_id')}:{match.get('ordinal')}")
            by_chunk[chunk_key] = match
            doc_key = str(match.get("document_id") or match.get("document_name") or chunk_key)
            doc = documents.setdefault(
                doc_key,
                {
                    "document_id": match.get("document_id", ""),
                    "document_name": match.get("document_name", ""),
                    "match_count": 0,
                },
            )
            doc["match_count"] = int(doc.get("match_count") or 0) + 1
            if match.get("document_name"):
                doc["document_name"] = match["document_name"]
    matches = list(by_chunk.values())
    return {
        "match_count": len(matches),
        "knowledge_base_ids": sorted(knowledge_base_ids),
        "documents": sorted(
            documents.values(),
            key=lambda item: (str(item.get("document_name") or item.get("document_id") or ""), -int(item.get("match_count") or 0)),
        )[:20],
        "matches": matches[:20],
    }


def _is_public_task_event(value: dict[str, Any]) -> bool:
    """Unknown event types are private until an explicit projector is added."""

    return str(value.get("type") or "").lower() in _PUBLIC_TASK_EVENT_TYPES


def _task_event_cursor(
    request: Request,
    *,
    cursor: int | None,
    after_id: int | None,
) -> int:
    """Resolve an SSE resume cursor from the query string or Last-Event-ID."""
    explicit = cursor if cursor is not None else after_id
    raw: Any = explicit if explicit is not None else request.headers.get("last-event-id", "0")
    try:
        resolved = int(raw or 0)
    except (TypeError, ValueError) as exc:
        raise HTTPException(status_code=400, detail="事件游标必须是非负整数") from exc
    if resolved < 0:
        raise HTTPException(status_code=400, detail="事件游标必须是非负整数")
    return resolved


def _public_runtime_record(value: dict[str, Any]) -> dict[str, Any]:
    """Remove duplicate storage columns from TaskState API records."""
    internal_suffix_fields = {
        "result_json", "error_json", "metadata_json", "input_json",
        "output_json", "payload_json", "state_json",
        "last_restore_metadata_json", "command_type",
    }
    return {
        key: _sanitize_public_payload(item)
        for key, item in value.items()
        if key not in internal_suffix_fields
    }


def _public_task_run(value: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value.get(key)
        for key in (
            "id", "task_id", "attempt", "status", "current_node_id",
            "resumed_from_checkpoint_id", "started_at", "finished_at",
            "created_at", "updated_at",
        )
        if value.get(key) not in (None, "")
    }


def _public_task_node(value: dict[str, Any], *, attempt: int = 1) -> dict[str, Any]:
    has_error = bool(value.get("error"))
    metadata = value.get("metadata") if isinstance(value.get("metadata"), dict) else {}
    kind = str(value.get("kind") or "step").lower()
    logical_id = str(metadata.get("logical_id") or "")
    capability: dict[str, str] | None = None
    if kind in {"skill", "mcp", "tool", "model", "agent", "knowledge", "memory"}:
        capability_id = logical_id
        for prefix in ("skill:", "tool:", "model:", "agent:", "knowledge:", "memory:"):
            if capability_id.startswith(prefix):
                capability_id = capability_id[len(prefix):]
                break
        if re.fullmatch(r"[A-Za-z0-9_.:@/-]{1,160}", capability_id):
            capability = {
                "type": "mcp" if kind == "tool" else kind,
                "id": capability_id,
                "label": str(value.get("title") or capability_id)[:160],
            }
    status_message = ""
    if not has_error:
        status_message = str(metadata.get("last_message") or "")[:300]
    return {
        **{
            key: value.get(key)
            for key in (
                "id", "run_id", "task_id", "node_key", "parent_node_id",
                "title", "kind", "sequence", "status", "started_at",
                "finished_at", "created_at", "updated_at",
            )
            if value.get(key) not in (None, "")
        },
        "output_summary": str((value.get("output") or {}).get("summary") or "")[:500],
        "status_message": status_message,
        "error_summary": (
            "该执行步骤未完成。请检查公开错误提示、参数或能力配置后重试。"
            if has_error
            else ""
        ),
        "capability": capability,
        "attempt": attempt,
    }


def _public_checkpoint(value: dict[str, Any], *, attempt: int = 1) -> dict[str, Any]:
    return {
        **{
            key: value.get(key)
            for key in (
                "id", "task_id", "run_id", "node_id", "sequence", "reason",
                "restored_at", "restore_count", "created_at",
            )
            if value.get(key) not in (None, "")
        },
        "label": value.get("reason") or f"检查点 {value.get('sequence', '')}",
        "restorable": True,
        "attempt": attempt,
    }


def _public_task_command(value: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value.get(key)
        for key in (
            "id", "task_id", "run_id", "type", "status", "priority",
            "available_at", "created_at", "claimed_at", "completed_at", "updated_at",
        )
        if value.get(key) not in (None, "")
    }


def _two_level_node_tree(nodes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Build exactly roots plus one child layer; never expose nested node internals."""

    by_id = {str(item.get("id") or ""): item for item in nodes if item.get("id")}
    roots = [
        item
        for item in nodes
        if not item.get("parent_node_id")
        or str(item.get("parent_node_id")) not in by_id
    ]
    result: list[dict[str, Any]] = []
    for root in roots:
        root_id = str(root.get("id") or "")
        result.append(
            {
                **root,
                "children": [
                    {**child, "children": []}
                    for child in nodes
                    if str(child.get("parent_node_id") or "") == root_id
                ],
            }
        )
    return result


def _active_goal_projection(task_id: str, run_id: str = "") -> dict[str, Any] | None:
    goal = task_state.latest_goal_spec(run_id=run_id) if run_id else None
    if not goal:
        goal = task_state.latest_goal_spec(task_id=task_id)
    if not goal:
        return None
    summary = _public_goal_summary(goal.get("public_summary"))
    return {
        "id": goal.get("id"),
        "goal_id": summary.get("goal_id"),
        "version": goal.get("version"),
        "status": goal.get("status"),
        "spec_hash": goal.get("spec_hash"),
        "summary": summary,
    }


def _verification_projection(
    *,
    task_status: str,
    run_id: str,
    active_goal: dict[str, Any] | None,
) -> dict[str, Any]:
    if active_goal:
        reports = task_state.list_verifications(
            run_id=run_id or None,
            goal_spec_id=str(active_goal.get("id") or "") or None,
            limit=1,
        )
    else:
        reports = []
    if reports:
        report = reports[0]
        public_report = _public_verification_report(report.get("public_report"))
        if public_report.get("passed") is True:
            state = "passed"
        elif public_report.get("verdict") == "failed" or report.get("status") == "failed":
            state = "failed"
        else:
            state = "inconclusive"
        return {
            "state": state,
            **{
                key: report.get(key)
                for key in (
                    "id", "goal_spec_id", "attempt", "mode", "status",
                    "started_at", "finished_at", "created_at",
                )
                if report.get(key) not in (None, "")
            },
            "public_report": public_report,
        }
    terminal = task_status in _TASK_STREAM_TERMINAL_STATUSES
    if active_goal:
        return {"state": "verification_missing" if terminal else "pending"}
    return {"state": "legacy_unverified" if terminal else "not_started"}


def _runtime_trace_summary(
    *,
    task: dict[str, Any],
    runs: list[dict[str, Any]],
    nodes: list[dict[str, Any]],
    active_run: dict[str, Any] | None,
    current_node: dict[str, Any] | None,
    active_goal: dict[str, Any] | None,
    verification: dict[str, Any],
) -> dict[str, Any]:
    capabilities: dict[str, dict[str, str]] = {}
    capability_calls: list[dict[str, Any]] = []
    for node in nodes:
        capability = node.get("capability")
        if not isinstance(capability, dict):
            continue
        capability_type = str(capability.get("type") or "node")
        capability_id = str(capability.get("id") or "").strip()
        if not capability_id:
            continue
        key = f"{capability_type}:{capability_id}"
        capabilities[key] = {
            "type": capability_type,
            "id": capability_id,
            "label": str(capability.get("label") or capability_id)[:160],
            "status": str(node.get("status") or ""),
        }
        capability_calls.append(
            {
                key: value
                for key, value in {
                    "node_id": node.get("id"),
                    "parent_node_id": node.get("parent_node_id"),
                    "title": node.get("title"),
                    "type": capability_type,
                    "id": capability_id,
                    "label": str(capability.get("label") or capability_id)[:160],
                    "status": node.get("status"),
                    "status_message": node.get("status_message"),
                    "output_summary": node.get("output_summary"),
                    "error_summary": node.get("error_summary"),
                    "started_at": node.get("started_at"),
                    "finished_at": node.get("finished_at"),
                    "attempt": node.get("attempt"),
                }.items()
                if value not in (None, "")
            }
        )

    node_status_counts: dict[str, int] = {}
    for node in nodes:
        status = str(node.get("status") or "unknown")
        node_status_counts[status] = node_status_counts.get(status, 0) + 1

    goal_summary = (
        active_goal.get("summary")
        if active_goal and isinstance(active_goal.get("summary"), dict)
        else {}
    )
    objective = goal_summary.get("objective") if isinstance(goal_summary, dict) else {}
    if isinstance(objective, dict):
        objective_text = str(objective.get("statement") or "")
    else:
        objective_text = str(objective or "")
    deliverables = (
        goal_summary.get("deliverables")
        if isinstance(goal_summary.get("deliverables"), list)
        else []
    )
    artifacts = [
        item
        for item in db.json_loads(task.get("artifacts_json"), [])
        if isinstance(item, dict)
    ]
    public_artifacts = [_public_artifact(item) for item in artifacts]
    published_artifacts = [
        item for item in public_artifacts if str(item.get("delivery_status") or "") == "published"
    ]
    return {
        "task_status": str(task.get("status") or ""),
        "active_run_status": str((active_run or {}).get("status") or ""),
        "attempts": len(runs),
        "current_node": (
            {
                key: current_node.get(key)
                for key in ("id", "title", "kind", "status", "status_message", "capability")
                if current_node.get(key) not in (None, "")
            }
            if current_node
            else None
        ),
        "goal": {
            "status": str((active_goal or {}).get("status") or ""),
            "version": (active_goal or {}).get("version"),
            "objective": objective_text[:300],
            "deliverable_count": len(deliverables),
        },
        "capabilities": sorted(capabilities.values(), key=lambda item: (item["type"], item["id"])),
        "capability_calls": sorted(
            capability_calls,
            key=lambda item: (
                int(item.get("attempt") or 0),
                str(item.get("started_at") or item.get("finished_at") or ""),
                str(item.get("node_id") or ""),
            ),
        )[:50],
        "node_status_counts": node_status_counts,
        "verification_state": str(verification.get("state") or ""),
        "knowledge": _task_knowledge_summary(str(task.get("id") or "")),
        "artifacts": {
            "total": len(artifacts),
            "published": len(published_artifacts),
            "formats": sorted(
                {
                    str(item.get("kind") or item.get("format") or "").lower()
                    for item in artifacts
                    if str(item.get("kind") or item.get("format") or "").strip()
                }
            ),
            "items": public_artifacts[:12],
        },
    }


def _public_task(
    value: dict[str, Any],
    *,
    include_result: bool = True,
    include_attachments: bool = True,
) -> dict[str, Any]:
    """Return the stable public task shape and remove internal JSON/path fields."""
    internal_fields = {"result_json", "artifacts_json", "attachments_json"}
    result = {key: item for key, item in value.items() if key not in internal_fields}
    result['schema_version'] = 1
    if include_result:
        stored_result = db.json_loads(value.get("result_json"), {})
        if str(value.get("status") or "") == "failed":
            code = str(stored_result.get("error_code") or "")
            result["result"] = {
                "error_code": (
                    code
                    if re.fullmatch(r"[a-z0-9_.-]{1,80}", code)
                    else "task_execution_failed"
                )
            }
        else:
            result["result"] = _sanitize_public_payload(stored_result)
    result["artifacts"] = [
        _public_artifact(item)
        for item in db.json_loads(value.get("artifacts_json"), [])
        if isinstance(item, dict)
    ]
    if include_attachments:
        result["attachments"] = [
            _public_attachment(item)
            for item in db.json_loads(value.get("attachments_json"), [])
            if isinstance(item, dict)
        ]
    return result


def _task_or_404(task_id: str, *, write: bool = False) -> dict[str, Any]:
    task = db.query_one("SELECT * FROM tasks WHERE id = ?", (task_id,))
    if not task:
        raise HTTPException(status_code=404, detail="任务不存在")
    identity = auth_service.current_identity.get()
    if identity and identity["role"] != "admin" and (task.get("user_id") != identity["user_id"] or task.get("organization_id") != "local-org"):
        raise HTTPException(status_code=403, detail="无权访问此任务")
    _require_workspace_access(task.get("workspace") or "default", write=write)
    return task


def _runtime_projection(task_id: str) -> dict[str, Any]:
    task = _task_or_404(task_id)
    raw_runs = task_state.list_runs(task_id=task_id, limit=100)
    runs = [_public_task_run(item) for item in raw_runs]
    attempts = {item["id"]: item["attempt"] for item in runs}
    nodes: list[dict[str, Any]] = []
    for run in sorted(runs, key=lambda item: item["attempt"]):
        nodes.extend(
            _public_task_node(node, attempt=attempts.get(run["id"], 1))
            for node in task_state.list_nodes(run["id"])
        )
    checkpoints = [
        _public_checkpoint(
            item, attempt=attempts.get(item.get("run_id", ""), 1)
        )
        for item in task_state.list_checkpoints(task_id=task_id, include_state=False, limit=200)
    ]
    commands = [
        _public_task_command(item)
        for item in task_state.list_commands(task_id=task_id, limit=200)
    ]
    active_run = next(
        (
            item
            for item in runs
            if item["status"] in {"queued", "running", "paused", "waiting_approval"}
        ),
        None,
    )
    active_run_id = str((active_run or {}).get("id") or "")
    current_node: dict[str, Any] | None = None
    if active_run_id:
        current_node_id = str((active_run or {}).get("current_node_id") or "")
        current_node = next(
            (
                item for item in nodes
                if str(item.get("run_id") or "") == active_run_id
                and str(item.get("id") or "") == current_node_id
            ),
            None,
        )
        if current_node is None:
            running_nodes = [
                item for item in nodes
                if str(item.get("run_id") or "") == active_run_id
                and item.get("status") == "running"
            ]
            if running_nodes:
                current_node = max(
                    running_nodes,
                    key=lambda item: (
                        bool(item.get("parent_node_id")),
                        str(item.get("updated_at") or item.get("started_at") or ""),
                        int(item.get("sequence") or 0),
                    ),
                )
    focused_run = active_run or (runs[0] if runs else None)
    focused_run_id = str((focused_run or {}).get("id") or "")
    active_goal = _active_goal_projection(task_id, focused_run_id)
    verification = _verification_projection(
        task_status=str(task.get("status") or ""),
        run_id=focused_run_id,
        active_goal=active_goal,
    )
    return {
        "runs": runs,
        "nodes": nodes,
        "node_tree": _two_level_node_tree(nodes),
        "checkpoints": checkpoints,
        "commands": commands,
        "active_run": active_run,
        "current_node": current_node,
        "active_goal": active_goal,
        "verification": verification,
        "trace_summary": _runtime_trace_summary(
            task=task,
            runs=runs,
            nodes=nodes,
            active_run=active_run,
            current_node=current_node,
            active_goal=active_goal,
            verification=verification,
        ),
    }


def _active_run_or_409(task_id: str) -> dict[str, Any]:
    runtime_state = _runtime_projection(task_id)
    active = runtime_state.get("active_run")
    if not isinstance(active, dict):
        raise HTTPException(status_code=409, detail="当前任务没有可控制的运行尝试")
    return active


def _ensure_no_active_run(task_id: str) -> None:
    active = _runtime_projection(task_id).get("active_run")
    if isinstance(active, dict):
        raise HTTPException(
            status_code=409,
            detail=f"任务已有 {active.get('status', 'active')} 状态的运行尝试，请先取消或等待结束",
        )


def _complete_control_command(
    task_id: str,
    command_type: str,
    *,
    payload: dict[str, Any],
    run_id: str | None = None,
    result: dict[str, Any] | None = None,
) -> dict[str, Any]:
    command = task_state.enqueue_command(
        task_id,
        command_type,
        payload=payload,
        run_id=run_id,
    )
    claimed = task_state.claim_command(
        "api-control",
        task_id=task_id,
        run_id=run_id,
        command_types=[command_type],
    )
    if claimed:
        task_state.complete_command(claimed["id"], result=result or {"accepted": True})
    return task_state.get_command(command["id"]) or command


async def _retry_task(task_id: str) -> dict[str, Any]:
    task = _task_or_404(task_id, write=True)
    _ensure_no_active_run(task_id)
    if task.get("status") not in {"failed", "cancelled", "completed"}:
        raise HTTPException(status_code=409, detail="只有已结束的任务才能重试")
    run = task_state.create_run(task_id, metadata={"trigger": "retry", "dispatch_backend": "redis" if task_queue.enabled() else "local"})
    command = _complete_control_command(
        task_id,
        "retry",
        payload={},
        run_id=run["id"],
        result={"run_id": run["id"]},
    )
    db.execute(
        "UPDATE tasks SET status = 'queued', result_json = '{}', artifacts_json = '[]', "
        "updated_at = ? WHERE id = ?",
        (db.utc_now(), task_id),
    )
    db.insert_event(
        task_id,
        "retry_scheduled",
        "已创建重试运行",
        f"这是第 {run['attempt']} 次运行尝试。",
        {"run_id": run["id"], "attempt": run["attempt"]},
    )
    _schedule_runtime(task_id, run["id"])
    return {
        "ok": True,
        "run": _public_task_run(run),
        "command": _public_task_command(command),
    }


async def _resume_task_from_checkpoint(
    task_id: str,
    checkpoint_id: str | None,
    *,
    trigger: str,
) -> dict[str, Any]:
    _task_or_404(task_id, write=True)
    _ensure_no_active_run(task_id)
    checkpoint = (
        task_state.get_checkpoint(checkpoint_id, include_state=False)
        if checkpoint_id
        else (task_state.list_checkpoints(task_id=task_id, include_state=False, limit=1) or [None])[0]
    )
    if not checkpoint:
        raise HTTPException(status_code=404, detail="当前任务没有可恢复的检查点")
    if checkpoint.get("task_id") != task_id:
        raise HTTPException(status_code=400, detail="该检查点不属于当前任务")
    restored = task_state.restore_checkpoint(
        checkpoint["id"],
        restore_metadata={"requested_by": "api", "trigger": trigger},
    )
    run = task_state.create_run(
        task_id,
        resumed_from_checkpoint_id=checkpoint["id"],
        metadata={"trigger": trigger, "checkpoint_restore_audited": True, "dispatch_backend": "redis" if task_queue.enabled() else "local"},
    )
    command = _complete_control_command(
        task_id,
        "restore_checkpoint" if trigger == "restore_checkpoint" else "resume",
        payload={"checkpoint_id": checkpoint["id"]},
        run_id=run["id"],
        result={"run_id": run["id"], "checkpoint_id": checkpoint["id"]},
    )
    db.execute(
        "UPDATE tasks SET status = 'queued', result_json = '{}', artifacts_json = '[]', "
        "updated_at = ? WHERE id = ?",
        (db.utc_now(), task_id),
    )
    db.insert_event(
        task_id,
        "resume_scheduled",
        "已创建恢复运行",
        "新的运行尝试将从所选安全检查点继续。",
        {"run_id": run["id"], "checkpoint_id": checkpoint["id"], "restore_count": restored["restore_count"]},
    )
    _schedule_runtime(task_id, run["id"])
    return {
        "ok": True,
        "run": _public_task_run(run),
        "checkpoint": _public_checkpoint(
            restored, attempt=int(run.get("attempt") or 1)
        ),
        "command": _public_task_command(command),
    }


@app.get("/api/tasks/{task_id}/runtime")
def get_task_runtime(task_id: str) -> dict[str, Any]:
    return _runtime_projection(task_id)


@app.post("/api/tasks/{task_id}/commands", status_code=202)
async def send_task_command(task_id: str, payload: TaskCommandRequest) -> dict[str, Any]:
    _task_or_404(task_id, write=True)
    command_type = payload.type.strip().lower()
    if command_type not in {"message", "cancel", "retry", "resume", "restore_checkpoint"}:
        raise HTTPException(status_code=400, detail=f"不支持的任务指令：{command_type}")
    if command_type == "retry":
        return await _retry_task(task_id)
    if command_type in {"resume", "restore_checkpoint"}:
        return await _resume_task_from_checkpoint(
            task_id,
            str(payload.payload.get("checkpoint_id") or "") or None,
            trigger=command_type,
        )
    active = _active_run_or_409(task_id)
    if command_type == "cancel":
        try:
            command = task_state.request_cancel(
                task_id,
                run_id=active["id"],
                reason=str(payload.payload.get("reason") or "用户请求取消"),
                requested_by="user",
            )
            await task_queue.request_cancel(str(active["id"]))
        except RunIntakeClosed as exc:
            raise HTTPException(
                status_code=409,
                detail="当前运行已经结束，取消请求未加入旧运行。",
            ) from exc
        # Approval pauses do not have a worker loop left to claim the queued
        # cancel command.  Close that run immediately so the UI's Stop action
        # is deterministic even while a Skill/MCP approval card is visible.
        if active.get("status") in {"waiting_approval", "paused"}:
            try:
                task_state.commit_cancellation(
                    task_id=task_id,
                    run_id=active["id"],
                    result={"cancelled": True, "reason": "用户请求取消"},
                )
            except PublicationConflict as exc:
                raise HTTPException(status_code=409, detail="任务状态已变化，请刷新后重试") from exc
            command = task_state.get_command(command["id"]) or command
            active = task_state.get_run(active["id"]) or active
        return {
            "ok": True,
            "command": _public_task_command(command),
            "run": _public_task_run(active),
        }
    message = str(payload.payload.get("message") or "").strip()
    if not message:
        raise HTTPException(status_code=400, detail="追加指令不能为空")
    try:
        command = task_state.enqueue_command(
            task_id,
            "message",
            run_id=active["id"],
            payload={"message": message},
            priority=20,
        )
    except RunIntakeClosed as exc:
        raise HTTPException(
            status_code=409,
            detail=(
                "当前运行已经完成最终交付，这条追加要求未加入旧任务。"
                "请将其作为下一条消息发送。"
            ),
        ) from exc
    return {
        "ok": True,
        "command": _public_task_command(command),
        "run": _public_task_run(active),
    }


@app.post("/api/tasks/{task_id}/cancel", status_code=202)
async def cancel_task(task_id: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    _task_or_404(task_id, write=True)
    active = _active_run_or_409(task_id)
    values = payload or {}
    try:
        command = task_state.request_cancel(
            task_id,
            run_id=active["id"],
            reason=str(values.get("reason") or "用户请求取消"),
            requested_by="user",
        )
        await task_queue.request_cancel(active['id'])
    except RunIntakeClosed as exc:
        raise HTTPException(
            status_code=409,
            detail="当前运行已经结束，取消请求未加入旧运行。",
        ) from exc
    if active.get("status") in {"waiting_approval", "paused"}:
        try:
            task_state.commit_cancellation(
                task_id=task_id,
                run_id=active["id"],
                result={"cancelled": True, "reason": str(values.get("reason") or "用户请求取消")},
            )
        except PublicationConflict as exc:
            raise HTTPException(status_code=409, detail="任务状态已变化，请刷新后重试") from exc
        command = task_state.get_command(command["id"]) or command
        active = task_state.get_run(active["id"]) or active
    return {
        "ok": True,
        "command": _public_task_command(command),
        "run": _public_task_run(active),
    }


@app.post("/api/tasks/{task_id}/retry", status_code=202)
async def retry_task(task_id: str) -> dict[str, Any]:
    return await _retry_task(task_id)


@app.post("/api/tasks/{task_id}/resume", status_code=202)
async def resume_task(task_id: str, payload: TaskResumeRequest | None = None) -> dict[str, Any]:
    return await _resume_task_from_checkpoint(
        task_id,
        payload.checkpoint_id if payload else None,
        trigger="resume",
    )


@app.post("/api/tasks/{task_id}/checkpoints/{checkpoint_id}/restore", status_code=202)
async def restore_task_checkpoint(
    task_id: str,
    checkpoint_id: str,
    payload: CheckpointRestoreRequest | None = None,
) -> dict[str, Any]:
    _ = payload
    return await _resume_task_from_checkpoint(
        task_id,
        checkpoint_id,
        trigger="restore_checkpoint",
    )


def _validate_loop_bindings(agent_id: str, model_id: str, workspace_id: str = 'default') -> None:
    identity = auth_service.current_identity.get()
    if identity and not auth_service.agent_access(agent_id, identity, workspace_id):
        raise HTTPException(status_code=403, detail='无权使用所选智能体')
    if not db.query_one("SELECT id FROM agents WHERE id = ?", (agent_id,)):
        raise HTTPException(status_code=400, detail="所选智能体不存在")
    _ensure_model_ready(model_id, label="所选模型")


def _validate_loop_trigger(config: dict[str, Any], webhook_secret_ciphertext: str) -> None:
    try:
        validate_trigger_config(
            str(config.get("trigger_type") or "interval"),
            cron_expression=str(config.get("cron_expression") or ""),
            once_at=str(config.get("once_at") or ""),
            webhook_secret_configured=bool(webhook_secret_ciphertext),
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.get("/api/loops")
def list_loops(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
) -> list[dict[str, Any]]:
    scope = _api_scope(organization_id, workspace_id, user_id)
    return [
        serialize_loop(row)
        for row in db.query_all(
            """SELECT * FROM loops WHERE organization_id = ? AND workspace_id = ? AND user_id = ?
               ORDER BY created_at DESC""",
            (scope.organization_id, scope.workspace_id, scope.user_id),
        )
    ]


@app.post("/api/loops")
def create_loop(payload: LoopCreate) -> dict[str, Any]:
    scope = _api_scope(payload.organization_id, payload.workspace_id, payload.user_id)
    payload = payload.model_copy(update={"organization_id": scope.organization_id, "user_id": scope.user_id})
    _validate_loop_bindings(payload.agent_id, payload.model_id, payload.workspace_id)
    loop_id = (payload.id or ("loop_" + uuid.uuid4().hex[:12])).strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{2,80}", loop_id):
        raise HTTPException(status_code=400, detail="Loop ID 只能包含字母、数字、下划线和连字符")
    if db.query_one("SELECT id FROM loops WHERE id = ?", (loop_id,)):
        raise HTTPException(status_code=409, detail="Loop ID 已存在")
    encrypted_secret = secret_store.encrypt(payload.webhook_secret) if payload.webhook_secret else ""
    config = payload.model_dump(exclude={"webhook_secret", "initial_state", "auto_start", "id"})
    _validate_loop_trigger(config, encrypted_secret)
    now = db.utc_now()
    status = "active" if payload.auto_start else "paused"
    next_run_at = ""
    if payload.auto_start:
        try:
            next_run_at = next_schedule_at(config, immediate_interval=True)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.execute(
        """INSERT INTO loops(
               id, name, prompt, agent_id, model_id, trigger_type, interval_seconds,
               cron_expression, once_at, organization_id, workspace_id, user_id,
               webhook_secret_ciphertext, webhook_tolerance_seconds, status, max_runs,
               max_failures, max_attempts, retry_backoff_seconds, state_json, last_diff_json,
               next_run_at, created_at, updated_at
           ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '{}', ?, ?, ?)""",
        (
            loop_id, payload.name, payload.prompt, payload.agent_id, payload.model_id,
            payload.trigger_type, payload.interval_seconds, payload.cron_expression, payload.once_at,
            payload.organization_id, payload.workspace_id, payload.user_id, encrypted_secret,
            payload.webhook_tolerance_seconds, status, payload.max_runs, payload.max_failures,
            payload.max_attempts, payload.retry_backoff_seconds, db.json_dumps(payload.initial_state),
            next_run_at, now, now,
        ),
    )
    return serialize_loop(db.query_one("SELECT * FROM loops WHERE id = ?", (loop_id,)) or {})


def _require_owned_resource(row: dict[str, Any]) -> None:
    identity = auth_service.current_identity.get()
    if identity and identity["role"] != "admin" and (row.get("user_id") != identity["user_id"] or row.get("organization_id") != "local-org"):
        raise HTTPException(status_code=403, detail="无权访问此资源")
    _require_workspace_access(row.get("workspace_id") or "default", write=bool(identity and identity.get("request_method") not in {"GET", "HEAD", "OPTIONS"}))


def _loop_or_404(loop_id: str) -> dict[str, Any]:
    row = db.query_one("SELECT * FROM loops WHERE id=?", (loop_id,))
    if not row:
        raise HTTPException(status_code=404, detail="Loop not found")
    _require_owned_resource(row)
    return row


@app.get("/api/loops/{loop_id}")
def get_loop(loop_id: str) -> dict[str, Any]:
    row = _loop_or_404(loop_id)
    if not row:
        raise HTTPException(status_code=404, detail="Loop not found")
    result = serialize_loop(row)
    result["runs"] = [
        serialize_run(item)
        for item in db.query_all(
            "SELECT * FROM loop_runs WHERE loop_id = ? ORDER BY run_number DESC, attempt DESC",
            (loop_id,),
        )
    ]
    return result


@app.put("/api/loops/{loop_id}")
def update_loop(loop_id: str, payload: LoopUpdate) -> dict[str, Any]:
    current = _loop_or_404(loop_id)
    if not current:
        raise HTTPException(status_code=404, detail="Loop not found")
    incoming = payload.model_dump(exclude_unset=True)
    webhook_secret = incoming.pop("webhook_secret", None)
    state = incoming.pop("state", None)
    merged = {**current, **incoming}
    _validate_loop_bindings(merged["agent_id"], merged["model_id"], merged.get('workspace_id') or 'default')
    encrypted_secret = str(current.get("webhook_secret_ciphertext") or "")
    if webhook_secret is not None:
        encrypted_secret = secret_store.encrypt(webhook_secret)
    _validate_loop_trigger(merged, encrypted_secret)
    state_json = db.json_dumps(state) if state is not None else str(current.get("state_json") or "{}")
    next_run_at = str(current.get("next_run_at") or "")
    if current["status"] == "active":
        try:
            next_run_at = next_schedule_at(merged)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.execute(
        """UPDATE loops SET name = ?, prompt = ?, agent_id = ?, model_id = ?, trigger_type = ?,
               interval_seconds = ?, cron_expression = ?, once_at = ?, webhook_secret_ciphertext = ?,
               webhook_tolerance_seconds = ?, max_runs = ?, max_failures = ?, max_attempts = ?,
               retry_backoff_seconds = ?, state_json = ?, next_run_at = ?, updated_at = ?
           WHERE id = ?""",
        (
            merged["name"], merged["prompt"], merged["agent_id"], merged["model_id"],
            merged.get("trigger_type") or "interval", merged["interval_seconds"],
            merged.get("cron_expression") or "", merged.get("once_at") or "", encrypted_secret,
            merged.get("webhook_tolerance_seconds") or 300, merged["max_runs"],
            merged["max_failures"], merged.get("max_attempts") or 1,
            merged.get("retry_backoff_seconds") or 0, state_json, next_run_at, db.utc_now(), loop_id,
        ),
    )
    return serialize_loop(db.query_one("SELECT * FROM loops WHERE id = ?", (loop_id,)) or {})


@app.delete("/api/loops/{loop_id}")
def delete_loop(loop_id: str) -> dict[str, bool]:
    current = _loop_or_404(loop_id)
    if not current:
        raise HTTPException(status_code=404, detail="Loop not found")
    if current["status"] == "running":
        raise HTTPException(status_code=409, detail="循环任务运行中，请先等待本轮结束并暂停")
    db.execute("DELETE FROM loop_runs WHERE loop_id = ?", (loop_id,))
    db.execute("DELETE FROM automation_trigger_events WHERE loop_id = ?", (loop_id,))
    db.execute("DELETE FROM loops WHERE id = ?", (loop_id,))
    return {"ok": True}


@app.post("/api/loops/{loop_id}/start")
def start_loop(loop_id: str) -> dict[str, Any]:
    current = _loop_or_404(loop_id)
    if not current:
        raise HTTPException(status_code=404, detail="Loop not found")
    if current["status"] == "running":
        raise HTTPException(status_code=409, detail="循环任务正在运行")
    if int(current["run_count"]) >= int(current["max_runs"]):
        raise HTTPException(status_code=409, detail="已达到最大轮数，请提高最大轮数后再启动")
    if (current.get("trigger_type") or "interval") == "once" and int(current["run_count"]):
        raise HTTPException(status_code=409, detail="一次性自动化已经执行过，请新建自动化")
    _validate_loop_trigger(current, str(current.get("webhook_secret_ciphertext") or ""))
    now = db.utc_now()
    try:
        next_run_at = next_schedule_at(current, immediate_interval=True)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    db.execute(
        "UPDATE loops SET status = 'active', next_run_at = ?, updated_at = ? WHERE id = ?",
        (next_run_at, now, loop_id),
    )
    return serialize_loop(db.query_one("SELECT * FROM loops WHERE id = ?", (loop_id,)) or {})


@app.post("/api/loops/{loop_id}/pause")
def pause_loop(loop_id: str) -> dict[str, Any]:
    _loop_or_404(loop_id)
    if not db.query_one("SELECT id FROM loops WHERE id = ?", (loop_id,)):
        raise HTTPException(status_code=404, detail="Loop not found")
    db.execute("UPDATE loops SET status = 'paused', next_run_at = '', updated_at = ? WHERE id = ?", (db.utc_now(), loop_id))
    return serialize_loop(db.query_one("SELECT * FROM loops WHERE id = ?", (loop_id,)) or {})


@app.post("/api/loops/{loop_id}/run", status_code=202)
async def run_loop_now(loop_id: str) -> dict[str, Any]:
    _loop_or_404(loop_id)
    if not db.query_one("SELECT id FROM loops WHERE id = ?", (loop_id,)):
        raise HTTPException(status_code=404, detail="Loop not found")
    try:
        background = loop_scheduler.dispatch_once(loop_id)
        _runtime_tasks.add(background)
        background.add_done_callback(_runtime_tasks.discard)
        return {"accepted": True, "loop_id": loop_id, "status": "queued"}
    except RuntimeError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc


@app.get("/api/loops/{loop_id}/runs")
def list_loop_runs(loop_id: str) -> list[dict[str, Any]]:
    _loop_or_404(loop_id)
    if not db.query_one("SELECT id FROM loops WHERE id = ?", (loop_id,)):
        raise HTTPException(status_code=404, detail="Loop not found")
    return [
        serialize_run(item)
        for item in db.query_all(
            "SELECT * FROM loop_runs WHERE loop_id = ? ORDER BY run_number DESC, attempt DESC",
            (loop_id,),
        )
    ]


@app.get("/api/loops/{loop_id}/trigger-events")
def list_loop_trigger_events(loop_id: str) -> list[dict[str, Any]]:
    _loop_or_404(loop_id)
    if not db.query_one("SELECT id FROM loops WHERE id = ?", (loop_id,)):
        raise HTTPException(status_code=404, detail="Loop not found")
    return [
        serialize_trigger_event(item)
        for item in db.query_all(
            "SELECT * FROM automation_trigger_events WHERE loop_id = ? ORDER BY received_at DESC",
            (loop_id,),
        )
    ]


def _webhook_timestamp(value: str) -> float:
    try:
        return float(value)
    except ValueError:
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            if not parsed.tzinfo:
                parsed = parsed.replace(tzinfo=timezone.utc)
            return parsed.timestamp()
        except ValueError as exc:
            raise HTTPException(status_code=401, detail="Webhook 时间戳无效") from exc


@app.post("/api/loops/{loop_id}/webhook", status_code=202)
async def trigger_loop_webhook(loop_id: str, request: Request) -> dict[str, Any]:
    loop = _loop_or_404(loop_id)
    if not loop:
        raise HTTPException(status_code=404, detail="Loop not found")
    if (loop.get("trigger_type") or "interval") != "webhook":
        raise HTTPException(status_code=409, detail="这个自动化不是 Webhook 触发类型")
    if loop["status"] != "active":
        raise HTTPException(status_code=409, detail="Webhook 自动化已暂停")
    raw_body = await request.body()
    if len(raw_body) > 1024 * 1024:
        raise HTTPException(status_code=413, detail="Webhook 请求正文不能超过 1MB")
    timestamp = request.headers.get("x-automation-timestamp") or request.headers.get("x-webhook-timestamp") or ""
    provided_signature = request.headers.get("x-automation-signature") or request.headers.get("x-webhook-signature") or ""
    idempotency_key = request.headers.get("idempotency-key", "").strip()
    if not timestamp or not provided_signature:
        raise HTTPException(status_code=401, detail="缺少 Webhook 时间戳或签名")
    if not idempotency_key or len(idempotency_key) > 200:
        raise HTTPException(status_code=400, detail="必须提供 1-200 位 Idempotency-Key")
    tolerance = int(loop.get("webhook_tolerance_seconds") or 300)
    if abs(time.time() - _webhook_timestamp(timestamp)) > tolerance:
        raise HTTPException(status_code=401, detail="Webhook 时间戳超出允许时间窗")
    encrypted_secret = str(loop.get("webhook_secret_ciphertext") or "")
    if not encrypted_secret:
        raise HTTPException(status_code=409, detail="Webhook 签名密钥未配置")
    try:
        secret = secret_store.decrypt(encrypted_secret)
    except RuntimeError as exc:
        raise HTTPException(status_code=500, detail="Webhook 签名密钥无法解密") from exc
    expected = hmac.new(
        secret.encode("utf-8"), timestamp.encode("utf-8") + b"." + raw_body, hashlib.sha256
    ).hexdigest()
    supplied = provided_signature.removeprefix("sha256=").strip().lower()
    if not hmac.compare_digest(expected, supplied):
        raise HTTPException(status_code=401, detail="Webhook 签名验证失败")
    if auth_service.enabled():
        owner = db.query_one('SELECT id AS user_id,role FROM users WHERE id=? AND enabled=1', (loop['user_id'],))
        if not owner or loop['organization_id'] != 'local-org' or auth_service.workspace_access(loop['workspace_id'], owner) not in {'owner', 'member'}:
            raise HTTPException(status_code=403, detail='自动化所属用户已失去执行权限')
        request.state.audit_user_id = owner['user_id']
    try:
        event, duplicate = create_webhook_event(loop, idempotency_key, raw_body)
    except IdempotencyConflictError as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc
    queued = True
    if event.get("status") == "accepted" and not loop_scheduler.is_busy(loop_id):
        try:
            background = loop_scheduler.dispatch_once(
                loop_id, scheduled=True, trigger_event_id=event["id"], trigger_type="webhook"
            )
            _runtime_tasks.add(background)
            background.add_done_callback(_runtime_tasks.discard)
            queued = False
        except RuntimeError:
            queued = True
    return {"accepted": True, "duplicate": duplicate, "queued": queued, "event": event}


@app.get("/api/notifications")
def list_notifications(
    organization_id: str = "local-org",
    workspace_id: str = "default",
    user_id: str = "local-user",
    status: str = "",
    limit: int = 100,
) -> list[dict[str, Any]]:
    scope = _api_scope(organization_id, workspace_id, user_id)
    if status and status not in {"unread", "read"}:
        raise HTTPException(status_code=400, detail="通知状态只能是 unread 或 read")
    capped_limit = max(1, min(int(limit), 500))
    sql = """SELECT * FROM notifications
             WHERE organization_id = ? AND workspace_id = ? AND user_id = ?"""
    params: list[Any] = [scope.organization_id, scope.workspace_id, scope.user_id]
    if status:
        sql += " AND status = ?"
        params.append(status)
    sql += " ORDER BY created_at DESC LIMIT ?"
    params.append(capped_limit)
    return [serialize_notification(item) for item in db.query_all(sql, params)]


@app.post("/api/notifications/{notification_id}/read")
def read_notification(notification_id: str) -> dict[str, Any]:
    current = db.query_one("SELECT * FROM notifications WHERE id = ?", (notification_id,))
    if not current:
        raise HTTPException(status_code=404, detail="Notification not found")
    _require_owned_resource(current)
    now = db.utc_now()
    db.execute(
        "UPDATE notifications SET status = 'read', read_at = ? WHERE id = ?",
        (now, notification_id),
    )
    return serialize_notification(
        db.query_one("SELECT * FROM notifications WHERE id = ?", (notification_id,)) or {}
    )


@app.post("/api/tasks")
async def create_task(payload: TaskCreate, request: Request = None) -> dict[str, Any]:
    _require_workspace_access(payload.workspace, write=True)
    identity = auth_service.current_identity.get()
    if identity:
        payload = payload.model_copy(update={"user_id": identity["user_id"], "organization_id": "local-org"})
        if payload.parent_task_id:
            _task_or_404(payload.parent_task_id)
    submission_key = request.headers.get('idempotency-key', '') if request is not None else ''
    if len(submission_key) > 200:
        raise HTTPException(status_code=400, detail='Idempotency-Key 不能超过 200 字符')
    key_hash = hashlib.sha256(submission_key.encode()).hexdigest() if submission_key else ''
    payload_hash = hashlib.sha256(payload.model_dump_json().encode()).hexdigest()
    attachments = []
    total_attachment_bytes = 0
    max_attachment_bytes = int(os.getenv('APP_MAX_TASK_ATTACHMENTS_MB', '40')) * 1024 * 1024
    for upload_id in payload.attachment_ids:
        upload = db.query_one("SELECT * FROM uploads WHERE id = ?", (upload_id,))
        if identity and identity["role"] != "admin":
            owner = db.query_one("SELECT user_id FROM upload_owners WHERE upload_id=?", (upload_id,))
            if not owner or owner["user_id"] != identity["user_id"]:
                raise HTTPException(status_code=403, detail="无权使用此附件")
        if not upload:
            raise HTTPException(status_code=404, detail='附件不存在，请重新上传')
        total_attachment_bytes += int(upload.get('size') or 0)
        if total_attachment_bytes > max_attachment_bytes:
            raise HTTPException(status_code=413, detail='本次任务的附件总大小超过限制')
        attachments.append(upload)
    if payload.executor_type == "agent" and not db.query_one(
        "SELECT id FROM agents WHERE id = ?", (payload.agent_id,)
    ):
        raise HTTPException(status_code=400, detail="所选智能体不存在")
    scope = _api_scope(payload.organization_id, payload.workspace, payload.user_id)
    executor_id = payload.executor_id or payload.agent_id
    selected_agent_id = payload.agent_id
    expert_selection: dict[str, Any] | None = None
    if payload.executor_type == "team":
        recommendation: dict[str, Any] | None = None
        if payload.executor_id:
            team = expert_team_service.get_team(payload.executor_id, scope)
        else:
            try:
                recommendation = expert_team_service.recommend_team(payload.message, scope)
                team = recommendation["team"]
            except (ExpertNotFoundError, ExpertValidationError) as exc:
                raise _expert_http_error(exc) from exc
        if not team or not team.get("enabled"):
            raise HTTPException(status_code=400, detail="所选专家团不存在、不可见或已停用")
        selected_agent_id = team["supervisor_agent_id"]
        executor_id = str(team["id"])
        expert_selection = _public_expert_selection(
            team,
            scope,
            automatic=not bool(payload.executor_id),
            recommendation=recommendation,
        )
    selected_agent = db.query_one("SELECT * FROM agents WHERE id = ?", (selected_agent_id,))
    if identity and not auth_service.agent_access(selected_agent_id, identity, payload.workspace):
        raise HTTPException(status_code=403, detail='无权使用所选智能体')
    fallback_model_id = str((selected_agent or {}).get("model") or "")
    _ensure_model_ready(payload.model_id or fallback_model_id or "deterministic", label="所选模型")
    if payload.executor_type == "team":
        try:
            task, run, team_run = expert_team_service.create_task_and_run(
                executor_id,
                scope,
                message=payload.message,
                model_id=payload.model_id,
                conversation_id=payload.conversation_id,
                attachments=attachments,
                parent_task_id=payload.parent_task_id or "",
                submission_key=key_hash,
                submission_hash=payload_hash,
            )
        except (ExpertNotFoundError, ExpertConflictError, ExpertValidationError) as exc:
            raise _expert_http_error(exc) from exc
        reused_submission = task.pop('_submission_reused', False)
        if expert_selection and not reused_submission:
            emit(
                task["id"],
                "expert_selection",
                "已选择参与专家",
                (
                    f"已自动匹配“{expert_selection['team_name']}”，"
                    if expert_selection["selection_mode"] == "automatic"
                    else f"使用已指定的“{expert_selection['team_name']}”，"
                )
                + f"由 {len(expert_selection['members'])} 位专家并行分析，再由主管汇总。",
                expert_selection,
            )
        if not reused_submission:
            _schedule_team_run(team_run["id"])
        return {
            **_public_task(task),
            "result": {},
            "artifacts": [],
            "run": _public_runtime_record(run),
            "team_run": _public_runtime_record(team_run),
            "expert_selection": expert_selection,
        }
    duplicate = False
    with task_state.transaction(write=True) as conn:
        previous = conn.execute('SELECT * FROM task_submissions WHERE organization_id=? AND user_id=? AND key_hash=?', (payload.organization_id,payload.user_id,key_hash)).fetchone() if key_hash else None
        if previous:
            if previous['payload_hash'] != payload_hash:
                raise HTTPException(status_code=409, detail='相同 Idempotency-Key 已用于不同任务内容')
            task = dict(conn.execute('SELECT * FROM tasks WHERE id=?', (previous['task_id'],)).fetchone())
            run = task_state._serialize_run(dict(conn.execute('SELECT * FROM task_runs WHERE id=?', (previous['run_id'],)).fetchone()))
            duplicate = True
        else:
            task = create_task_record(
                payload.message,
                selected_agent_id,
                payload.workspace,
                connection=conn,
                attachments=attachments,
                model_id=payload.model_id,
                conversation_id=payload.conversation_id,
                organization_id=payload.organization_id,
                user_id=payload.user_id,
                parent_task_id=payload.parent_task_id or "",
                executor_type=payload.executor_type,
                executor_id=executor_id,
                execution_engine=payload.execution_engine,
            )
            run = task_state.create_run_in_transaction(
                conn,
                task["id"],
                metadata={
                    "trigger": "user",
                    "dispatch_backend": "redis" if task_queue.enabled() else "local",
                    "agent_id": selected_agent_id,
                    "workspace": payload.workspace,
                    "organization_id": payload.organization_id,
                    "user_id": payload.user_id,
                    "executor_type": payload.executor_type,
                    "executor_id": executor_id,
                    "execution_engine": payload.execution_engine,
                },
            )
            if key_hash:
                conn.execute('INSERT INTO task_submissions(organization_id,user_id,key_hash,payload_hash,task_id,run_id) VALUES(?,?,?,?,?,?)', (payload.organization_id,payload.user_id,key_hash,payload_hash,task['id'],run['id']))
    if not duplicate:
        _schedule_runtime(task["id"], run["id"])
    return {
        **_public_task(task),
        "result": {},
        "artifacts": [],
        "run": _public_runtime_record(run),
    }


@app.get("/api/tasks/{task_id}")
def get_task(task_id: str) -> dict[str, Any]:
    task = _task_or_404(task_id)
    events = db.query_all("SELECT * FROM task_events WHERE task_id = ? AND type != 'answer_delta' ORDER BY id", (task_id,))
    return {
        **_public_task(task),
        "events": [_public_event(e) for e in events if _is_public_task_event(e)],
        "runtime": _runtime_projection(task_id),
    }


@app.get("/api/conversations/{conversation_id}/messages")
def get_conversation_messages(conversation_id: str) -> dict[str, Any]:
    identity = auth_service.current_identity.get()
    owner_clause = " AND user_id=? AND organization_id='local-org'" if identity and identity["role"] != "admin" else ""
    rows = db.query_all(
        f"SELECT * FROM tasks WHERE conversation_id = ?{owner_clause} ORDER BY created_at, id LIMIT 100",
        (conversation_id, identity["user_id"]) if owner_clause else (conversation_id,),
    )
    if identity:
        rows = [row for row in rows if auth_service.workspace_access(row.get("workspace") or "default", identity)]
    messages: list[dict[str, Any]] = []
    for task in rows:
        messages.append({"role": "user", "content": task["message"], "task_id": task["id"]})
        answer = db.query_one(
            "SELECT id, content FROM task_events WHERE task_id = ? AND type = 'answer' ORDER BY id DESC LIMIT 1",
            (task["id"],),
        )
        if answer:
            messages.append({"role": "assistant", "content": answer["content"], "task_id": task["id"], "event_id": answer["id"]})
            continue
        # Keep failures visible after a page refresh without presenting them
        # as a successful assistant answer.  The frontend renders this
        # structured message as the same error card used by the live stream.
        error = db.query_one(
            "SELECT * FROM task_events WHERE task_id = ? AND type = 'error' ORDER BY id DESC LIMIT 1",
            (task["id"],),
        )
        if error:
            public_error = _public_event(error)
            messages.append({
                "role": "system",
                "message_type": "error",
                "content": public_error.get("content") or "任务执行未完成，请检查模型、参数或工具配置后重试。",
                "title": public_error.get("title") or "任务未完成",
                "data": public_error.get("data") or {},
                "task_id": task["id"],
                "event_id": public_error.get("id"),
            })
    return {"conversation_id": conversation_id, "messages": messages}


@app.post("/api/uploads")
async def upload_file(file: UploadFile = File(...)) -> dict[str, Any]:
    max_bytes = int(os.getenv("APP_MAX_UPLOAD_MB", "20")) * 1024 * 1024
    raw = await file.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise HTTPException(status_code=413, detail="文件超过上传大小限制")
    original = Path(file.filename or "upload.bin").name
    safe_name = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]", "_", original)[:180] or "upload.bin"
    upload_id = "upl_" + uuid.uuid4().hex[:12]
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    path = UPLOAD_DIR / f"{upload_id}_{safe_name}"
    path.write_bytes(raw)
    record = {"id": upload_id, "name": original, "content_type": file.content_type or "application/octet-stream", "size": len(raw), "path": str(path), "created_at": db.utc_now()}
    db.execute("INSERT INTO uploads(id, name, content_type, size, path, created_at) VALUES (?, ?, ?, ?, ?, ?)", tuple(record.values()))
    identity = auth_service.current_identity.get()
    if identity:
        db.execute("INSERT INTO upload_owners(upload_id,user_id) VALUES(?,?)", (upload_id, identity["user_id"]))
    return _public_attachment(record)


@app.get("/api/tasks/{task_id}/events")
def get_task_events(task_id: str, after_id: int = 0) -> list[dict[str, Any]]:
    _task_or_404(task_id)
    if after_id < 0:
        raise HTTPException(status_code=400, detail="事件游标必须是非负整数")
    events = db.query_all("SELECT * FROM task_events WHERE task_id = ? AND id > ? ORDER BY id", (task_id, after_id))
    return [_public_event(e) for e in events if _is_public_task_event(e)]


@app.get("/api/tasks/{task_id}/events/stream")
async def stream_task_events(
    task_id: str,
    request: Request,
    cursor: int | None = None,
    after_id: int | None = None,
):
    _task_or_404(task_id)
    initial_cursor = _task_event_cursor(request, cursor=cursor, after_id=after_id)

    async def event_generator():
        async with event_notifications.notification_waiter(task_id) as wait_for_event:
            last_id = initial_cursor
            idle_rounds = 0
            while True:
                if await request.is_disconnected():
                    break
                if auth_service.enabled():
                    if not auth_service.get_session(request.cookies.get(auth_service.SESSION_COOKIE)):
                        return
                    try:
                        _task_or_404(task_id)
                    except HTTPException:
                        return
                events = db.query_all("SELECT * FROM task_events WHERE task_id = ? AND id > ? ORDER BY id", (task_id, last_id))
                for event in events:
                    last_id = int(event["id"])
                    if not _is_public_task_event(event):
                        continue
                    payload = _public_event(event)
                    yield f"id: {last_id}\nevent: task_event\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
                task = db.query_one("SELECT status FROM tasks WHERE id = ?", (task_id,))
                status = str((task or {}).get("status") or "")
                should_close = status in _TASK_STREAM_TERMINAL_STATUSES or status == "waiting_approval"
                if should_close and not events:
                    idle_rounds += 1
                    if idle_rounds > 1:
                        payload = {
                            "task_id": task_id,
                            "status": status,
                            "terminal": status in _TASK_STREAM_TERMINAL_STATUSES,
                            "cursor": last_id,
                        }
                        yield f"event: task_status\ndata: {json.dumps(payload, ensure_ascii=False)}\n\n"
                        break
                else:
                    idle_rounds = 0
                await wait_for_event()
    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "X-Accel-Buffering": "no",
        },
    )


async def _resume_after_approval_safely(
    task_id: str,
    approved: bool,
    note: str,
    command_id: str,
) -> None:
    try:
        await runtime.resume_after_approval(
            task_id,
            approved,
            note,
            command_id=command_id,
        )
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        message = "审批后续处理失败，任务已安全终止。"
        task = db.query_one("SELECT * FROM tasks WHERE id = ?", (task_id,)) or {}
        if not task:
            return
        result = db.json_loads(task.get("result_json"), {})
        failure_result = {
            **result,
            "error": "approval_continuation_failed",
            "error_type": exc.__class__.__name__,
            "summary": message,
        }
        active = next(
            (
                item
                for item in task_state.list_runs(task_id=task_id)
                if item["status"] in {"running", "paused", "waiting_approval"}
            ),
            None,
        )
        if task.get("status") not in {"completed", "failed", "cancelled"} and active:
            try:
                task_state.commit_failure(
                    task_id=task_id,
                    run_id=active["id"],
                    error={
                        "message": message,
                        "error_type": exc.__class__.__name__,
                    },
                    result=failure_result,
                )
                return
            except Exception:
                pass
        emit(
            task_id,
            "error",
            "审批后续处理失败",
            message,
            {"error_type": exc.__class__.__name__},
        )
        if task.get("status") not in {"completed", "failed", "cancelled"}:
            db.update_task_status(task_id, "failed", result=failure_result)


def _latest_approval_command(
    task_id: str,
    *,
    run_id: str = "",
    approval_id: str = "",
) -> dict[str, Any] | None:
    params: list[Any] = [task_id]
    sql = (
        "SELECT id FROM task_commands "
        "WHERE task_id = ? AND command_type = 'approval'"
    )
    if run_id:
        sql += " AND (run_id = ? OR run_id IS NULL)"
        params.append(run_id)
    sql += " ORDER BY created_at DESC, id DESC LIMIT 100"
    for row in db.query_all(sql, tuple(params)):
        command = task_state.get_command(str(row["id"]))
        if command is None:
            continue
        command_approval_id = str(
            (command.get("payload") or {}).get("approval_id") or ""
        )
        if approval_id and command_approval_id != approval_id:
            continue
        return command
    return None


def _assert_same_approval_decision(
    command: Mapping[str, Any], payload: ApprovalRequest
) -> None:
    stored_payload = command.get("payload") or {}
    stored_approved = stored_payload.get("approved")
    stored_note = str(stored_payload.get("note") or "")
    if not isinstance(stored_approved, bool):
        raise HTTPException(
            status_code=409,
            detail="已有审批决定格式异常，无法覆盖。",
        )
    if stored_approved != payload.approved:
        raise HTTPException(
            status_code=409,
            detail="已有相反的审批决定，不能重复修改。",
        )
    if stored_note != payload.note:
        raise HTTPException(
            status_code=409,
            detail="审批决定已存在，不能修改原审批备注。",
        )


def _schedule_approval_continuation(
    task_id: str,
    payload: ApprovalRequest,
    command_id: str,
) -> None:
    command = task_state.get_command(command_id)
    run = task_state.get_run(command["run_id"]) if command and command.get("run_id") else None
    if run and (run.get("metadata") or {}).get("dispatch_backend") == "redis":
        db.execute("INSERT OR IGNORE INTO approval_dispatch(command_id,task_id,run_id,created_at) VALUES(?,?,?,?)", (command_id, task_id, run["id"], db.utc_now()))
        return
    background = asyncio.create_task(
        _resume_after_approval_safely(
            task_id,
            payload.approved,
            payload.note,
            command_id,
        )
    )
    _runtime_tasks.add(background)
    background.add_done_callback(_runtime_tasks.discard)


@app.post("/api/tasks/{task_id}/approve")
async def approve_task(task_id: str, payload: ApprovalRequest) -> dict[str, Any]:
    task = _task_or_404(task_id, write=True)
    if task.get("status") != "waiting_approval":
        previous = _latest_approval_command(task_id)
        if previous and previous.get("status") == "completed":
            _assert_same_approval_decision(previous, payload)
            return {"ok": True, "duplicate": True, "command": previous}
        raise HTTPException(status_code=409, detail="当前任务不处于等待审批状态")
    active = _active_run_or_409(task_id)
    result = db.json_loads(task.get("result_json"), {})
    active_metadata = active.get("metadata") or {}
    recommendation_decision = result.get("skill_recommendation_decision")
    generic_decision = result.get("approval_resolution_proof")
    pending_policy = active_metadata.get("pending_policy_approval")
    pending_recommendation = active_metadata.get("pending_skill_recommendation")
    durable_approval_id = str(
        result.get("policy_approval_id")
        or result.get("skill_recommendation_approval_id")
        or (
            recommendation_decision.get("approval_id")
            if isinstance(recommendation_decision, Mapping)
            else ""
        )
        or (
            generic_decision.get("approval_id")
            if isinstance(generic_decision, Mapping)
            else ""
        )
        or (
            pending_policy.get("approval_id")
            if isinstance(pending_policy, Mapping)
            else ""
        )
        or (
            pending_recommendation.get("approval_id")
            if isinstance(pending_recommendation, Mapping)
            else ""
        )
        or ""
    )
    existing = _latest_approval_command(
        task_id,
        run_id=str(active["id"]),
        approval_id=durable_approval_id,
    )
    if existing is not None:
        existing_result = existing.get("result") or {}
        current_generic_command = str(
            (generic_decision or {}).get("command_id")
            if isinstance(generic_decision, Mapping)
            else ""
        )
        belongs_to_current_wait = bool(
            durable_approval_id
            or existing.get("status") in {"queued", "claimed"}
            or current_generic_command == str(existing.get("id") or "")
            or (
                isinstance(existing_result, Mapping)
                and bool(existing_result.get("superseded"))
            )
        )
        if belongs_to_current_wait:
            _assert_same_approval_decision(existing, payload)
            if (
                existing.get("status") == "completed"
                and result.get("pending_action") != "policy_approval"
            ):
                # The decision transaction committed but its in-process
                # continuation may have been interrupted.  Re-dispatching is
                # safe because the durable command proof is idempotent.
                _schedule_approval_continuation(
                    task_id, payload, str(existing["id"])
                )
            return {"ok": True, "duplicate": True, "command": existing}
    approval_event = db.query_one(
        "SELECT id FROM task_events WHERE task_id = ? AND type = 'approval_required' "
        "ORDER BY id DESC LIMIT 1",
        (task_id,),
    )
    approval_scope = (
        f"approval:{durable_approval_id}"
        if durable_approval_id
        else f"event:{approval_event['id']}"
        if approval_event
        else "state:"
        + str(task.get("updated_at") or "")
        + ":"
        + json.dumps(result, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    )
    decision_payload = {
        "approved": payload.approved,
        "note": payload.note,
    }
    request_token = uuid.uuid4().hex
    command_id = (
        "tcmd_approval_"
        + hashlib.sha256(
            f"{task_id}\x1f{active['id']}\x1f{approval_scope}".encode("utf-8")
        ).hexdigest()[:24]
    )
    persisted_payload = {
        **decision_payload,
        "decision_request_id": request_token,
    }
    if durable_approval_id:
        persisted_payload["approval_id"] = durable_approval_id
    try:
        command = task_state.enqueue_command(
            task_id,
            "approval",
            run_id=active["id"],
            payload=persisted_payload,
            priority=90,
            command_id=command_id,
            deduplicate=True,
        )
    except sqlite3.IntegrityError:
        # A completed decision is outside enqueue_command's active-command
        # deduplication window, but its stable primary key remains the durable
        # single-decision fence.
        command = task_state.get_command(command_id)
        if command is None:
            raise
    except RunIntakeClosed as exc:
        raise HTTPException(
            status_code=409,
            detail="当前审批运行已经结束，未接受新的审批决定。",
        ) from exc
    except PublicationConflict as exc:
        raise HTTPException(status_code=409, detail=str(exc)) from exc

    _assert_same_approval_decision(command, payload)
    stored_payload = command.get("payload") or {}

    first_decision = stored_payload.get("decision_request_id") == request_token
    if not first_decision:
        return {"ok": True, "duplicate": True, "command": command}

    if result.get("pending_action") == "policy_approval":
        return {"ok": True, "duplicate": False, "command": command}

    _schedule_approval_continuation(task_id, payload, str(command["id"]))
    return {"ok": True, "duplicate": False, "command": command}


def _artifact_row_to_public(row: dict[str, Any]) -> dict[str, Any]:
    metadata = db.json_loads(row.get("metadata_json"), {})
    return _public_artifact({**row, "metadata": metadata})


def _artifact_path(row: dict[str, Any]) -> Path:
    relative = str(row.get("relative_path") or "")
    try:
        if relative:
            return resolve_artifact_path(relative)
        legacy = Path(str(row.get("path") or "")).resolve(strict=True)
        root = ARTIFACT_DIR.resolve(strict=True)
        derived = legacy.relative_to(root).as_posix()
        return resolve_artifact_path(derived)
    except (FileNotFoundError, OSError, RuntimeError, ToolError, ValueError) as exc:
        raise HTTPException(status_code=404, detail="产物文件不存在或不在受控目录中") from exc


def _artifact_or_404(artifact_id: str) -> dict[str, Any]:
    artifact = db.query_one("SELECT * FROM artifacts WHERE id = ?", (artifact_id,))
    if not artifact:
        raise HTTPException(status_code=404, detail="Artifact not found")
    if auth_service.current_identity.get():
        _task_or_404(artifact["task_id"])
    if str(artifact.get("delivery_status") or "") != "published":
        raise HTTPException(
            status_code=409,
            detail="产物尚未通过最终验收，暂不可访问。",
        )
    return artifact


@app.get("/api/artifacts")
def list_artifacts(
    task_id: str = "",
    run_id: str = "",
    workspace_id: str = "",
    kind: str = "",
    limit: int = 200,
) -> list[dict[str, Any]]:
    limit = max(1, min(int(limit), 1000))
    clauses: list[str] = ["delivery_status = 'published'"]
    params: list[Any] = []
    identity = auth_service.current_identity.get()
    if identity and identity["role"] != "admin":
        clauses.append("task_id IN (SELECT t.id FROM tasks t JOIN workspaces w ON w.id=t.workspace LEFT JOIN workspace_members m ON m.workspace_id=w.id AND m.user_id=? WHERE t.user_id=? AND t.organization_id='local-org' AND w.organization_id='local-org' AND (w.owner_user_id=? OR (w.enabled=1 AND m.user_id IS NOT NULL)))")
        params.extend([identity["user_id"]] * 3)
    for column, value in (
        ("task_id", task_id),
        ("run_id", run_id),
        ("workspace_id", workspace_id),
        ("kind", kind),
    ):
        if value:
            clauses.append(f"{column} = ?")
            params.append(value)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    rows = db.query_all(
        f"SELECT * FROM artifacts{where} ORDER BY created_at DESC LIMIT ?",  # noqa: S608 - fixed columns
        (*params, limit),
    )
    return [_artifact_row_to_public(row) for row in rows]


@app.get("/api/tasks/{task_id}/artifacts")
def list_task_artifacts(task_id: str, run_id: str = "") -> list[dict[str, Any]]:
    _task_or_404(task_id)
    return list_artifacts(task_id=task_id, run_id=run_id)


@app.get("/api/artifacts/{artifact_id}")
def get_artifact(artifact_id: str) -> dict[str, Any]:
    return _artifact_row_to_public(_artifact_or_404(artifact_id))


def _sanitise_html_preview(value: str) -> str:
    cleaned = re.sub(
        r"<\s*(script|iframe|object|embed|base|form|link|meta)[^>]*>[\s\S]*?<\s*/\s*\1\s*>",
        "",
        value,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(
        r"<\s*(script|iframe|object|embed|base|form|link|meta)\b[^>]*?/?>",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    cleaned = re.sub(r"\s+on[a-z]+\s*=\s*(['\"]).*?\1", "", cleaned, flags=re.IGNORECASE | re.DOTALL)
    cleaned = re.sub(
        r"\s+(href|src)\s*=\s*(['\"])(?!#|data:)[\s\S]*?\2",
        "",
        cleaned,
        flags=re.IGNORECASE,
    )
    csp = "<meta http-equiv=\"Content-Security-Policy\" content=\"default-src 'none'; img-src data:; style-src 'unsafe-inline'; font-src data:\">"
    return csp + cleaned


@app.get("/api/artifacts/{artifact_id}/preview")
def preview_artifact(artifact_id: str) -> dict[str, Any]:
    artifact = _artifact_or_404(artifact_id)
    path = _artifact_path(artifact)
    kind = str(artifact.get("kind") or path.suffix.lstrip(".")).lower()
    public = _artifact_row_to_public(artifact)
    if path.stat().st_size > 25 * 1024 * 1024:
        return {"artifact": public, "preview_kind": "unavailable", "message": "文件超过 25MB，请下载后查看。"}
    try:
        if kind in {"markdown", "md"}:
            return {"artifact": public, "preview_kind": "markdown", "content": path.read_text(encoding="utf-8", errors="replace")[:500_000]}
        if kind == "html":
            raw = path.read_text(encoding="utf-8", errors="replace")[:500_000]
            return {"artifact": public, "preview_kind": "html", "content": _sanitise_html_preview(raw), "sandbox": ""}
        if kind in {"csv", "tsv"}:
            delimiter = "\t" if kind == "tsv" else ","
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as stream:
                rows = [row[:30] for _, row in zip(range(101), csv.reader(stream, delimiter=delimiter))]
            return {"artifact": public, "preview_kind": "spreadsheet", "sheets": [{"name": path.stem, "rows": rows}]}
        if kind == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for sheet in workbook.worksheets[:10]:
                rows = [
                    ["" if cell is None else str(cell) for cell in row[:30]]
                    for _, row in zip(range(101), sheet.iter_rows(values_only=True))
                ]
                sheets.append({"name": sheet.title, "rows": rows})
            workbook.close()
            return {"artifact": public, "preview_kind": "spreadsheet", "sheets": sheets}
        if kind == "docx":
            from docx import Document

            document = Document(path)
            paragraphs = [item.text for item in document.paragraphs if item.text.strip()][:500]
            tables = [
                [[cell.text for cell in row.cells] for row in table.rows[:100]]
                for table in document.tables[:20]
            ]
            return {"artifact": public, "preview_kind": "document", "paragraphs": paragraphs, "tables": tables}
        if kind == "pptx":
            from pptx import Presentation

            presentation = Presentation(path)
            slides = []
            for index, slide in enumerate(presentation.slides, start=1):
                if index > 100:
                    break
                texts = [
                    str(shape.text).strip()
                    for shape in slide.shapes
                    if hasattr(shape, "text") and str(shape.text).strip()
                ]
                slides.append({"number": index, "title": texts[0] if texts else f"第 {index} 页", "texts": texts})
            return {"artifact": public, "preview_kind": "slides", "slides": slides}
        if kind == "pdf":
            return {"artifact": public, "preview_kind": "pdf", "url": public["download_url"] + "?inline=true"}
        if kind in {"txt", "text", "json", "yaml", "yml"}:
            return {"artifact": public, "preview_kind": "text", "content": path.read_text(encoding="utf-8", errors="replace")[:500_000]}
    except Exception:
        return {
            "artifact": public,
            "preview_kind": "error",
            "message": "平台暂时无法生成该文件的预览，请下载后使用对应应用打开。",
        }
    return {"artifact": public, "preview_kind": "unavailable", "message": "该格式暂不支持平台内预览，可下载原文件。"}


@app.get("/api/artifacts/{artifact_id}/download")
def download_artifact(artifact_id: str, inline: bool = False):
    artifact = _artifact_or_404(artifact_id)
    path = _artifact_path(artifact)
    media_type = str(artifact.get("mime_type") or "application/octet-stream")
    if inline:
        return FileResponse(path, media_type=media_type)
    return FileResponse(path, filename=artifact["name"], media_type=media_type)


app.mount("/", StaticFiles(directory=str(WEB_DIR), html=True), name="web")
