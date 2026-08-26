from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Iterator, Mapping

from app import db
from app.services.context_service import ExecutionScope


KNOWLEDGE_SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS knowledge_bases (
    id TEXT PRIMARY KEY,
    organization_id TEXT NOT NULL DEFAULT 'local-org',
    workspace_id TEXT NOT NULL DEFAULT 'default',
    owner_user_id TEXT NOT NULL DEFAULT 'local-user',
    name TEXT NOT NULL,
    description TEXT NOT NULL DEFAULT '',
    visibility TEXT NOT NULL DEFAULT 'workspace',
    enabled INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS knowledge_documents (
    id TEXT PRIMARY KEY,
    knowledge_base_id TEXT NOT NULL,
    upload_id TEXT NOT NULL DEFAULT '',
    name TEXT NOT NULL,
    mime_type TEXT NOT NULL DEFAULT 'text/plain',
    source_type TEXT NOT NULL DEFAULT 'upload',
    source_ref TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'indexed',
    error TEXT NOT NULL DEFAULT '',
    content_sha256 TEXT NOT NULL DEFAULT '',
    chunk_count INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS knowledge_chunks (
    id TEXT PRIMARY KEY,
    knowledge_base_id TEXT NOT NULL,
    document_id TEXT NOT NULL,
    ordinal INTEGER NOT NULL,
    content TEXT NOT NULL,
    content_sha256 TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    UNIQUE(document_id, ordinal)
);

CREATE INDEX IF NOT EXISTS idx_knowledge_bases_scope
    ON knowledge_bases(organization_id, workspace_id, owner_user_id, visibility, enabled);
CREATE INDEX IF NOT EXISTS idx_knowledge_documents_base
    ON knowledge_documents(knowledge_base_id, created_at DESC);
CREATE INDEX IF NOT EXISTS idx_knowledge_chunks_base
    ON knowledge_chunks(knowledge_base_id, document_id, ordinal);
"""


_WORD_RE = re.compile(r"[\w\u4e00-\u9fff]+", re.UNICODE)
_WHITESPACE_RE = re.compile(r"\s+")


class KnowledgeBaseError(RuntimeError):
    pass


class KnowledgeBaseNotFoundError(KnowledgeBaseError):
    pass


class KnowledgeDocumentError(KnowledgeBaseError):
    pass


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean_text(value: Any, *, field: str, allow_empty: bool = True, max_length: int = 10_000) -> str:
    if not isinstance(value, str):
        raise TypeError(f"{field} must be a string")
    result = value.strip()
    if not allow_empty and not result:
        raise ValueError(f"{field} cannot be empty")
    if len(result) > max_length:
        raise ValueError(f"{field} is too long")
    if any(ord(character) < 32 and character not in "\n\r\t" for character in result):
        raise ValueError(f"{field} cannot contain control characters")
    return result


def _clean_id(value: Any, *, field: str, allow_empty: bool = False) -> str:
    result = _clean_text(value, field=field, allow_empty=allow_empty, max_length=160)
    if not result and allow_empty:
        return ""
    if not re.fullmatch(r"[A-Za-z0-9_-]+", result):
        raise ValueError(f"{field} must contain only letters, digits, '_' or '-'")
    return result


def _tokenize(value: str) -> set[str]:
    return {item.casefold() for item in _WORD_RE.findall(value) if len(item.strip()) >= 2}


def _normalise_content(value: str) -> str:
    return _WHITESPACE_RE.sub(" ", value).strip()


def _chunk_text(value: str, *, target_chars: int = 1200, overlap_chars: int = 160) -> list[str]:
    text = _normalise_content(value)
    if not text:
        return []
    chunks: list[str] = []
    start = 0
    length = len(text)
    while start < length and len(chunks) < 500:
        end = min(length, start + target_chars)
        if end < length:
            split_at = max(text.rfind("。", start, end), text.rfind(".", start, end), text.rfind("\n", start, end))
            if split_at > start + 400:
                end = split_at + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= length:
            break
        start = max(end - overlap_chars, start + 1)
    return chunks


def _extract_upload_text(upload: Mapping[str, Any], *, max_bytes: int = 2_000_000) -> str:
    path = Path(str(upload.get("path") or ""))
    if not path.exists() or not path.is_file():
        raise KnowledgeDocumentError("上传文件不存在，无法建立知识库索引")
    if path.stat().st_size > max_bytes:
        raise KnowledgeDocumentError("知识库单文件索引上限为 2MB")
    suffix = path.suffix.lower()
    content_type = str(upload.get("content_type") or "").lower()
    if suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml", ".html", ".htm", ".log"} or content_type.startswith("text/"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
        table_cells = [
            cell.text.strip()
            for table in document.tables[:50]
            for row in table.rows[:200]
            for cell in row.cells[:50]
            if cell.text.strip()
        ]
        return "\n".join([*paragraphs, *table_cells])
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            lines: list[str] = []
            for sheet in workbook.worksheets[:10]:
                lines.append(f"工作表：{sheet.title}")
                for row in sheet.iter_rows(max_row=500, max_col=50, values_only=True):
                    values = ["" if cell is None else str(cell) for cell in row]
                    if any(item.strip() for item in values):
                        lines.append("\t".join(values))
            return "\n".join(lines)
        finally:
            workbook.close()
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        lines = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines.append(f"第 {index} 页")
            for shape in slide.shapes:
                if hasattr(shape, "text") and str(shape.text).strip():
                    lines.append(str(shape.text).strip())
        return "\n".join(lines)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except Exception as exc:  # pragma: no cover - optional dependency
            raise KnowledgeDocumentError("PDF 解析组件未安装，暂不能索引该文件") from exc
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages[:80])
    raise KnowledgeDocumentError("该文件格式暂不支持知识库索引")


class KnowledgeBaseService:
    def __init__(
        self,
        connection: sqlite3.Connection | Callable[[], sqlite3.Connection] | None = None,
        *,
        clock: Callable[[], str] | None = None,
        auto_init: bool = True,
    ) -> None:
        self._shared_connection = connection if isinstance(connection, sqlite3.Connection) else None
        self._connection_factory = connection if callable(connection) else db.get_conn
        self._clock = clock or _utc_now
        self._lock = threading.RLock()
        if auto_init:
            self.init_schema()

    def init_schema(self) -> None:
        with self._connection(write=True) as conn:
            db._execute_schema_script(conn, KNOWLEDGE_SCHEMA_SQL)

    def _connection(self, *, write: bool = False) -> Iterator[sqlite3.Connection]:
        class _Context:
            def __init__(inner_self, outer: KnowledgeBaseService) -> None:
                inner_self.outer = outer
                inner_self.conn: sqlite3.Connection | None = None
                inner_self.owns = False
                inner_self.original_row_factory: Any = None

            def __enter__(inner_self) -> sqlite3.Connection:
                outer = inner_self.outer
                outer._lock.acquire()
                conn = outer._shared_connection or outer._connection_factory()
                inner_self.conn = conn
                inner_self.owns = outer._shared_connection is None
                inner_self.original_row_factory = conn.row_factory
                conn.row_factory = sqlite3.Row
                if write:
                    conn.execute("BEGIN IMMEDIATE")
                return conn

            def __exit__(inner_self, exc_type: Any, exc: Any, tb: Any) -> None:
                assert inner_self.conn is not None
                try:
                    if write:
                        if exc_type is None:
                            inner_self.conn.commit()
                        else:
                            inner_self.conn.rollback()
                finally:
                    if inner_self.owns:
                        inner_self.conn.close()
                    else:
                        inner_self.conn.row_factory = inner_self.original_row_factory
                    inner_self.outer._lock.release()

        return _Context(self)

    @staticmethod
    def _scope(value: ExecutionScope | Mapping[str, Any] | None) -> ExecutionScope:
        return ExecutionScope.normalise(value)

    @staticmethod
    def _base_public(row: Mapping[str, Any]) -> dict[str, Any]:
        return {
            "id": row["id"],
            "organization_id": row["organization_id"],
            "workspace_id": row["workspace_id"],
            "owner_user_id": row["owner_user_id"],
            "name": row["name"],
            "description": row["description"],
            "visibility": row["visibility"],
            "enabled": bool(row["enabled"]),
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
        }

    @staticmethod
    def _document_public(row: Mapping[str, Any]) -> dict[str, Any]:
        return {
            "id": row["id"],
            "knowledge_base_id": row["knowledge_base_id"],
            "upload_id": row["upload_id"],
            "name": row["name"],
            "mime_type": row["mime_type"],
            "source_type": row["source_type"],
            "source_ref": row["source_ref"],
            "status": row["status"],
            "error": row["error"],
            "content_sha256": row["content_sha256"],
            "chunk_count": int(row["chunk_count"] or 0),
            "created_at": row["created_at"],
            "updated_at": row["updated_at"],
        }

    @staticmethod
    def _base_visible(row: Mapping[str, Any], scope: ExecutionScope) -> bool:
        if str(row["organization_id"]) != scope.organization_id:
            return False
        visibility = str(row["visibility"] or "workspace")
        if visibility == "organization":
            return True
        if str(row["workspace_id"]) != scope.workspace_id:
            return False
        if visibility == "workspace":
            return True
        return str(row["owner_user_id"]) == scope.user_id

    def create_base(
        self,
        execution_scope: ExecutionScope | Mapping[str, Any] | None,
        *,
        name: str,
        description: str = "",
        visibility: str = "workspace",
        enabled: bool = True,
        base_id: str | None = None,
    ) -> dict[str, Any]:
        scope = self._scope(execution_scope)
        base_id = _clean_id(base_id or "kb_" + uuid.uuid4().hex[:12], field="base_id")
        name = _clean_text(name, field="name", allow_empty=False, max_length=120)
        description = _clean_text(description, field="description", max_length=2000)
        if visibility not in {"private", "workspace", "organization"}:
            raise ValueError("visibility must be private, workspace or organization")
        if not isinstance(enabled, bool):
            raise TypeError("enabled must be a boolean")
        now = self._clock()
        with self._connection(write=True) as conn:
            conn.execute(
                """
                INSERT INTO knowledge_bases(
                    id, organization_id, workspace_id, owner_user_id, name,
                    description, visibility, enabled, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    base_id,
                    scope.organization_id,
                    scope.workspace_id,
                    scope.user_id,
                    name,
                    description,
                    visibility,
                    int(enabled),
                    now,
                    now,
                ),
            )
            row = conn.execute("SELECT * FROM knowledge_bases WHERE id = ?", (base_id,)).fetchone()
        return self._base_public(dict(row or {}))

    def list_bases(self, execution_scope: ExecutionScope | Mapping[str, Any] | None, *, include_disabled: bool = True) -> list[dict[str, Any]]:
        scope = self._scope(execution_scope)
        with self._connection() as conn:
            rows = conn.execute(
                "SELECT * FROM knowledge_bases WHERE organization_id = ? ORDER BY updated_at DESC, id",
                (scope.organization_id,),
            ).fetchall()
        bases = [self._base_public(dict(row)) for row in rows if self._base_visible(dict(row), scope)]
        if not include_disabled:
            bases = [item for item in bases if item["enabled"]]
        return bases

    def get_base(self, base_id: str, execution_scope: ExecutionScope | Mapping[str, Any] | None) -> dict[str, Any] | None:
        base_id = _clean_id(base_id, field="base_id")
        scope = self._scope(execution_scope)
        with self._connection() as conn:
            row = conn.execute("SELECT * FROM knowledge_bases WHERE id = ?", (base_id,)).fetchone()
        if row is None:
            return None
        raw = dict(row)
        return self._base_public(raw) if self._base_visible(raw, scope) else None

    def update_base(
        self,
        base_id: str,
        execution_scope: ExecutionScope | Mapping[str, Any] | None,
        *,
        name: str,
        description: str = "",
        visibility: str = "workspace",
        enabled: bool = True,
    ) -> dict[str, Any]:
        base_id = _clean_id(base_id, field="base_id")
        scope = self._scope(execution_scope)
        name = _clean_text(name, field="name", allow_empty=False, max_length=120)
        description = _clean_text(description, field="description", max_length=2000)
        if visibility not in {"private", "workspace", "organization"}:
            raise ValueError("visibility must be private, workspace or organization")
        if not isinstance(enabled, bool):
            raise TypeError("enabled must be a boolean")
        now = self._clock()
        with self._connection(write=True) as conn:
            self._require_base(conn, base_id, scope)
            conn.execute(
                """
                UPDATE knowledge_bases
                SET name = ?, description = ?, visibility = ?, enabled = ?, updated_at = ?
                WHERE id = ?
                """,
                (name, description, visibility, int(enabled), now, base_id),
            )
            row = conn.execute("SELECT * FROM knowledge_bases WHERE id = ?", (base_id,)).fetchone()
        return self._base_public(dict(row or {}))

    def delete_base(
        self,
        base_id: str,
        execution_scope: ExecutionScope | Mapping[str, Any] | None,
    ) -> dict[str, Any]:
        base_id = _clean_id(base_id, field="base_id")
        scope = self._scope(execution_scope)
        with self._connection(write=True) as conn:
            self._require_base(conn, base_id, scope)
            conn.execute("DELETE FROM knowledge_chunks WHERE knowledge_base_id = ?", (base_id,))
            conn.execute("DELETE FROM knowledge_documents WHERE knowledge_base_id = ?", (base_id,))
            conn.execute("DELETE FROM knowledge_bases WHERE id = ?", (base_id,))
        return {"id": base_id, "deleted": True}

    def _require_base(self, conn: sqlite3.Connection, base_id: str, scope: ExecutionScope) -> dict[str, Any]:
        row = conn.execute("SELECT * FROM knowledge_bases WHERE id = ?", (base_id,)).fetchone()
        if row is None or not self._base_visible(dict(row), scope):
            raise KnowledgeBaseNotFoundError("知识库不存在或当前作用域不可见")
        return dict(row)

    def index_upload(
        self,
        base_id: str,
        execution_scope: ExecutionScope | Mapping[str, Any] | None,
        *,
        upload: Mapping[str, Any],
        document_id: str | None = None,
    ) -> dict[str, Any]:
        base_id = _clean_id(base_id, field="base_id")
        scope = self._scope(execution_scope)
        document_id = _clean_id(document_id or "kdoc_" + uuid.uuid4().hex[:12], field="document_id")
        upload_id = _clean_text(str(upload.get("id") or ""), field="upload_id", max_length=160)
        name = _clean_text(str(upload.get("name") or "uploaded"), field="name", allow_empty=False, max_length=300)
        mime_type = _clean_text(str(upload.get("content_type") or "application/octet-stream"), field="mime_type", max_length=200)
        raw_text = _extract_upload_text(upload)
        chunks = _chunk_text(raw_text)
        if not chunks:
            raise KnowledgeDocumentError("文件没有可索引的文本内容")
        digest = hashlib.sha256(raw_text.encode("utf-8", errors="replace")).hexdigest()
        now = self._clock()
        with self._connection(write=True) as conn:
            self._require_base(conn, base_id, scope)
            conn.execute(
                """
                INSERT INTO knowledge_documents(
                    id, knowledge_base_id, upload_id, name, mime_type,
                    source_type, source_ref, status, error, content_sha256,
                    chunk_count, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, 'upload', ?, 'indexed', '', ?, ?, ?, ?)
                """,
                (
                    document_id,
                    base_id,
                    upload_id,
                    name,
                    mime_type,
                    upload_id,
                    digest,
                    len(chunks),
                    now,
                    now,
                ),
            )
            for ordinal, chunk in enumerate(chunks):
                chunk_id = f"kchunk_{uuid.uuid4().hex[:12]}"
                conn.execute(
                    """
                    INSERT INTO knowledge_chunks(
                        id, knowledge_base_id, document_id, ordinal, content,
                        content_sha256, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        chunk_id,
                        base_id,
                        document_id,
                        ordinal,
                        chunk,
                        hashlib.sha256(chunk.encode("utf-8")).hexdigest(),
                        now,
                    ),
                )
            conn.execute(
                "UPDATE knowledge_bases SET updated_at = ? WHERE id = ?",
                (now, base_id),
            )
            row = conn.execute(
                "SELECT * FROM knowledge_documents WHERE id = ?", (document_id,)
            ).fetchone()
        return self._document_public(dict(row or {}))

    def list_documents(self, base_id: str, execution_scope: ExecutionScope | Mapping[str, Any] | None) -> list[dict[str, Any]]:
        base_id = _clean_id(base_id, field="base_id")
        scope = self._scope(execution_scope)
        with self._connection() as conn:
            self._require_base(conn, base_id, scope)
            rows = conn.execute(
                "SELECT * FROM knowledge_documents WHERE knowledge_base_id = ? ORDER BY created_at DESC",
                (base_id,),
            ).fetchall()
        return [self._document_public(dict(row)) for row in rows]

    def search(
        self,
        execution_scope: ExecutionScope | Mapping[str, Any] | None,
        *,
        query: str,
        base_id: str = "",
        limit: int = 5,
    ) -> dict[str, Any]:
        scope = self._scope(execution_scope)
        query = _clean_text(query, field="query", allow_empty=False, max_length=4000)
        if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 20:
            raise ValueError("limit must be an integer between 1 and 20")
        query_tokens = _tokenize(query)
        if not query_tokens:
            return {"query": query, "matches": [], "used_knowledge_base_ids": []}
        selected_base_id = _clean_id(base_id, field="base_id", allow_empty=True) if base_id else ""
        with self._connection() as conn:
            bases = [
                dict(row)
                for row in conn.execute(
                    "SELECT * FROM knowledge_bases WHERE organization_id = ? AND enabled = 1",
                    (scope.organization_id,),
                ).fetchall()
            ]
            visible_base_ids = {
                row["id"]
                for row in bases
                if self._base_visible(row, scope) and (not selected_base_id or row["id"] == selected_base_id)
            }
            if selected_base_id and selected_base_id not in visible_base_ids:
                raise KnowledgeBaseNotFoundError("知识库不存在或当前作用域不可见")
            if not visible_base_ids:
                return {"query": query, "matches": [], "used_knowledge_base_ids": []}
            placeholders = ",".join("?" for _ in visible_base_ids)
            rows = conn.execute(
                f"""
                SELECT c.*, d.name AS document_name, d.source_type, d.source_ref
                FROM knowledge_chunks AS c
                JOIN knowledge_documents AS d ON d.id = c.document_id
                WHERE c.knowledge_base_id IN ({placeholders})
                  AND d.status = 'indexed'
                """,
                tuple(sorted(visible_base_ids)),
            ).fetchall()
        scored: list[dict[str, Any]] = []
        for row in rows:
            raw = dict(row)
            content = str(raw["content"])
            tokens = _tokenize(content)
            overlap = query_tokens & tokens
            if not overlap:
                continue
            score = len(overlap) * 10 + sum(content.casefold().count(token) for token in overlap)
            scored.append(
                {
                    "chunk_id": raw["id"],
                    "knowledge_base_id": raw["knowledge_base_id"],
                    "document_id": raw["document_id"],
                    "document_name": raw["document_name"],
                    "ordinal": int(raw["ordinal"]),
                    "content": content,
                    "score": score,
                    "matched_terms": sorted(overlap)[:20],
                    "source": {
                        "type": raw.get("source_type") or "upload",
                        "ref": raw.get("source_ref") or "",
                    },
                }
            )
        scored.sort(key=lambda item: (-int(item["score"]), item["document_name"], int(item["ordinal"])))
        matches = scored[:limit]
        return {
            "query": query,
            "matches": matches,
            "used_knowledge_base_ids": sorted({item["knowledge_base_id"] for item in matches}),
        }

    @staticmethod
    def format_context(search_result: Mapping[str, Any], *, max_chars: int = 6000) -> str:
        matches = search_result.get("matches") if isinstance(search_result, Mapping) else []
        if not isinstance(matches, list) or not matches:
            return ""
        lines: list[str] = []
        used = 0
        for index, match in enumerate(matches, start=1):
            if not isinstance(match, Mapping):
                continue
            content = _normalise_content(str(match.get("content") or ""))
            prefix = (
                f"[{index}] 文档 {match.get('document_name')} "
                f"片段 {int(match.get('ordinal') or 0) + 1} "
                f"(chunk_id={match.get('chunk_id')})"
            )
            item = f"{prefix}\n{content}"
            if used + len(item) > max_chars:
                break
            lines.append(item)
            used += len(item)
        return "\n\n".join(lines)
