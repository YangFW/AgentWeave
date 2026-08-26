from __future__ import annotations

import builtins
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from app.services.agent_runtime import AgentRuntime


class AttachmentContextTests(unittest.TestCase):
    def setUp(self) -> None:
        self.runtime = object.__new__(AgentRuntime)
        self.temp_directory = TemporaryDirectory()
        self.directory = Path(self.temp_directory.name)

    def tearDown(self) -> None:
        self.temp_directory.cleanup()

    @staticmethod
    def _attachment(path: Path, content_type: str) -> dict[str, object]:
        return {
            "name": path.name,
            "path": str(path),
            "content_type": content_type,
            "size": path.stat().st_size,
        }

    def test_docx_extracts_paragraphs_and_tables_without_following_links(self) -> None:
        path = self.directory / "brief.docx"
        document = Document()
        document.add_heading("项目简报", level=1)
        document.add_paragraph("这是 Word 正文。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "风险"
        table.cell(0, 1).text = "负责人"
        table.cell(1, 0).text = "进度延期"
        table.cell(1, 1).text = "张三"
        document.save(path)

        context = self.runtime._attachment_context(
            [self._attachment(path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")]
        )

        self.assertIn("--- brief.docx ---", context)
        self.assertIn("项目简报", context)
        self.assertIn("这是 Word 正文。", context)
        self.assertIn("风险 | 负责人", context)
        self.assertIn("进度延期 | 张三", context)

    def test_xlsx_extracts_cells_marks_formulas_and_limits_worksheets(self) -> None:
        path = self.directory / "numbers.xlsx"
        workbook = Workbook()
        first = workbook.active
        first.title = "汇总"
        first.append(["项目", "金额", "合计"])
        first.append(["A", 12, "=SUM(B2:B2)"])
        for index in range(1, self.runtime.ATTACHMENT_MAX_WORKSHEETS + 2):
            sheet = workbook.create_sheet(f"附表{index}")
            sheet["A1"] = f"内容{index}"
        workbook.save(path)

        context = self.runtime._attachment_context(
            [self._attachment(path, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")]
        )

        self.assertIn("[工作表：汇总]", context)
        self.assertIn("项目 | 金额 | 合计", context)
        self.assertIn("A | 12 | [公式未执行]", context)
        self.assertNotIn("SUM(B2:B2)", context)
        self.assertIn(
            f"[工作簿仅提取前 {self.runtime.ATTACHMENT_MAX_WORKSHEETS} 个工作表]",
            context,
        )
        self.assertNotIn(f"[工作表：附表{self.runtime.ATTACHMENT_MAX_WORKSHEETS}]", context)

    def test_pptx_extracts_slide_text_and_tables(self) -> None:
        path = self.directory / "review.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "平台评审"
        slide.placeholders[1].text = "第一项建议"
        table_shape = slide.shapes.add_table(2, 2, 0, 0, 4_000_000, 1_000_000)
        table_shape.table.cell(0, 0).text = "等级"
        table_shape.table.cell(0, 1).text = "事项"
        table_shape.table.cell(1, 0).text = "P0"
        table_shape.table.cell(1, 1).text = "安全整改"
        presentation.save(path)

        context = self.runtime._attachment_context(
            [self._attachment(path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")]
        )

        self.assertIn("[幻灯片 1]", context)
        self.assertIn("平台评审", context)
        self.assertIn("第一项建议", context)
        self.assertIn("等级 | 事项", context)
        self.assertIn("P0 | 安全整改", context)

    def test_pptx_slide_count_is_bounded(self) -> None:
        path = self.directory / "many-slides.pptx"
        presentation = Presentation()
        for index in range(4):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = f"标题 {index + 1}"
        presentation.save(path)

        with patch.object(self.runtime, "ATTACHMENT_MAX_SLIDES", 2):
            context = self.runtime._attachment_context(
                [self._attachment(path, "application/vnd.openxmlformats-officedocument.presentationml.presentation")]
            )

        self.assertIn("标题 1", context)
        self.assertIn("标题 2", context)
        self.assertNotIn("标题 3", context)
        self.assertIn("[演示文稿仅提取前 2 张幻灯片]", context)

    def test_pdf_extracts_text_when_pypdf_is_available(self) -> None:
        try:
            import pypdf  # noqa: F401
        except ImportError:
            self.skipTest("pypdf is declared in requirements but is not installed in this environment")

        path = self.directory / "notice.pdf"
        document = canvas.Canvas(str(path))
        document.drawString(72, 760, "PDF attachment extraction works")
        document.save()

        context = self.runtime._attachment_context(
            [self._attachment(path, "application/pdf")]
        )

        self.assertIn("[第 1 页]", context)
        self.assertIn("PDF attachment extraction works", context)

    def test_pdf_page_count_is_bounded(self) -> None:
        try:
            import pypdf  # noqa: F401
        except ImportError:
            self.skipTest("pypdf is declared in requirements but is not installed in this environment")

        path = self.directory / "many-pages.pdf"
        document = canvas.Canvas(str(path))
        for index in range(4):
            document.drawString(72, 760, f"PDF page {index + 1}")
            document.showPage()
        document.save()

        with patch.object(self.runtime, "ATTACHMENT_MAX_PDF_PAGES", 2):
            context = self.runtime._attachment_context(
                [self._attachment(path, "application/pdf")]
            )

        self.assertIn("PDF page 1", context)
        self.assertIn("PDF page 2", context)
        self.assertNotIn("PDF page 3", context)
        self.assertIn("[PDF 仅提取前 2 页]", context)

    def test_pdf_missing_optional_parser_returns_friendly_message(self) -> None:
        path = self.directory / "missing-parser.pdf"
        path.write_bytes(b"%PDF-1.4\n")
        real_import = builtins.__import__

        def import_without_pypdf(name: str, *args: object, **kwargs: object):
            if name == "pypdf":
                raise ImportError("test-missing-pypdf")
            return real_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=import_without_pypdf):
            context = self.runtime._attachment_context(
                [self._attachment(path, "application/pdf")]
            )

        self.assertIn("PDF 正文解析组件未安装", context)
        self.assertNotIn("test-missing-pypdf", context)

    def test_corrupt_office_file_and_missing_file_are_friendly(self) -> None:
        corrupt = self.directory / "corrupt.docx"
        corrupt.write_bytes(b"not-a-zip")
        missing = self.directory / "missing.xlsx"

        context = self.runtime._attachment_context(
            [
                self._attachment(corrupt, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                {"name": missing.name, "path": str(missing), "content_type": "application/octet-stream"},
            ]
        )

        self.assertIn("Word 正文提取失败", context)
        self.assertIn("文件可能已损坏", context)
        self.assertIn("附件文件不存在或不可访问", context)
        self.assertNotIn("Traceback", context)

    def test_total_context_and_attachment_count_are_bounded(self) -> None:
        attachments: list[dict[str, object]] = []
        for index in range(self.runtime.ATTACHMENT_MAX_FILES + 2):
            path = self.directory / f"attachment-{index}.txt"
            path.write_text("x" * self.runtime.ATTACHMENT_MAX_FILE_CHARS, encoding="utf-8")
            attachments.append(self._attachment(path, "text/plain"))

        context = self.runtime._attachment_context(attachments)

        self.assertLessEqual(len(context), self.runtime.ATTACHMENT_MAX_CONTEXT_CHARS)
        self.assertIn("附件正文已达到总字符上限", context)
        self.assertNotIn(f"--- attachment-{self.runtime.ATTACHMENT_MAX_FILES}.txt ---", context)

    def test_oversized_file_is_not_opened(self) -> None:
        path = self.directory / "oversized.txt"
        path.write_text("small fixture", encoding="utf-8")
        with patch.object(self.runtime, "ATTACHMENT_MAX_FILE_BYTES", 1):
            context = self.runtime._attachment_context(
                [{"name": path.name, "path": str(path), "content_type": "text/plain"}]
            )

        self.assertIn("附件超过正文解析大小上限", context)


if __name__ == "__main__":
    unittest.main()
