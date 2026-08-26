from __future__ import annotations

import asyncio
import contextvars
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from reportlab.pdfgen import canvas

from app.services.agent_runtime import AgentRuntime, sanitize_public_answer
from app.services.goal_spec_service import (
    AcceptanceCriterion,
    ConfirmationSpec,
    DeliverableSpec,
    GoalSpec,
    InputSpec,
    ObjectiveSpec,
    ProvenanceRef,
)
from app.services.mcp_gateway import ToolError


class NoopTaskState:
    def raise_if_cancel_requested(self, *_: object, **__: object) -> None:
        return None


class ExecutionGuardTests(unittest.TestCase):
    def setUp(self) -> None:
        self.runtime = object.__new__(AgentRuntime)
        self.runtime._execution_context = contextvars.ContextVar(
            f"execution_guard_test_{id(self)}", default=None
        )
        self.runtime.task_state = NoopTaskState()

    def _contract_plan(
        self, plan: dict[str, object], tools: list[tuple[str, str]]
    ) -> dict[str, object]:
        """Attach the immutable references present on every runtime plan."""

        goal_ref = {
            "id": "goal-spec-test",
            "goal_id": "goal-test",
            "version": 1,
            "spec_hash": "a" * 64,
        }
        bound = {
            **plan,
            "goal_spec_ref": goal_ref,
            "allowed_tools": [
                {"server_id": server_id, "tool_name": tool_name}
                for server_id, tool_name in tools
            ],
        }
        self.runtime._execution_context.set(
            {"task_id": "task_guard", "run_id": "run_guard", "state": {"goal_spec_ref": goal_ref}}
        )
        return bound

    def test_missing_city_returns_one_friendly_question(self) -> None:
        answer = self.runtime._clarification_for_missing({"intent": "weather_query", "missing_information": ["city"]})
        self.assertEqual(answer, "还差一个信息：请告诉我城市或地区，例如“宁波”或“上海浦东”。")
        self.assertNotIn("执行过程", answer)
        self.assertNotIn("任务理解", answer)

    def test_internal_acceptance_code_is_not_public_answer(self) -> None:
        answer = sanitize_public_answer("你好！\n本次验收代号：星云-218\n\n我在，可以继续帮你。")
        self.assertEqual(answer, "你好！\n\n我在，可以继续帮你。")
        self.assertNotIn("星云-218", answer)

    def test_arbitrary_short_required_fields_are_requested(self) -> None:
        answer = self.runtime._clarification_for_missing({
            "intent": "prepare_contract",
            "missing_information": ["合同主体", "签署方名称"],
        })
        self.assertEqual(answer, "还需要你补充：合同主体、签署方名称。补充后我会继续当前任务。")

    def test_missing_information_never_exposes_parser_json_or_reasoning(self) -> None:
        answer = self.runtime._clarification_for_missing({
            "intent": "general",
            "missing_information": [
                '{"field":"合同主体"}',
                "因为没有材料，所以我需要先分析原因。",
                "parameters",
                "合同主体",
            ],
        })
        self.assertEqual(answer, "还差一个信息：请告诉我合同主体。")
        self.assertNotIn("parameters", answer)
        self.assertNotIn("因为", answer)

    def test_document_follow_up_with_weather_text_is_not_weather_intent(self) -> None:
        request = self.runtime._weather_request_for_intent(
            "能给我将这个行程出个文档吗，word版的，让我下载",
            "将杭州行程整理成 Word，包含天气不好时的室内推荐",
            [],
            {"intent": "create_travel_itinerary_docx", "parameters": {"city": "杭州", "format": "docx"}},
        )
        self.assertIsNone(request)

    def test_itinerary_with_bad_weather_alternative_is_not_weather_lookup(self) -> None:
        message = "给我安排杭州一日行程，包含天气不好时的室内备选"
        self.assertFalse(self.runtime._looks_like_weather_lookup(message))
        request = self.runtime._weather_request_for_intent(
            message,
            message,
            [],
            {"intent": "travel_itinerary", "source": "model", "parameters": {"city": "杭州"}},
        )
        self.assertIsNone(request)

    def test_real_weather_question_is_weather_lookup(self) -> None:
        self.assertTrue(self.runtime._looks_like_weather_lookup("杭州今天天气怎么样"))
        self.assertTrue(self.runtime._looks_like_weather_lookup("明天会不会下雨"))

    def test_query_is_not_extracted_as_a_city(self) -> None:
        self.assertEqual(self.runtime._weather_request("查询今天天气怎么样", []), {"city": "", "day": "today"})

    def test_document_plan_allows_report_and_blocks_weather(self) -> None:
        task = {
            "message": "整理成 Word",
            "resolved_message": "把上一轮杭州行程整理成 Word 文档",
            "intent_resolution": {"intent": "create_document", "standalone_request": "把上一轮杭州行程整理成 Word 文档"},
        }
        plan = self.runtime._build_execution_plan(task, [], "docx", False)
        self.assertIn("report", plan["allowed_servers"])
        self.assertNotIn("weather", plan["allowed_servers"])
        plan = self._contract_plan(plan, [("report", "generate_document")])
        with self.assertRaises(ToolError):
            self.runtime._validate_tool_against_plan(plan, "weather", "forecast")
        reason, message = self.runtime._tool_plan_denial(plan, "weather", "forecast")
        self.assertEqual(reason, "irrelevant_weather_tool")
        self.assertIn("生成文档", message)

    def test_document_plan_respects_explicit_no_network_constraint(self) -> None:
        message = "把前一轮结果整理成 Markdown 文件。不要联网，也不要调用天气工具。"
        task = {
            "message": message,
            "resolved_message": message,
            "intent_resolution": {
                "intent": "create_document",
                "standalone_request": message,
            },
        }
        plan = self.runtime._build_execution_plan(task, [], "md", False)
        self.assertEqual(plan["allowed_servers"], ["report"])
        self.assertTrue(self.runtime._web_search_explicitly_negated(message))

    def test_markdown_tool_artifact_satisfies_md_delivery_without_duplicate_call(self) -> None:
        artifact = {
            "kind": "markdown",
            "name": "AgentNexus模型连接验收.md",
            "download_url": "/api/artifacts/art_test/download",
        }
        self.assertTrue(
            self.runtime._artifact_matches_requested_format(artifact, "md")
        )

    def test_markdown_alias_extracts_file_content_for_semantic_verification(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "验收.md"
            path.write_text(
                "# AgentNexus 模型连接验收\n\n"
                "验收代号：星云-218\n\n模型连接测试通过。\n",
                encoding="utf-8",
            )
            with patch.object(self.runtime, "_artifact_file", return_value=path):
                readable, _, content = self.runtime._artifact_content_check(
                    {"id": "art_markdown"}, "markdown"
                )

        self.assertTrue(readable)
        self.assertIn("AgentNexus 模型连接验收", content)
        self.assertIn("星云-218", content)
        self.assertIn("模型连接测试通过", content)

    def test_goal_contract_does_not_freeze_generated_document_content(self) -> None:
        class DocumentContract:
            @staticmethod
            def tool_definition(*_: object) -> dict[str, object]:
                return {
                    "input_schema": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "content": {"type": "string"},
                            "format": {"type": "string"},
                        },
                    }
                }

        provenance = (
            ProvenanceRef(source_type="user_message", source_id="task_guard"),
        )
        goal = GoalSpec(
            goal_id="goal-document-contract",
            task_id="task_guard",
            objective=ObjectiveSpec(statement="生成 Markdown 文件"),
            inputs=(
                InputSpec(
                    key="title",
                    label="标题",
                    value="AgentNexus 模型连接验收",
                    provenance=provenance,
                ),
                InputSpec(
                    key="content",
                    label="正文来源",
                    value="模型连接测试通过。当前验收代号是星云-218。",
                    provenance=provenance,
                ),
                InputSpec(
                    key="format",
                    label="格式",
                    value="md",
                    provenance=provenance,
                ),
            ),
            deliverables=(
                DeliverableSpec(
                    id="primary_artifact", kind="artifact", format="md"
                ),
            ),
            acceptance=(
                AcceptanceCriterion(
                    id="artifact_format",
                    title="格式正确",
                    kind="format",
                    target="deliverable:primary_artifact",
                ),
            ),
        )
        self.runtime.contract_service = DocumentContract()

        constraints = self.runtime._goal_argument_constraints(
            goal, [("report", "generate_document")]
        )[("report", "generate_document")]

        self.assertEqual(
            {item["argument_path"] for item in constraints},
            {"title", "format"},
        )

    def test_markdown_tool_artifact_is_bound_to_md_goal_deliverable(self) -> None:
        goal = GoalSpec(
            goal_id="goal-markdown",
            task_id="task_guard",
            status="confirmed",
            objective=ObjectiveSpec(statement="生成 Markdown 文件"),
            deliverables=(
                DeliverableSpec(id="answer", kind="answer", format="text"),
                DeliverableSpec(
                    id="primary_artifact",
                    kind="artifact",
                    format="md",
                    download_required=True,
                ),
            ),
            acceptance=(
                AcceptanceCriterion(
                    id="goal_semantics",
                    title="结果与目标一致",
                    kind="semantic_match",
                    target="answer_and_artifacts",
                ),
            ),
            confirmation=ConfirmationSpec(
                status="auto_confirmed",
                mode="automatic",
                confidence=1.0,
            ),
        )
        self.runtime._execution_context.set(
            {
                "task_id": "task_guard",
                "run_id": "run_guard",
                "state": {"goal_spec_ref": {}, "tool_evidence": []},
            }
        )
        artifact = {
            "id": "art_markdown",
            "kind": "markdown",
            "name": "AgentNexus模型连接验收.md",
        }

        def candidate_artifact(
            item: dict[str, object], *, deliverable_id: str
        ) -> dict[str, object]:
            return {
                "id": item["id"],
                "deliverable_id": deliverable_id,
                "name": item["name"],
                "kind": item["kind"],
            }

        with patch.object(
            self.runtime,
            "_candidate_artifact_output",
            side_effect=candidate_artifact,
        ):
            candidate, _ = self.runtime._candidate_and_evidence(
                goal,
                "文档已生成。",
                [artifact],
                {},
            )

        self.assertEqual(
            candidate.artifacts[0].deliverable_id, "primary_artifact"
        )

    def test_plan_guard_emits_public_block_events_before_dispatch(self) -> None:
        task = {
            "message": "整理成 Word",
            "resolved_message": "把上一轮杭州行程整理成 Word 文档",
            "intent_resolution": {"intent": "create_document", "standalone_request": "把上一轮杭州行程整理成 Word 文档"},
        }
        plan = self.runtime._build_execution_plan(task, [], "docx", False)
        plan = self._contract_plan(plan, [("report", "generate_document")])
        plan["plan_id"] = "plan-docx-guard"
        with patch("app.services.agent_runtime.emit") as emit_mock:
            with self.assertRaises(ToolError):
                asyncio.run(
                    self.runtime._tool(
                        "task_guard",
                        "weather",
                        "forecast",
                        {"city": "杭州", "day": "today"},
                        plan=plan,
                    )
                )
        event_types = [call.args[1] for call in emit_mock.call_args_list]
        self.assertEqual(event_types, ["plan_check", "tool_blocked"])
        plan_check_data = emit_mock.call_args_list[0].args[4]
        blocked_data = emit_mock.call_args_list[1].args[4]
        self.assertFalse(plan_check_data["passed"])
        self.assertEqual(plan_check_data["reason"], "irrelevant_weather_tool")
        self.assertEqual(plan_check_data["source"], "execution_plan")
        self.assertEqual(blocked_data["reason"], "irrelevant_weather_tool")
        self.assertEqual(blocked_data["source"], "execution_plan")

    def test_plan_uses_two_levels_with_skill_children(self) -> None:
        task = {
            "message": "整理成 Word",
            "resolved_message": "把行程整理成 Word 文档",
            "intent_resolution": {"intent": "create_document", "standalone_request": "把行程整理成 Word 文档"},
        }
        skills = [{"id": "report_generation", "name": "报告生成", "required_mcps": ["report"]}]
        plan = self.runtime._build_execution_plan(task, skills, "docx", False)
        self.assertEqual([node["id"] for node in plan["nodes"]], ["understand", "prepare", "execute", "artifact", "validate"])
        self.assertEqual(plan["nodes"][0]["children"][0]["id"], "skill:report_generation")
        self.assertTrue(all("children" not in child for node in plan["nodes"] for child in node["children"]))
        self.assertTrue(all(node["status"] == "pending" for node in plan["nodes"]))
        self.assertEqual(plan["goal_confirmation"]["status"], "auto_confirmed")
        self.assertEqual([item["id"] for item in plan["acceptance_criteria"]], ["goal", "response", "format", "content", "download"])

    def test_non_document_plan_has_no_demo_specific_nodes(self) -> None:
        for message in ("春节还有多久", "安排宁波一日旅游路线"):
            with self.subTest(message=message):
                plan = self.runtime._build_execution_plan(
                    {
                        "message": message,
                        "resolved_message": message,
                        "intent_resolution": {"intent": "general", "standalone_request": message},
                    },
                    [],
                    "",
                    False,
                )
                self.assertEqual(
                    [node["title"] for node in plan["nodes"]],
                    ["确认当前目标与约束", "生成任务结果", "核对结果与当前目标"],
                )

    def test_common_download_formats_are_detected(self) -> None:
        cases = {
            "给我生成一个可下载的 Word 文档": "docx",
            "做成 PPTX 让我下载": "pptx",
            "导出为 Excel 文件": "xlsx",
            "生成一份 CSV 文件": "csv",
            "生成一个 Markdown 文件": "md",
            "制作一个 HTML 网页文档": "html",
        }
        for message, expected in cases.items():
            with self.subTest(message=message):
                self.assertEqual(self.runtime._requested_document_format(message), expected)

    def test_csv_plan_uses_spreadsheet_tool_and_csv_labels(self) -> None:
        task = {
            "message": "导出数据清单",
            "resolved_message": "将项目清单导出为 CSV 文件",
            "intent_resolution": {
                "intent": "export_data",
                "standalone_request": "将项目清单导出为 CSV 文件",
                "parameters": {"filename": "项目清单.csv", "topic": "项目清单"},
            },
        }
        plan = self.runtime._build_execution_plan(task, [], "csv", False)

        self.assertEqual(plan["output_format"], "csv")
        self.assertEqual(plan["allowed_servers"], ["spreadsheet"])
        self.assertEqual(plan["tool_node_id"], "artifact")
        self.assertIn("确认 CSV 交付要求", [node["title"] for node in plan["nodes"]])
        self.assertIn("生成可下载的 CSV 文件", [node["title"] for node in plan["nodes"]])
        self.assertEqual(plan["requirements"]["filename"], "项目清单.csv")
        self.assertTrue({"format", "content", "download", "filename"}.issubset(
            {item["id"] for item in plan["acceptance_criteria"]}
        ))
        plan = self._contract_plan(plan, [("spreadsheet", "create_excel")])
        self.runtime._validate_tool_against_plan(plan, "spreadsheet", "create_excel")
        with self.assertRaises(ToolError):
            self.runtime._validate_tool_against_plan(plan, "report", "generate_document")

    def test_ppt_plan_is_dynamic_and_uses_artifact_node(self) -> None:
        task = {
            "message": "做个 PPT",
            "resolved_message": "制作宁波旅游介绍 PPTX",
            "intent_resolution": {"intent": "create_presentation", "standalone_request": "制作宁波旅游介绍 PPTX", "parameters": {"filename": "宁波旅游.pptx", "topic": "宁波旅游", "chapters": ["景点", "行程"]}},
        }
        plan = self.runtime._build_execution_plan(task, [], "pptx", False)
        self.assertEqual(plan["tool_node_id"], "artifact")
        self.assertIn("生成可下载的 PowerPoint 文件", [node["title"] for node in plan["nodes"]])
        self.assertIn("report", plan["allowed_servers"])
        self.assertEqual(plan["requirements"]["filename"], "宁波旅游.pptx")
        self.assertTrue({"filename", "topic", "sections"}.issubset({item["id"] for item in plan["acceptance_criteria"]}))

    def test_document_output_requires_matching_artifact(self) -> None:
        plan = {"output_format": "docx"}
        failed = self.runtime._validate_output_against_plan(plan, "正文", [])
        self.assertFalse(failed["passed"])
        self.assertTrue(any(item["status"] == "failed" for item in failed["criteria"]))
        with TemporaryDirectory() as directory, patch(
            "app.services.mcp_gateway.ARTIFACT_DIR", Path(directory)
        ), patch("app.services.agent_runtime.ARTIFACT_DIR", Path(directory)):
            path = Path(directory) / "result.docx"
            document = Document()
            document.add_heading("验收测试文档", 0)
            document.add_paragraph("这是用于验证文档内容、文件格式和下载链路的有效正文内容。")
            document.save(path)
            passed = self.runtime._validate_output_against_plan(
                {"goal": "生成 Word 验收文档", "output_format": "docx"},
                "文档已经生成，可以下载。",
                [{"kind": "docx", "name": "result.docx", "relative_path": "result.docx", "download_url": "/api/artifacts/example/download"}],
            )
        self.assertTrue(passed["passed"])
        self.assertEqual(len(passed["criteria"]), 5)
        self.assertTrue(all(item["status"] == "passed" for item in passed["criteria"]))

    def test_attachment_document_must_preserve_sampled_source_content(self) -> None:
        with TemporaryDirectory() as directory, patch(
            "app.services.mcp_gateway.ARTIFACT_DIR", Path(directory)
        ), patch("app.services.agent_runtime.ARTIFACT_DIR", Path(directory)):
            path = Path(directory) / "result.docx"
            document = Document()
            document.add_heading("泛化结果", 0)
            document.add_paragraph("这是格式正确但完全没有引用用户附件的泛化内容。")
            document.save(path)
            artifact = {
                "kind": "docx",
                "name": "result.docx",
                "relative_path": "result.docx",
                "download_url": "/api/artifacts/example/download",
            }
            plan = {
                "goal": "根据附件生成 Word",
                "output_format": "docx",
                "attachment_requirements": ["唯一标记：NEXUS-E2E-20260812"],
            }
            failed = self.runtime._validate_output_against_plan(
                plan, "文档已经生成。", [artifact]
            )
            self.assertFalse(failed["passed"])
            source_check = next(
                item for item in failed["criteria"] if item["id"] == "source_consistency"
            )
            self.assertEqual(source_check["status"], "failed")

            document.add_paragraph("唯一标记：NEXUS-E2E-20260812")
            document.save(path)
            passed = self.runtime._validate_output_against_plan(
                plan, "文档已经生成。", [artifact]
            )
            self.assertTrue(passed["passed"])
            self.assertEqual(
                next(item for item in passed["criteria"] if item["id"] == "source_consistency")["status"],
                "passed",
            )

    def test_csv_output_is_read_and_validated_before_delivery(self) -> None:
        with TemporaryDirectory() as directory, patch(
            "app.services.mcp_gateway.ARTIFACT_DIR", Path(directory)
        ), patch("app.services.agent_runtime.ARTIFACT_DIR", Path(directory)):
            path = Path(directory) / "项目清单.csv"
            path.write_text(
                "项目,状态,说明\n智枢,进行中,CSV 交付已验证\n",
                encoding="utf-8-sig",
            )
            artifact = {
                "kind": "csv",
                "name": "项目清单.csv",
                "relative_path": "项目清单.csv",
                "download_url": "/api/artifacts/example/download",
            }
            plan = {
                "goal": "生成项目清单 CSV",
                "output_format": "csv",
                "requirements": {"filename": "项目清单.csv", "topic": "智枢"},
            }
            passed = self.runtime._validate_output_against_plan(
                plan, "CSV 已生成，可以下载。", [artifact]
            )
            path.write_text("项目,状态\n智枢\n", encoding="utf-8-sig")
            failed = self.runtime._validate_output_against_plan(
                plan, "CSV 已生成，可以下载。", [artifact]
            )

        self.assertTrue(passed["passed"])
        self.assertTrue(all(item["status"] == "passed" for item in passed["criteria"]))
        self.assertFalse(failed["passed"])
        self.assertEqual(
            next(item for item in failed["criteria"] if item["id"] == "content")["status"],
            "failed",
        )

    def test_attachment_acceptance_requirements_prefer_distinctive_lines(self) -> None:
        requirements = self.runtime._attachment_acceptance_requirements(
            "--- input.docx ---\n项目介绍\n唯一标记：NEXUS-E2E-20260812\n需求：生成 Word"
        )
        self.assertEqual(requirements[0], "唯一标记：NEXUS-E2E-20260812")
        self.assertIn("需求：生成 Word", requirements)

    def test_attachment_pdf_source_consistency_uses_extracted_pdf_text(self) -> None:
        with TemporaryDirectory() as directory, patch(
            "app.services.mcp_gateway.ARTIFACT_DIR", Path(directory)
        ), patch("app.services.agent_runtime.ARTIFACT_DIR", Path(directory)):
            path = Path(directory) / "result.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.drawString(72, 760, "Attachment marker NEXUS-PDF-E2E-20260812")
            pdf.drawString(72, 740, "Generated document retains the uploaded source.")
            pdf.save()
            result = self.runtime._validate_output_against_plan(
                {
                    "goal": "根据附件生成 PDF",
                    "output_format": "pdf",
                    "attachment_requirements": ["NEXUS-PDF-E2E-20260812"],
                },
                "PDF 已生成。",
                [{
                    "kind": "pdf",
                    "name": "result.pdf",
                    "relative_path": "result.pdf",
                    "download_url": "/api/artifacts/example/download",
                }],
            )
        self.assertTrue(result["passed"])
        self.assertEqual(
            next(item for item in result["criteria"] if item["id"] == "source_consistency")["status"],
            "passed",
        )


if __name__ == "__main__":
    unittest.main()
